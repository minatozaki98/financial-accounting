from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
STRICT_TEMPLATE = Path(r"C:\Users\Asus\Downloads\conference-template-letter.docx")
TEMPLATE = ROOT / "tmp" / "ieee-template-convert" / "conference-template-letter-transitional.docx"
SOURCE = ROOT / "Document" / "paper" / "incit2026-ieee-submission.docx"
OUT_REPO = ROOT / "Document" / "paper" / "incit2026-ieee-submission-ieee-template.docx"
OUT_DOWNLOADS = Path(r"C:\Users\Asus\Downloads\incit2026 ieee submission IEEE template rebuilt.docx")

TITLE = (
    "Evaluating ChatGPT (Codex 5.4) for Security, Performance, and "
    "Code-Quality Improvement of a Financial Accounting REST API: "
    "A Branch-Isolated Baseline Study"
)
AUTHORS = "Zaw Ye Htut Ko, Dr Darun Kesrarat"
FIGURES = [
    ROOT / "Document" / "paper" / "figures" / "workflow-phase1.png",
    ROOT / "Document" / "paper" / "figures" / "workflow-phase2.png",
    ROOT / "Document" / "paper" / "figures" / "workflow-phase3.png",
    ROOT / "Document" / "paper" / "figures" / "workflow-phase4.png",
]
COLUMN_WIDTH_DXA = 4968


def style_name(doc, preferred, fallback=None):
    try:
        doc.styles[preferred]
        return preferred
    except KeyError:
        return fallback


def clear_document_body(doc):
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_ieee_page(section, cols=1):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.62 if cols == 1 else 0.63)
    section.right_margin = Inches(0.62 if cols == 1 else 0.63)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    sect_pr = section._sectPr
    for old in list(sect_pr):
        if old.tag == qn("w:cols"):
            sect_pr.remove(old)
    cols_el = OxmlElement("w:cols")
    cols_el.set(qn("w:num"), str(cols))
    cols_el.set(qn("w:space"), "360" if cols == 2 else "720")
    sect_pr.append(cols_el)


def set_doc_metadata(doc):
    props = doc.core_properties
    props.author = AUTHORS
    props.title = TITLE
    props.subject = "InCIT 2026 IEEE conference paper"
    props.keywords = "ChatGPT; LLM; REST API; SonarQube; OWASP ZAP; JMeter"
    props.comments = ""


def add_styled_paragraph(doc, text, style, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.add_run(text)
    return p


def add_body_paragraph(doc, text):
    p = doc.add_paragraph(style=style_name(doc, "Body Text", "Normal"))
    lead_match = re.match(
        r"^(Why this work matters|Outcome summary|Contributions|What is [^?]+\?|"
        r"Why this dimension\?|Experiment\.|Reading the table\.|Findings\.|"
        r"Internal\.|External\.|Construct\.)(\s+)(.*)$",
        text,
    )
    if lead_match:
        lead, space, rest = lead_match.groups()
        r = p.add_run(lead + space)
        r.bold = True
        p.add_run(rest)
        return p

    italic_leads = [
        "LLM code generation and productivity.",
        "LLMs and security.",
        "APIs, static analysis, and LLMs for SE.",
    ]
    for lead in italic_leads:
        if text.startswith(lead):
            r = p.add_run(lead + " ")
            r.italic = True
            p.add_run(text[len(lead) :].strip())
            return p

    p.add_run(text)
    return p


def strip_template_number(text, pattern):
    return re.sub(pattern, "", text).strip()


def iter_source_blocks(source_doc):
    body = source_doc._body._element
    for child in body:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in child.iter() if t.tag.rsplit("}", 1)[-1] == "t")
            has_drawing = any(x.tag.rsplit("}", 1)[-1] == "drawing" for x in child.iter())
            yield ("p", text.strip(), has_drawing)
        elif tag == "tbl":
            rows = []
            for tr in [x for x in child if x.tag.rsplit("}", 1)[-1] == "tr"]:
                row = []
                for tc in [x for x in tr if x.tag.rsplit("}", 1)[-1] == "tc"]:
                    texts = [
                        t.text or ""
                        for t in tc.iter()
                        if t.tag.rsplit("}", 1)[-1] == "t"
                    ]
                    row.append("".join(texts).strip())
                rows.append(row)
            yield ("tbl", rows, False)


def table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def set_table_width(table, width_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(width_dxa))

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_dxa):
    cell.width = Inches(width_dxa / 1440)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(width_dxa))


def set_table_grid(table, widths):
    old_grid = table._tbl.tblGrid
    if old_grid is not None:
        table._tbl.remove(old_grid)
    tbl_grid = OxmlElement("w:tblGrid")
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    table._tbl.insert(1, tbl_grid)


def table_column_widths(rows):
    cols = max(len(r) for r in rows)
    header = rows[0] if rows else []
    if cols == 4 and header and header[0].startswith("Endpoint"):
        weights = [2.0, 0.75, 0.75, 0.65]
    elif cols == 4:
        weights = [2.0, 0.7, 0.7, 0.45]
    elif cols == 5:
        weights = [1.15, 0.8, 0.8, 0.8, 0.8]
    elif cols == 6:
        weights = [1.0, 0.8, 0.8, 0.6, 0.55, 0.55]
    else:
        weights = [1.0] * cols
    total = sum(weights)
    widths = [int(COLUMN_WIDTH_DXA * w / total) for w in weights]
    widths[-1] += COLUMN_WIDTH_DXA - sum(widths)
    return widths


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    widths = table_column_widths(rows)
    set_table_width(table, COLUMN_WIDTH_DXA)
    set_table_grid(table, widths)
    table_borders(table)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = style_name(doc, "table col head" if i == 0 else "table copy", "Normal")
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            run.bold = i == 0
            if len(widths) >= 5:
                run.font.size = Pt(7)


def add_figure(doc, figure_path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(figure_path), width=Inches(2.0))


def build():
    if not TEMPLATE.exists():
        raise FileNotFoundError(
            f"{TEMPLATE} is required. Create it by opening {STRICT_TEMPLATE} "
            "in Word and saving it as a standard DOCX."
        )
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    source_doc = Document(str(SOURCE))
    doc = Document(str(TEMPLATE))
    clear_document_body(doc)
    set_doc_metadata(doc)

    set_ieee_page(doc.sections[0], cols=1)
    add_styled_paragraph(doc, TITLE, style_name(doc, "paper title", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_paragraph(doc, AUTHORS, style_name(doc, "Author", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_ieee_page(body_section, cols=2)

    paragraphs = [block for block in iter_source_blocks(source_doc)]
    source_texts = [text for kind, text, _ in paragraphs if kind == "p" and text]
    abstract_text = ""
    keywords_text = ""
    if "Abstract—" in source_texts:
        idx = source_texts.index("Abstract—")
        abstract_text = source_texts[idx + 1] if idx + 1 < len(source_texts) else ""
    for text in source_texts:
        if text.startswith("Index Terms—") or text.startswith("Keywords—"):
            keywords_text = text.split("—", 1)[1].strip()
            break

    p = doc.add_paragraph(style=style_name(doc, "Abstract", "Normal"))
    r = p.add_run("Abstract—")
    r.bold = True
    r.italic = True
    p.add_run(abstract_text)

    p = doc.add_paragraph(style=style_name(doc, "Keywords", "Normal"))
    r = p.add_run("Keywords—")
    r.bold = True
    r.italic = True
    p.add_run(keywords_text)

    started = False
    refs = False
    fig_index = 0
    for kind, payload, has_drawing in paragraphs:
        if kind == "p":
            text = payload
            if not started:
                if text == "I. INTRODUCTION":
                    started = True
                else:
                    continue
            if not text and has_drawing:
                if fig_index < len(FIGURES):
                    add_figure(doc, FIGURES[fig_index])
                    fig_index += 1
                continue
            if not text:
                continue

            if text in {"Abstract—", abstract_text} or text.startswith("Index Terms—"):
                continue

            if text.lower() == "references":
                refs = True
                add_styled_paragraph(doc, "REFERENCES", style_name(doc, "Normal", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)
            elif refs or re.match(r"^\[\d+\]", text):
                ref_text = strip_template_number(text, r"^\[\d+\]\s*")
                add_styled_paragraph(doc, ref_text, style_name(doc, "references", "Normal"))
            elif text.startswith("TABLE "):
                if re.match(r"^TABLE\s+II\.", text):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                elif re.match(r"^TABLE\s+III\.", text):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
                table_text = strip_template_number(text, r"^TABLE\s+[IVXLCDM]+\.\s*")
                p = add_styled_paragraph(doc, table_text, style_name(doc, "table head", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
            elif re.match(r"^Fig\.\s+\d+\.", text):
                caption_text = strip_template_number(text, r"^Fig\.\s+\d+\.\s*")
                add_styled_paragraph(doc, caption_text, style_name(doc, "figurecaption", "Normal"))
            elif re.match(r"^(I|II|III|IV|V|VI)\.", text):
                heading_text = strip_template_number(text, r"^[IVXLCDM]+\.\s*")
                add_styled_paragraph(doc, heading_text, style_name(doc, "heading 1", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)
            elif re.match(r"^[A-Z]\.\s", text):
                heading_text = strip_template_number(text, r"^[A-Z]\.\s*")
                add_styled_paragraph(doc, heading_text, style_name(doc, "heading 2", "Normal"))
            elif text in {"ACKNOWLEDGMENT AND AI USAGE DISCLOSURE"}:
                p = add_styled_paragraph(doc, text, style_name(doc, "Normal", "Normal"), WD_ALIGN_PARAGRAPH.CENTER)
                for run in p.runs:
                    run.bold = False
            else:
                add_body_paragraph(doc, text)
        elif kind == "tbl" and started:
            add_table(doc, payload)

    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_REPO))
    shutil.copy2(OUT_REPO, OUT_DOWNLOADS)
    print(OUT_REPO)
    print(OUT_DOWNLOADS)


if __name__ == "__main__":
    build()
