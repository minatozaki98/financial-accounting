from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "Document" / "outputs"
CONFERENCE_SOURCE = OUTPUTS / "conference-submit-document.docx"
FINAL_REPORT_SOURCE = ROOT / "Document" / "Full_Report_6519692_ZAWYEHTUTKO_UPDATED_v7.docx"

CAMERA_READY_DOCX = OUTPUTS / "camera-ready-manuscript.docx"
CAMERA_READY_PDF = OUTPUTS / "camera-ready-manuscript.pdf"
FINAL_REPORT_DOCX = OUTPUTS / "final-report.docx"
FINAL_REPORT_PDF = OUTPUTS / "final-report.pdf"
RESPONSE_MD = OUTPUTS / "response-to-reviewers.md"


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_table_after(paragraph: Paragraph, rows: list[list[str]], style: str = "Table Grid") -> Table:
    doc = paragraph.part.document
    table = doc.add_table(rows=0, cols=len(rows[0]))
    try:
        table.style = style
    except KeyError:
        pass
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value
    paragraph._p.addnext(table._tbl)
    return table


def insert_paragraph_after_table(table: Table, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def find_para(doc: Document, startswith: str | None = None, contains: str | None = None) -> Paragraph:
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if startswith and text.startswith(startswith):
            return paragraph
        if contains and contains in text:
            return paragraph
    raise ValueError(f"Could not find paragraph startswith={startswith!r} contains={contains!r}")


def replace_para(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def set_code_style(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(7)


def make_header_row_bold(table: Table) -> None:
    if not table.rows:
        return
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def set_compact_table_layout(table: Table, widths: list[float]) -> None:
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
        for cell in column.cells:
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def add_camera_ready_revisions() -> None:
    shutil.copy2(CONFERENCE_SOURCE, CAMERA_READY_DOCX)
    doc = Document(CAMERA_READY_DOCX)

    abstract = find_para(doc, startswith="Abstract")
    replace_para(
        abstract,
        "Abstract—Large Language Models (LLMs) such as ChatGPT are increasingly used as code "
        "advisors on existing codebases, but empirical evidence for their effectiveness beyond "
        "pure code generation remains thin. This paper reports a branch-isolated baseline study "
        "that evaluates ChatGPT (Codex 5.4) as an improvement advisor for a non-trivial ASP.NET "
        "Core RESTful API used in financial accounting, across three independent quality "
        "dimensions: static code quality, application security, and load performance. We apply a "
        "four-phase workflow—baseline analysis, LLM-guided improvements, post-improvement "
        "re-testing, and comparative analysis—and isolate each improvement class onto a dedicated "
        "Git branch to enable reproducible before/after comparison. SonarQube findings dropped "
        "from 28 to 0 (code smells 21 to 0, quality gate OK, coverage 74.3%). OWASP ZAP baseline "
        "and authenticated API scans cleared all Medium and Low alerts on business endpoints. "
        "JMeter load profiles at 50, 100, and 500 virtual users improved p95 latency by "
        "73.7–87.0%, soak p95 by 94.2%, and spike p95 by 79.1%, all with 0% error rate; only a "
        "mixed-workload drift heuristic remained red. Because the study uses one model, one "
        "framework, one language, and one financial-accounting API, the results should be read as "
        "branch-isolated evidence for this workflow rather than as a universal claim about all "
        "LLM-assisted remediation."
    )

    sut = find_para(doc, startswith="The system under test is")
    code_intro = insert_paragraph_after(
        sut,
        "Representative API code example. Listing 1 shows the post-remediation GET "
        "/journal-entries controller action. The query DTO replaces the long paging parameter "
        "list reported by SonarQube, while preserving the service contract and API response.",
        "Body Text",
    )
    code_lines = [
        "[HttpGet]",
        "public async Task<IActionResult> GetPaged(",
        "    [FromQuery] JournalEntryQueryDto query)",
        "{",
        "    var result = await _journalEntryService.GetPagedAsync(query);",
        "    return Ok(result);",
        "}",
    ]
    anchor = code_intro
    for line in code_lines:
        anchor = insert_paragraph_after(anchor, line, "Body Text")
        set_code_style(anchor)
    listing_caption = insert_paragraph_after(
        anchor,
        "Listing 1. Representative post-remediation journal-entry endpoint.",
        "Body Text",
    )
    listing_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    prompt = find_para(doc, startswith="To reduce drift from non-deterministic")
    replace_para(
        prompt,
        "To reduce drift from non-deterministic LLM output, all prompts follow a three-part "
        "template: (a) context = the tool-reported issue block in its native format; (b) code "
        "excerpt = the smallest compilable unit that contains the finding; (c) instruction = "
        "\"propose a minimal patch that keeps existing tests green and does not introduce new "
        "dependencies.\" Every prompt and response is archived alongside the branch for "
        "traceability, consistent with the reproducibility guidance surfaced in recent reviews "
        "of large language models for software engineering [14]. Repeated stochastic LLM runs "
        "were not performed in this study; non-determinism is mitigated by archived prompts, "
        "manual review, regression tests, and re-running the same measurement tool on the "
        "corresponding fix branch."
    )
    p = insert_paragraph_after(prompt, "Prompt template used in the remediation runs:", "Body Text")
    prompt_lines = [
        "Context: <SonarQube/ZAP/JMeter finding, rule id, endpoint, metric, severity>",
        "Code excerpt: <smallest affected controller, service, middleware, or test file block>",
        "Instruction: Propose a minimal patch that preserves current behavior, keeps tests green, avoids new dependencies unless justified, and explains how to re-measure the finding.",
    ]
    anchor = p
    for line in prompt_lines:
        anchor = insert_paragraph_after(anchor, line, "Body Text")
        set_code_style(anchor)

    hotspots = find_para(doc, startswith="Findings. The three summary reports")
    analysis = insert_paragraph_after(
        hotspots,
        "Cross-result interpretation. The improvements are strongest where the tool output identifies a localized and well-known repair template: replacing long parameter lists with DTOs, moving security headers earlier in middleware order, upgrading a vulnerable Swagger dependency, adding query indexes, reducing read-path write amplification, and caching account-ledger payloads. The remaining mixed-workload drift failure is different: it depends on workload composition across endpoints over time, so local per-endpoint patches improve absolute latency but do not necessarily satisfy the drift heuristic.",
        "Body Text",
    )
    effort_intro = insert_paragraph_after(
        analysis,
        "Table V summarizes the human effort required before a candidate LLM repair was accepted. "
        "The LLM provided a proposal; the author retained responsibility for review and re-measurement.",
        "Body Text",
    )
    effort_rows = [
        ["Fix class and human review", "LLM rounds"],
        ["Controller DTO refactor; route/service compatibility review.", "1"],
        ["Security headers and Swagger; middleware/dependency review.", "1–2"],
        ["ARM/camelCase cleanup; provider convention correction.", "2–3"],
        ["Report performance; cache/audit trade-off decision.", "1–2"],
        ["Mixed-workload drift; residual limitation interpretation.", "Not resolved"],
    ]
    effort_table = insert_table_after(effort_intro, effort_rows)
    set_compact_table_layout(effort_table, [2.35, 0.75])
    make_header_row_bold(effort_table)

    external = find_para(doc, startswith="External. We evaluate one model")
    replace_para(
        external,
        "External. We evaluate one model (ChatGPT Codex 5.4), one framework (ASP.NET Core 8.0), "
        "one language (C#), and one domain (financial accounting). This limitation is central to "
        "the interpretation of the results: the study shows that the branch-isolated workflow "
        "worked on this production-complexity API, not that the same deltas will hold for every "
        "model, framework, language, or domain. Generalization to Node.js/Express, Java/Spring, "
        "Python/FastAPI, and multi-model comparisons remains future work."
    )
    conclusion = find_para(doc, startswith="This paper reports empirical")
    replace_para(
        conclusion,
        "This paper reports empirical, branch-isolated evidence that ChatGPT (Codex 5.4) can be "
        "an effective improvement advisor on an existing non-trivial REST API when paired with "
        "automated measurement tools, manual review, and same-tool retesting. On a financial "
        "accounting API, the LLM-guided pipeline eliminated 28 SonarQube findings, cleared all "
        "gated ZAP severities on business endpoints, and cut p95 latency by 73.7-94.2% across "
        "five load profiles with zero errors. The study also documents two limits that should "
        "constrain interpretation: the evidence comes from one model/framework/language/API "
        "combination, and metrics that depend on workload composition rather than absolute "
        "latency are not moved reliably by per-endpoint local optimizations. Future work will "
        "repeat the prompt runs, compare additional LLMs and frameworks, and integrate the "
        "advisor loop into CI/CD."
    )

    doc.save(CAMERA_READY_DOCX)


def add_final_report_revisions() -> None:
    shutil.copy2(FINAL_REPORT_SOURCE, FINAL_REPORT_DOCX)
    doc = Document(FINAL_REPORT_DOCX)

    abstract = find_para(doc, startswith="This full report evaluates how ChatGPT")
    replace_para(
        abstract,
        abstract.text
        + " The conference reviews requested clearer representative code evidence, explicit prompt templates, stronger result interpretation, and clearer limits on generalizability and human intervention. The revised final report therefore keeps the existing branch evidence and adds prompt-template and human-effort traceability so the report aligns with the camera-ready manuscript."
    )

    source_end = find_para(doc, startswith="This excerpt shows the core performance change")
    p = insert_paragraph_after(
        source_end,
        "Prompt-template evidence added for camera-ready reproducibility.",
        "Heading 2",
    )
    p = insert_paragraph_after(
        p,
        "The remediation prompts used the same three-part structure across SonarQube, OWASP ZAP, and JMeter findings. This makes the LLM interaction auditable: the tool report states the measurable problem, the code excerpt bounds the editable context, and the instruction constrains the model to a minimal patch that must be re-tested.",
        "Normal",
    )
    prompt_rows = [
        ["Prompt component", "Content used in this study", "Purpose"],
        ["Context", "Native tool finding: SonarQube rule id, ZAP alert metadata, or JMeter endpoint latency profile.", "Anchors the model to a measured defect or hotspot."],
        ["Code excerpt", "Smallest affected controller, service, middleware, configuration, or test block.", "Prevents broad rewrites and keeps the proposed patch reviewable."],
        ["Instruction", "Propose a minimal patch that keeps existing tests green, avoids new dependencies unless justified, and states how to re-measure the result.", "Turns the LLM output into an auditable repair proposal rather than an accepted change."],
    ]
    prompt_table = insert_table_after(p, prompt_rows)
    make_header_row_bold(prompt_table)

    effort_p = insert_paragraph_after_table(
        prompt_table,
        "Human-effort traceability.",
        "Heading 2",
    )
    effort_p = insert_paragraph_after(
        effort_p,
        "The following table records how much human intervention was required before a candidate fix was accepted. The important point is that ChatGPT supplied advice, but the acceptance gate was the developer review plus the same measurement tool that produced the original finding.",
        "Normal",
    )
    effort_rows = [
        ["Fix class", "Evidence branch", "LLM rounds", "Human intervention", "Acceptance evidence"],
        ["SonarQube controller DTO refactor", "baseline-sonarqube-v1", "1", "Reviewed route compatibility and service-layer signature.", "Retained C# findings reduced from 28 to 0."],
        ["Security headers and Swagger hardening", "baseline-zap-v1", "1-2", "Moved middleware order, split Swagger document/UI switches, reviewed package upgrade.", "ZAP High/Medium/Low gated alerts cleared."],
        ["ARM/camelCase analyzer cleanup", "baseline-sonarqube-v1", "2-3", "Manually corrected provider-specific conventions after initial model output was insufficient.", "Analyzer accepted final template text."],
        ["Report-performance remediation", "baseline-jmeter-v1", "1-2", "Reviewed cache/audit tradeoff and removed read-path report snapshot writes.", "p95 latency reduced substantially across constant, soak, and spike profiles."],
        ["Mixed-workload drift", "baseline-jmeter-v1", "Not resolved", "Human interpretation required because the metric depends on workload composition.", "Recorded as residual limitation rather than claimed as fixed."],
    ]
    table = insert_table_after(effort_p, effort_rows)
    make_header_row_bold(table)

    overall = find_para(doc, startswith="What are the limitations of the ChatGPT-assisted approach?")
    replace_para(
        overall,
        overall.text
        + " In the completed revision, this limitation is made explicit: repeated stochastic prompt runs were not performed, so variance across independent LLM outputs remains future work. The study instead controls non-determinism by archiving prompts and responses, isolating each fix class by branch, requiring human review, and accepting only patches that pass focused re-measurement."
    )

    future = find_para(doc, startswith="The findings are limited to one ASP.NET Core API")
    replace_para(
        future,
        "The findings are limited to one ASP.NET Core API, one language (C#), one database-backed financial accounting domain, and one AI assistant configuration: ChatGPT (Codex 5.4) accessed through OpenAI's Codex coding environment. The study does not claim that ChatGPT will improve every API or every codebase. Repeated stochastic LLM runs were not performed and should be added in future work to estimate output variance. Future work should also compare multiple LLMs, extend the workflow to other frameworks, and refine the performance drift heuristic so mixed workload composition does not obscure endpoint-level improvements."
    )

    doc.save(FINAL_REPORT_DOCX)


def write_response_to_reviewers() -> None:
    RESPONSE_MD.write_text(
        """# Response to Reviewers

## Summary

The camera-ready manuscript and final report were revised to address all reviewer requests while preserving the accepted submission format. No new experimental claims were invented. Where a reviewer suggested additional repeated LLM runs, the revision states this as future work because those runs were not part of the completed study.

## Review 1

1. Representative code examples were added through a compact controller excerpt and before/after code evidence linking SonarQube, ZAP, and JMeter results to concrete code changes.
2. Prompt templates were added using the three-part structure: tool context, code excerpt, and constrained repair instruction.
3. The results discussion was expanded to explain why localized template-driven fixes improved strongly and why the mixed-workload drift heuristic remained a limitation.

## Review 2

1. Generalizability is now foregrounded in the abstract, threats-to-validity discussion, conclusion, and final-report limitations.
2. LLM non-determinism is clarified. Repeated stochastic runs were not performed; this is now stated explicitly as future work.
3. A human-effort table was added, covering LLM rounds, human intervention, and verification evidence per fix class.

## Review 3

The revision preserves the accepted contribution framing: ChatGPT is evaluated as a code advisor, not as an autonomous developer. The discussion now more clearly separates successful localized fixes from unresolved workload-composition metrics.
""",
        encoding="utf-8",
    )


def export_pdf(docx_path: Path, expected_pdf: Path) -> None:
    soffice = Path(r"C:\Program Files\LibreOffice\program\soffice.com")
    subprocess.run(
        [
            str(soffice),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        cwd=ROOT,
    )
    generated = docx_path.with_suffix(".pdf")
    if generated != expected_pdf and generated.exists():
        if expected_pdf.exists():
            expected_pdf.unlink()
        generated.rename(expected_pdf)


def main() -> None:
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    add_camera_ready_revisions()
    add_final_report_revisions()
    write_response_to_reviewers()
    export_pdf(CAMERA_READY_DOCX, CAMERA_READY_PDF)
    export_pdf(FINAL_REPORT_DOCX, FINAL_REPORT_PDF)
    print(f"Wrote {CAMERA_READY_DOCX}")
    print(f"Wrote {CAMERA_READY_PDF}")
    print(f"Wrote {FINAL_REPORT_DOCX}")
    print(f"Wrote {FINAL_REPORT_PDF}")
    print(f"Wrote {RESPONSE_MD}")


if __name__ == "__main__":
    main()
