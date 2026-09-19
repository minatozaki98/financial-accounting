from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
CONTENT_PATH = ROOT / "scripts" / "defense" / "defense_content.json"
OUTPUT_PATH = (
    ROOT
    / "Document"
    / "outputs"
    / "final-defense"
    / "G6519692_Zaw_Ye_Htut_Ko_Defense_Preparation_Guide.docx"
)

COLORS = {
    "text": "20262E",
    "muted": "66717E",
    "methodology": "0B6670",
    "security": "B93A3A",
    "performance": "C47B12",
    "verified": "2F7D4A",
    "line": "D9DEE5",
    "surface": "F2F5F7",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
    else:
        tc_pr.remove(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    following_tags = {
        qn("w:noWrap"),
        qn("w:tcMar"),
        qn("w:textDirection"),
        qn("w:tcFitText"),
        qn("w:vAlign"),
        qn("w:hideMark"),
    }
    insertion_index = next(
        (index for index, child in enumerate(tc_pr) if child.tag in following_tags),
        len(tc_pr),
    )
    tc_pr.insert(insertion_index, shading)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    tc_pr.remove(margins)
    following_tags = {
        qn("w:textDirection"),
        qn("w:tcFitText"),
        qn("w:vAlign"),
        qn("w:hideMark"),
    }
    insertion_index = next(
        (index for index, child in enumerate(tc_pr) if child.tag in following_tags),
        len(tc_pr),
    )
    tc_pr.insert(insertion_index, margins)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, result, end):
        run._r.append(element)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(COLORS["text"])
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")

    for style_name, size, color in (
        ("Title", 27, COLORS["text"]),
        ("Subtitle", 14, COLORS["muted"]),
        ("Heading 1", 18, COLORS["methodology"]),
        ("Heading 2", 14, COLORS["text"]),
        ("Heading 3", 11.5, COLORS["security"]),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    for current_section in doc.sections:
        header = current_section.header.paragraphs[0]
        header.text = "FINAL THESIS DEFENSE | G6519692"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(COLORS["muted"])
        add_page_number(current_section.footer.paragraphs[0])


def add_rule(paragraph, color: str = COLORS["methodology"]) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    following_tags = {
        qn("w:shd"),
        qn("w:tabs"),
        qn("w:spacing"),
        qn("w:ind"),
        qn("w:jc"),
        qn("w:textDirection"),
        qn("w:textAlignment"),
        qn("w:outlineLvl"),
        qn("w:rPr"),
        qn("w:sectPr"),
        qn("w:pPrChange"),
    }
    insertion_index = next(
        (index for index, child in enumerate(p_pr) if child.tag in following_tags),
        len(p_pr),
    )
    p_pr.insert(insertion_index, borders)


def add_label_paragraph(doc: Document, label: str, text: str, color: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    label_run.font.color.rgb = RGBColor.from_string(color)
    paragraph.add_run(text)


def add_callout(doc: Document, title: str, text: str, color: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    cell.width = Inches(6.65)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "EEF3F5")
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run(f"{title}: ")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(color)
    paragraph.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc: Document, content: dict) -> None:
    doc.add_paragraph("DEFENSE PREPARATION GUIDE", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph(content["metadata"]["title"], style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(55)
    title.paragraph_format.space_after = Pt(28)
    add_rule(title, COLORS["performance"])

    for line in (
        f"{content['metadata']['student']} | {content['metadata']['studentId']}",
        content["metadata"]["degree"],
        f"Advisor: {content['metadata']['advisor']}",
        content["metadata"]["dateLabel"],
    ):
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.bold = line.startswith(content["metadata"]["student"])

    doc.add_paragraph()
    add_callout(
        doc,
        "Defense principle",
        "Use the model to propose. Use evidence to decide.",
        COLORS["verified"],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_how_to_use(doc: Document) -> None:
    doc.add_heading("How to Use This Guide", level=1)
    add_callout(
        doc,
        "Your target",
        "Deliver the 22 main slides in 27-28 minutes, then use appendix slides only when a professor asks for detailed evidence.",
        COLORS["methodology"],
    )
    items = [
        "Do not memorize every sentence. Memorize each slide's key message and first sentence.",
        "Use Presenter View so the current notes and the next slide remain visible to you.",
        "Explain charts from left to right: baseline, after result, interpretation, limitation.",
        "When interrupted, answer the question directly and return with: 'I will connect that answer to the next result.'",
        "Keep an offline backup of the PPTX, presentation PDF, guide PDF, and evidence screenshots on a USB drive and cloud storage.",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Opening Script", level=2)
    doc.add_paragraph(
        "Good morning, respected chair and committee members. My name is Zaw Ye Htut Ko, student ID G6519692. "
        "Today I will present my thesis on improving a financial ASP.NET Core RESTful API using ChatGPT as a code-improvement advisor. "
        "The central idea is that the model proposes changes, but human review and repeatable tool evidence decide whether those changes are accepted."
    )
    doc.add_heading("Closing Script", level=2)
    doc.add_paragraph(
        "In conclusion, the study found measurable gains across static quality, security gates, and tail latency, while also identifying limitations in drift measurement and model comparison. "
        "The strongest contribution is the controlled audit-to-fix workflow. Thank you for your attention. I welcome your questions."
    )


def add_timing_map(doc: Document, content: dict) -> None:
    doc.add_heading("Presentation Timing Map", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ("Slide", "Section", "Title", "Target")
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, COLORS["methodology"])
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, slide in enumerate(content["slides"], start=1):
        row = table.add_row().cells
        values = (str(index), slide["section"], slide["title"], f"{slide['durationSeconds']} sec")
        for cell, value in zip(row, values):
            cell.text = value
            set_cell_margins(cell)
    total = sum(slide["durationSeconds"] for slide in content["slides"])
    minutes, seconds = divmod(total, 60)
    add_callout(doc, "Planned duration", f"{minutes} minutes {seconds} seconds, leaving time for pauses and brief clarification.", COLORS["verified"])


def add_slide_scripts(doc: Document, content: dict) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Slide-by-Slide Script", level=1)
    doc.add_paragraph(
        "Use these scripts for rehearsal. During the defense, look at the committee and use the notes as prompts rather than reading every word."
    )
    for index, slide in enumerate(content["slides"], start=1):
        if index > 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"Slide {index}: {slide['title']}", level=2)
        add_label_paragraph(doc, "Time: ", f"{slide['durationSeconds']} seconds", COLORS["performance"])
        add_label_paragraph(doc, "Key message: ", slide["keyMessage"], COLORS["verified"])
        doc.add_heading("What to say", level=3)
        doc.add_paragraph(slide["speakerNotes"])
        add_label_paragraph(doc, "Transition: ", slide["transition"], COLORS["methodology"])
        add_label_paragraph(doc, "Do not overclaim: ", slide["overclaimWarning"], COLORS["security"])
        add_label_paragraph(doc, "Evidence: ", " | ".join(slide["evidenceRefs"]), COLORS["muted"])


def add_question_bank(doc: Document, content: dict) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Professor Question Bank", level=1)
    doc.add_paragraph(
        "Answer first, then support the answer. A strong defense answer is normally shorter than one minute unless the professor asks for more detail."
    )
    names = {
        "chair": "Chair and Contribution",
        "methodology": "Research Methodology",
        "security": "Security",
        "performance": "Performance and Systems",
        "devilsAdvocate": "Devil's Advocate",
    }
    current = None
    number = 0
    for question in content["questions"]:
        if question["perspective"] != current:
            current = question["perspective"]
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(names[current], level=2)
        number += 1
        doc.add_heading(f"Question {number}: {question['question']}", level=3)
        add_label_paragraph(doc, "Short answer: ", question["shortAnswer"], COLORS["verified"])
        add_label_paragraph(doc, "Follow-up answer: ", question["followUpAnswer"], COLORS["methodology"])
        add_label_paragraph(doc, "Supporting slide: ", question["slideRef"], COLORS["performance"])
        add_label_paragraph(doc, "Caution: ", question["caution"], COLORS["security"])


def add_rehearsal_plan(doc: Document, content: dict) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Rehearsal Plan", level=1)
    for index, run in enumerate(content["rehearsal"]["runs"], start=1):
        doc.add_heading(f"Run {index}: {run['name']}", level=2)
        doc.add_paragraph(run["goal"])

    doc.add_heading("Four-Step Question Method", level=2)
    for step in content["rehearsal"]["answerMethod"]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Recovery Phrases", level=2)
    for phrase in content["rehearsal"]["recoveryPhrases"]:
        doc.add_paragraph(phrase, style="List Bullet")

    doc.add_heading("Three-Day Practice Schedule", level=2)
    schedule = [
        ("Day 1", "Read every note aloud. Mark words that feel unnatural and rehearse the opening and conclusion three times."),
        ("Day 2", "Run the full deck with a timer. Practice the five questions you find most difficult and use appendix slides."),
        ("Day 3", "Perform one final defense without stopping. Review only mistakes and sleep normally; do not rewrite the presentation."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("When", "Practice")):
        cell.text = value
        set_cell_shading(cell, COLORS["methodology"])
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for when, practice in schedule:
        row = table.add_row().cells
        row[0].text = when
        row[1].text = practice


def add_final_day_checklist(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Final-Day Checklist", level=1)
    groups = {
        "Files": [
            "Editable PowerPoint opens without repair warnings.",
            "Presentation PDF matches the final slide order.",
            "Preparation guide PDF is available offline.",
            "Evidence screenshots and the final thesis PDF are available as an offline backup.",
            "A second copy exists on a USB drive or another device.",
        ],
        "Room and display": [
            "Confirm the projector uses 16:9 widescreen mode.",
            "Test the HDMI or USB-C adapter before the committee arrives.",
            "Disable notifications, chat popups, automatic updates, and screen sleep.",
            "Use Presenter View and confirm the notes appear only on your screen.",
        ],
        "Speaking": [
            "Memorize the opening, conclusion, and first sentence of every results slide.",
            "Carry water and pause after each major result.",
            "Look at the professor who asked the question before looking at the appendix.",
            "If uncertain, state the measured boundary instead of guessing.",
        ],
    }
    for title, items in groups.items():
        doc.add_heading(title, level=2)
        for item in items:
            doc.add_paragraph(f"[ ] {item}")

    add_callout(
        doc,
        "Last reminder",
        "You do not need to know everything. You need to explain what was measured, what the evidence supports, and where the limitations begin.",
        COLORS["verified"],
    )


def build_guide(content: dict, output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc, content)
    add_how_to_use(doc)
    add_timing_map(doc, content)
    add_slide_scripts(doc, content)
    add_question_bank(doc, content)
    add_rehearsal_plan(doc, content)
    add_final_day_checklist(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    content = json.loads(CONTENT_PATH.read_text(encoding="utf-8"))
    build_guide(content, OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
