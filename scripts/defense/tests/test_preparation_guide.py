import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
GUIDE = (
    ROOT
    / "Document"
    / "outputs"
    / "final-defense"
    / "G6519692_Zaw_Ye_Htut_Ko_Defense_Preparation_Guide.docx"
)


class PreparationGuideTests(unittest.TestCase):
    def test_required_sections_and_question_count(self):
        doc = Document(GUIDE)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        for heading in (
            "Presentation Timing Map",
            "Slide-by-Slide Script",
            "Professor Question Bank",
            "Rehearsal Plan",
            "Final-Day Checklist",
        ):
            self.assertIn(heading, text)
        self.assertGreaterEqual(text.count("Short answer:"), 30)
        self.assertGreaterEqual(text.count("Follow-up answer:"), 30)

    def test_every_main_slide_has_script_and_transition(self):
        doc = Document(GUIDE)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        for slide_number in range(1, 23):
            self.assertIn(f"Slide {slide_number}:", text)
        self.assertGreaterEqual(text.count("Transition:"), 22)
        self.assertGreaterEqual(text.count("Key message:"), 22)

    def test_guide_has_substantial_rehearsal_support(self):
        doc = Document(GUIDE)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        for phrase in (
            "Answer the question in one sentence.",
            "Content run",
            "Timing run",
            "Defense run",
            "Presenter View",
            "offline backup",
        ):
            self.assertIn(phrase, text)

    def test_generated_xml_uses_schema_safe_order_and_attributes(self):
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(GUIDE) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
            settings = ET.fromstring(archive.read("word/settings.xml"))
        for shading in document.iter(f"{namespace}shd"):
            self.assertEqual("clear", shading.get(f"{namespace}val"))
        for properties in document.iter(f"{namespace}tcPr"):
            tags = [node.tag for node in properties]
            if f"{namespace}shd" in tags and f"{namespace}vAlign" in tags:
                self.assertLess(tags.index(f"{namespace}shd"), tags.index(f"{namespace}vAlign"))
            if f"{namespace}tcMar" in tags and f"{namespace}vAlign" in tags:
                self.assertLess(tags.index(f"{namespace}tcMar"), tags.index(f"{namespace}vAlign"))
        for properties in document.iter(f"{namespace}pPr"):
            tags = [node.tag for node in properties]
            if f"{namespace}pBdr" in tags:
                border_index = tags.index(f"{namespace}pBdr")
                for later_tag in (f"{namespace}spacing", f"{namespace}jc"):
                    if later_tag in tags:
                        self.assertLess(border_index, tags.index(later_tag))
        zoom = settings.find(f"{namespace}zoom")
        self.assertIsNotNone(zoom)
        self.assertEqual("100", zoom.get(f"{namespace}percent"))

    def test_question_headings_stay_with_their_answers(self):
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(GUIDE) as archive:
            document = ET.fromstring(archive.read("word/document.xml"))
        question_count = 0
        for paragraph in document.iter(f"{namespace}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{namespace}t"))
            if not text.startswith("Question "):
                continue
            question_count += 1
            properties = paragraph.find(f"{namespace}pPr")
            self.assertIsNotNone(properties)
            self.assertIsNotNone(properties.find(f"{namespace}keepNext"), text)
        self.assertGreaterEqual(question_count, 30)

    def test_guide_explains_core_and_future_work_scope(self):
        doc = Document(GUIDE)
        text_parts = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                text_parts.extend(cell.text for cell in row.cells)
        text = "\n".join(text_parts)
        self.assertIn("Primary evaluated model: ChatGPT/Codex 5.4", text)
        self.assertIn("Slides 20-21 are post-thesis future-work follow-ups", text)
        self.assertIn("Future Work Follow-Up: GPT-5.4 vs GPT-5.5", text)


if __name__ == "__main__":
    unittest.main()
