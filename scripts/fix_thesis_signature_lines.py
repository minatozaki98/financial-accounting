from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-23SEP.docx"
SIGNATURE_COLUMN_WIDTHS = (Twips(3960), Twips(1440), Twips(3960))


def signature_table(doc: DocumentObject):
    matches = [
        table for table in doc.tables
        if len(table.rows) == 7
        and len(table.columns) == 3
        and "Dr. Darun Kesrarat" in table.cell(1, 0).text
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one approval committee table; found {len(matches)}")
    return matches[0]


def set_signature_rule(paragraph) -> None:
    paragraph.clear()
    run = paragraph.add_run("\u00a0")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    for name, value in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{name}"), value)


def fix_signature_table(doc: DocumentObject) -> None:
    table = signature_table(doc)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, table_width)
    table_width.set(qn("w:w"), str(sum(width.twips for width in SIGNATURE_COLUMN_WIDTHS)))
    table_width.set(qn("w:type"), "dxa")

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, SIGNATURE_COLUMN_WIDTHS, strict=True):
        grid_column.set(qn("w:w"), str(width.twips))
    for row in table.rows:
        for cell, width in zip(row.cells, SIGNATURE_COLUMN_WIDTHS, strict=True):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(width.twips))
            tc_width.set(qn("w:type"), "dxa")

    for row_index in (0, 4):
        for column_index in (0, 2):
            set_signature_rule(table.cell(row_index, column_index).paragraphs[0])

    for row_index in (1, 2, 5, 6):
        for column_index in (0, 2):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            text = paragraph.text.strip()
            if row_index in (1, 5) and text.startswith("(") and text.endswith(")") and not text[1:-1].strip():
                text = "(" + "\u00a0" * 14 + ")"
            paragraph.clear()
            run = paragraph.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Inches(0)
            paragraph.paragraph_format.right_indent = Inches(0)
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0

    table.rows[3].height = Inches(0.4)
    table.rows[3].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def fix_report(path: Path) -> None:
    doc = Document(path)
    fix_signature_table(doc)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Replace wrapped approval signatures with stable rules.")
    parser.add_argument("target", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    target = args.target.resolve()
    fix_report(target)
    print(target)


if __name__ == "__main__":
    main()
