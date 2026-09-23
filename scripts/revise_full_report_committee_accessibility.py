from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-gpt55-comparison.docx"

ENDPOINT_FONT = "Courier New"
ENDPOINT_PATTERN = re.compile(
    r"(?:\b(?:GET|POST|PUT|DELETE|PATCH)\s+/(?:api/)?[A-Za-z0-9_{}?&=./:\-]+)"
    r"|(?:https?://[^\s,;)]+)"
    r"|(?:(?<![\w])/(?:api/)?(?:auth|accounts|journal-entries|journalentries|periods|reports|audit-logs|auditlogs|users)(?:/[A-Za-z0-9_{}?&=.\-]+)*)"
    r"|(?:(?:origin/)?(?:baseline-v0\.1|baseline-(?:sonarqube|zap|jmeter)-v1|codex/thesis-reproducibility-v1))"
    r"|(?:(?:[A-Za-z0-9_.-]+/)+[A-Za-z0-9_.{}-]+\.(?:cs|ps1|ts|tsx|json|sql|jmx|md|yml|yaml))",
    re.IGNORECASE,
)


def normalized(text: str) -> str:
    return " ".join(text.split())


def find_one(doc: DocumentObject, prefix: str):
    matches = [p for p in doc.paragraphs if normalized(p.text).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning with {prefix!r}; found {len(matches)}")
    return matches[0]


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def insert_paragraph_after(paragraph, text: str, *, bold_label: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    result = Paragraph(new_p, paragraph._parent)
    result.style = "Normal"
    if bold_label:
        run = result.add_run(bold_label)
        run.bold = True
    result.add_run(text)
    return result


def append_sentence(paragraph, sentence: str) -> None:
    if sentence in paragraph.text:
        return
    paragraph.add_run(" " + sentence)


def merge_section_3_3(doc: DocumentObject) -> None:
    intro = find_one(doc, "This research employs a combination of tools")
    set_paragraph_text(
        intro,
        "This section explains the three evaluation tools in plain language because each tool answers a different question. SonarQube asks whether the source code contains maintainability, reliability, or security concerns; OWASP ZAP asks how the running API behaves when examined like an external security tester; and Apache JMeter asks whether the API continues responding acceptably when many requests arrive together. ChatGPT (Codex 5.4) is the improvement advisor, while the three independent tools provide the measurements used to accept or reject its recommendations.",
    )
    intro.paragraph_format.keep_together = True

    sonar = find_one(doc, "What is SonarQube and why we focus on it?")
    sonar_duplicate = find_one(doc, "SonarQube (Community Edition):")
    set_paragraph_text(
        sonar,
        "SonarQube works like an automated code reviewer: it examines source code without requiring a committee member to read every file and reports possible bugs, security weaknesses, maintainability problems called code smells, duplicated code, and test coverage. A SonarQube issue is a warning that requires interpretation, not automatic proof that the application has failed. In this study, the .NET scanner examines the API, service, model, repository, and test projects, and the result is accepted only when the remediation branch reaches the defined Quality Gate without hiding changes in coverage or complexity. Because research has shown that static-analysis warnings do not always correspond to real defects (Lenarduzzi et al., 2020), the SonarQube result is interpreted together with ZAP and JMeter. The actual before-and-after implementation is reproduced in Appendix E, the dashboards are Figures L.1-L.4, and the result mapping is in Appendix M.",
    )
    remove_paragraph(sonar_duplicate)

    jmeter = find_one(doc, "What is Apache JMeter and why we focus on it?")
    jmeter_duplicate = find_one(doc, "Apache JMeter:")
    zap = find_one(doc, "What is OWASP ZAP and why we focus on it?")
    zap_duplicate = find_one(doc, "OWASP ZAP:")

    set_paragraph_text(
        jmeter,
        "OWASP ZAP works like an automated security tester positioned outside the application. Its passive scan observes pages and responses for missing headers or unsafe configuration, while its authenticated API scan signs in and exercises endpoints described by the OpenAPI document. Alerts are labelled High, Medium, Low, or Informational so that serious findings can be separated from observations. In this study, the configured gate passed, but one raw Medium HTTP Only Site observation remained because the local test used HTTP rather than production TLS; therefore the report does not claim that every raw Medium alert disappeared. The hardening code is reproduced in Appendix F, the scan dashboards are Figures L.5-L.8, and the result mapping is in Appendix M.",
    )
    remove_paragraph(jmeter_duplicate)

    set_paragraph_text(
        zap,
        "Apache JMeter works like a controlled crowd of users sending requests to the API at the same time. It records response time, throughput, and errors for 50 VU, 100 VU, 500 VU, soak, and spike profiles. The p95 value means that 95 out of every 100 measured requests completed at or below that time; a lower p95 is normally better. Soak testing checks sustained activity, while spike testing checks a sudden burst. The fresh remediation run improved the soak and spike aggregate p95 values but failed the core 100 VU and 500 VU thresholds, so the performance outcome is reported as Not Achieved for the core criterion. The performance implementation is reproduced in Appendix G, the dashboards are Figures L.9-L.18, and the result mapping is in Appendix M.",
    )
    remove_paragraph(zap_duplicate)

    branch_paragraph = insert_paragraph_after(
        zap,
        "baseline-v0.1 is the frozen pre-remediation baseline: it represents the code before the evaluated ChatGPT-guided changes. The branches baseline-sonarqube-v1, baseline-zap-v1, and baseline-jmeter-v1 are the first improvement versions for static quality, security, and performance respectively; each contains only the changes relevant to its evaluation so that one type of fix cannot distort another tool's comparison. The codex/thesis-reproducibility-v1 branch contains evidence automation and report support, not an additional improvement result. The suffixes v0.1 and v1 therefore identify controlled versions, while the exact commit IDs in Appendix A identify the immutable source actually tested.",
        bold_label="Branch roles in plain language. ",
    )
    branch_paragraph.paragraph_format.space_after = intro.paragraph_format.space_after


def revise_accessible_explanations(doc: DocumentObject) -> None:
    set_paragraph_text(
        find_one(doc, "Decision rule. In this report"),
        "Decision rule in plain language. Each improvement is judged against a rule chosen before interpreting the outcome. Achieved means the target was met without a material regression; Partially Achieved means some measurable benefit occurred but an important condition remained; and Not Achieved means the primary target failed. This prevents a visually positive dashboard or one improved number from being presented as success when another required measure became worse.",
    )
    set_paragraph_text(
        find_one(doc, "Code-quality comparison units:"),
        "How to read the code-quality evidence. SonarQube acts as an automated reviewer. Total issues count distinct warnings, while the newer Reliability, Security, and Maintainability impacts describe why an issue matters and may overlap. Coverage is the percentage of executable code exercised by tests, duplicated-line density indicates repeated code, and the Quality Gate is SonarQube's overall pass/fail decision. The report therefore considers both issue closure and changes in these control metrics rather than treating a lower issue count alone as proof of better software.",
    )
    set_paragraph_text(
        find_one(doc, "Security comparison units:"),
        "How to read the security evidence. ZAP examines the running API rather than only reading its source. High, Medium, and Low alerts indicate potential security risk, while Informational alerts describe observations that do not automatically fail the gate. The passive scan checks ordinary responses and browser-facing configuration; the authenticated API scan signs in and exercises protected endpoints. A remaining HTTP Only Site alert means the local reproduction used HTTP, not that a new application-code vulnerability was introduced.",
    )
    set_paragraph_text(
        find_one(doc, "Performance comparison units:"),
        "How to read the performance evidence. JMeter sends controlled workloads to the API and measures latency, throughput, and errors. The p95 latency is the time below which 95% of measured requests completed, so lower is normally better. Throughput is the number of completed transactions per second, and error rate is the proportion of failed requests. The core 50 VU, 100 VU, and 500 VU profiles determine the primary gate; soak and spike profiles provide additional evidence about sustained and sudden load but cannot override a failed core gate.",
    )


def cite_appendices(doc: DocumentObject) -> None:
    append_sentence(
        find_one(doc, "Therefore, the study evaluates the financial accounting API"),
        "The database schema and deterministic seed implementation are reproduced in Appendix H.",
    )
    append_sentence(
        find_one(doc, "Baseline Testing:"),
        "The exact baseline ref and commit are recorded in Appendix A.",
    )
    append_sentence(
        find_one(doc, "Post-Improvement Testing:"),
        "The automated tests are reproduced in Appendix I, and the exact-ref verification workflow is reproduced in Appendix K.",
    )
    set_paragraph_text(
        find_one(doc, "The evaluation compares baseline and remediated branches"),
        "The evaluation compares the frozen baseline with one tool-specific remediation branch at a time. Appendix A identifies the exact immutable commits, Appendix K shows how the scripts reject an incorrect or dirty checkout, Appendix M connects each conclusion to the corresponding machine-readable result, and Appendix N states the limits that must be considered when interpreting the outcome.",
    )

    sonar = find_one(doc, "Outcome judgment: Achieved with minor control-metric regressions.")
    set_paragraph_text(
        sonar,
        "Outcome judgment: Achieved with minor control-metric regressions. The fresh scan reduced 17 unique issues and 17 legacy Code Smells to 0, while the Quality Gate remained OK and duplicated-line density remained 0.0%. Coverage decreased from 74.4% to 74.3%, and cyclomatic complexity increased from 886 to 889, so these controls are reported as small regressions rather than described as unchanged. The before-and-after code is shown in Listings E.1-E.2 of Appendix E, the supporting dashboards are Figures L.1-L.4, the claim-to-evidence mapping is in Appendix M, and the interpretation limits are in Appendix N.",
    )
    security = find_one(doc, "Outcome judgment: Partially Achieved in raw-alert terms")
    set_paragraph_text(
        security,
        "Outcome judgment: Partially Achieved in raw-alert terms and Achieved for the configured gate. The passive scan improved from Medium 3 / Low 6 to Medium 0 / Low 0; the authenticated API scan improved from Medium 1 / Low 3 to Medium 1 / Low 0. Because one raw Medium HTTP Only Site observation remained, the report does not claim zero raw Medium alerts. The security-header and pipeline changes are shown in Listings F.1-F.2 of Appendix F, the supporting dashboards are Figures L.5-L.8, the claim-to-evidence mapping is in Appendix M, and the interpretation limits are in Appendix N.",
    )
    performance = find_one(doc, "Outcome judgment: Not Achieved for the core performance criterion")
    set_paragraph_text(
        performance,
        "Outcome judgment: Not Achieved for the core performance criterion and Partially Achieved for stress-profile behavior. The remediation core gate failed because 100 VU p95 was 1146.9 ms against a 500 ms threshold and 500 VU p95 was 1547.95 ms against a 1200 ms threshold. Soak total p95 improved from 443.95 ms to 171.0 ms, spike total p95 improved from 33401.75 ms to 4730.85 ms, and spike error rate fell from 0.4688% to 0.00%; these stress improvements do not compensate for the failed core gate. The implementation is shown in Listings G.1-G.3 of Appendix G, the execution controls are in Appendix K, the dashboards are Figures L.9-L.18, the traceability table is in Appendix M, and the limitations are in Appendix N.",
    )
    append_sentence(
        find_one(doc, "Overall, the fresh evidence is mixed."),
        "Appendix M summarizes the complete result chain, while Appendix N explains why these findings must not be generalized beyond this application and reproduction environment.",
    )
    append_sentence(
        find_one(doc, "The local seed process creates four demo accounts"),
        "The supporting code is reproduced in Appendices B-D, I, and J: Appendices B-D contain the server-side authentication, accounting, and audit implementations; Appendix I contains the automated role tests; and Appendix J contains the React/API integration.",
    )


def table_header(table) -> tuple[str, ...]:
    if not table.rows:
        return ()
    return tuple(normalized(cell.text) for cell in table.rows[0].cells)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def add_supporting_evidence_column(doc: DocumentObject) -> None:
    header_prefix = ("Evaluation area", "Measure", "Acceptance target", "Result")
    matches = [table for table in doc.tables if table_header(table)[:4] == header_prefix]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one evaluation-criteria table; found {len(matches)}")
    table = matches[0]
    if len(table.columns) == 4:
        table.add_column(Inches(1.35))
    if len(table.columns) != 5:
        raise RuntimeError(f"Evaluation-criteria table has {len(table.columns)} columns; expected 5")

    rows = [
        ["Evaluation area", "Measure", "Acceptance target", "Result", "Supporting evidence"],
        ["Static analysis", "SonarQube fresh issues and Quality Gate", "No retained blocker/critical issues after fixes; preserve Quality Gate OK", "17 unique issues reduced to 0; Quality Gate OK; coverage 74.4% -> 74.3%", "Appendix E; Figures L.1-L.4; Appendix M"],
        ["Security", "OWASP ZAP raw alerts and configured gate", "Clear actionable configured-rule alerts and disclose residual observations", "Configured gate passed; passive M3/L6 -> M0/L0; API M1/L3 -> M1/L0", "Appendix F; Figures L.5-L.8; Appendix M"],
        ["Performance", "JMeter p95, throughput, error rate, and gate", "50, 100, and 500 VU profiles pass with no material latency regression", "Remediation core gate FAIL; 100 VU and 500 VU exceeded thresholds", "Appendix G; Figures L.9-L.18; Appendices K and M"],
        ["Maintainability", "Representative before/after code evidence", "Selected code smells resolved with traceable branch evidence", "DTO, middleware-order, and report-persistence examples documented", "Listings E.1-E.2; Appendices E-G"],
    ]
    if len(table.rows) != len(rows):
        raise RuntimeError(f"Evaluation-criteria table has {len(table.rows)} rows; expected {len(rows)}")
    for row_index, (row, values) in enumerate(zip(table.rows, rows, strict=True)):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
            if row_index == 0:
                shade_cell(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(7.5)
                    run.bold = row_index == 0
    table.autofit = True

    note_text = "Supporting evidence is identified in the final column of Table 4.1. Appendix A records the exact branch and commit provenance, while Appendix N presents the cross-cutting limitations that apply to all four evaluation areas."
    old_matches = [p for p in doc.paragraphs if normalized(p.text).startswith("The evaluation compares the frozen baseline")]
    new_matches = [p for p in doc.paragraphs if normalized(p.text).startswith("Supporting evidence is identified in the final column")]
    if len(old_matches) == 1 and not new_matches:
        set_paragraph_text(old_matches[0], note_text)
    elif len(new_matches) == 1 and not old_matches:
        set_paragraph_text(new_matches[0], note_text)
    else:
        raise RuntimeError(
            f"Could not resolve Table 4.1 evidence note uniquely (old={len(old_matches)}, new={len(new_matches)})."
        )


def set_rfonts(run_element, font_name: str) -> None:
    r_pr = run_element.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        run_element.insert(0, r_pr)
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), font_name)


def split_endpoint_run(run) -> int:
    text = run.text
    matches = list(ENDPOINT_PATTERN.finditer(text))
    if not matches:
        return 0
    run_element = run._r
    parent = run_element.getparent()
    if parent is None:
        return 0
    allowed = {qn("w:rPr"), qn("w:t")}
    if any(child.tag not in allowed for child in run_element):
        set_rfonts(run_element, ENDPOINT_FONT)
        return 1

    segments: list[tuple[str, bool]] = []
    cursor = 0
    for match in matches:
        if match.start() > cursor:
            segments.append((text[cursor : match.start()], False))
        segments.append((match.group(0), True))
        cursor = match.end()
    if cursor < len(text):
        segments.append((text[cursor:], False))

    insert_at = parent.index(run_element)
    original_properties = run_element.find(qn("w:rPr"))
    for segment, is_endpoint in segments:
        if not segment:
            continue
        new_run = OxmlElement("w:r")
        if original_properties is not None:
            new_run.append(deepcopy(original_properties))
        if is_endpoint:
            set_rfonts(new_run, ENDPOINT_FONT)
        text_element = OxmlElement("w:t")
        if segment[0].isspace() or segment[-1].isspace():
            text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text_element.text = segment
        new_run.append(text_element)
        parent.insert(insert_at, new_run)
        insert_at += 1
    parent.remove(run_element)
    return len(matches)


def apply_identifier_typography(doc: DocumentObject) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Appendix Code":
            continue
        for run in list(paragraph.runs):
            count += split_endpoint_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in list(paragraph.runs):
                        count += split_endpoint_run(run)
    return count


def standardize_code_typography(doc: DocumentObject) -> None:
    if "Appendix Code" not in [style.name for style in doc.styles]:
        return
    style = doc.styles["Appendix Code"]
    style.font.name = "Courier New"
    style.font.size = Pt(9)
    for paragraph in doc.paragraphs:
        if paragraph.style.name != "Appendix Code":
            continue
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(9)


def revise_report(report_path: Path) -> int:
    doc = Document(report_path)
    marker = find_one(doc, "This section explains the three evaluation tools in plain language") if any(
        normalized(p.text).startswith("This section explains the three evaluation tools in plain language")
        for p in doc.paragraphs
    ) else None
    if marker is None:
        merge_section_3_3(doc)
        revise_accessible_explanations(doc)
        cite_appendices(doc)
    else:
        marker.paragraph_format.keep_together = True
    add_supporting_evidence_column(doc)
    endpoint_count = apply_identifier_typography(doc)
    standardize_code_typography(doc)
    doc.save(report_path)
    return endpoint_count


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Add appendix citations, accessible explanations, and identifier typography to the full report."
    )
    parser.add_argument("report", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    report_path = args.report.resolve()
    endpoint_count = revise_report(report_path)
    print(report_path)
    print(f"Endpoint/URL typography matches: {endpoint_count}")


if __name__ == "__main__":
    main()
