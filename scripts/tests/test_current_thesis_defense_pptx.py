from __future__ import annotations

import re
import unittest
import json
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
DECK = (
    ROOT
    / "Document"
    / "outputs"
    / "final-defense"
    / "G6519692_Zaw_Ye_Htut_Ko_Final_Thesis_Defense.pptx"
)


def slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


class CurrentThesisDefenseTests(unittest.TestCase):
    def test_committee_notes_are_plain_and_point_to_evidence(self):
        deck = Presentation(DECK)
        self.assertEqual(42, len(deck.slides))
        for index in range(22):
            notes = deck.slides[index].notes_slide.notes_text_frame.text
            self.assertIn("SHOW IF ASKED:", notes, f"slide {index + 1}")
            self.assertIn("APPENDIX:", notes, f"slide {index + 1}")

        dataset_rows = [
            " | ".join(cell.text for cell in row.cells)
            for shape in deck.slides[22].shapes if shape.has_table
            for row in shape.table.rows
        ]
        self.assertTrue(any("120 accounts" in row and "branch-verification.json" in row for row in dataset_rows))
        self.assertIn("not a measured peak from real customers", deck.slides[8].notes_slide.notes_text_frame.text)
        self.assertIn("Tool finding", slide_text(deck.slides[9]))
        self.assertIn("no recorded conditions", slide_text(deck.slides[12]).lower())
        self.assertIn("HTTPS retest", slide_text(deck.slides[13]))
        self.assertIn("set up HTTPS and repeat the scan", deck.slides[13].notes_slide.notes_text_frame.text)
        self.assertIn("cannot identify one cause", deck.slides[16].notes_slide.notes_text_frame.text)
        self.assertIn("SOAK p95 LOWER", slide_text(deck.slides[37]))
        self.assertIn("SPIKE p95 LOWER", slide_text(deck.slides[38]))
        self.assertIn("not load-test data", slide_text(deck.slides[39]).lower())

    def test_deck_carries_fresh_results_and_all_dashboard_captures(self):
        with ZipFile(DECK) as package:
            self.assertIsNone(package.testzip())
            media = [name for name in package.namelist() if name.startswith("ppt/media/")]
            self.assertGreaterEqual(len(media), 24)
            embedded_hashes = {sha256(package.read(name)).hexdigest() for name in media}
        manifest = json.loads(
            (ROOT / "docs" / "appendix" / "dashboard-capture-manifest.json").read_text(encoding="utf-8-sig")
        )
        for relative_path in manifest["requiredCaptures"]:
            source = ROOT / "docs" / "appendix" / "dashboards" / relative_path
            self.assertIn(sha256(source.read_bytes()).hexdigest(), embedded_hashes, relative_path)
        credential_screen = ROOT / "Document" / "outputs" / "web-app-evidence" / "01-login.png"
        self.assertNotIn(sha256(credential_screen.read_bytes()).hexdigest(), embedded_hashes)

        deck = Presentation(DECK)
        self.assertEqual(42, len(deck.slides))
        self.assertIn("17 to 0", slide_text(deck.slides[1]))
        self.assertIn("Target not met", slide_text(deck.slides[1]))
        self.assertIn("raw API Medium remains", slide_text(deck.slides[13]))
        self.assertIn("HTTP Only Site", slide_text(deck.slides[13]))
        self.assertIn("1,146.9", slide_text(deck.slides[16]))
        self.assertIn("1,547.95", slide_text(deck.slides[16]))
        self.assertIn("Core performance target was not met", slide_text(deck.slides[16]))
        self.assertIn("NOT MET:", slide_text(deck.slides[16]))
        self.assertIn("100 VU = planned peak", slide_text(deck.slides[8]))
        self.assertIn("500 VU = heavier load", slide_text(deck.slides[8]))
        self.assertIn("Exact cause of slowdown not isolated", slide_text(deck.slides[16]))
        self.assertIn("50 VU", slide_text(deck.slides[15]))
        self.assertIn("100 VU", slide_text(deck.slides[16]))
        self.assertIn("500 VU", slide_text(deck.slides[16]))
        self.assertNotIn("p100", slide_text(deck.slides[16]))
        self.assertNotIn("FAIL", "\n".join(slide_text(deck.slides[index]) for index in range(22)))
        self.assertIn("remediation core gate FAIL", slide_text(deck.slides[26]))
        self.assertIn("CPU/RAM and service tiers not controlled", slide_text(deck.slides[20]))
        self.assertIn("did not save CPU and memory readings", deck.slides[20].notes_slide.notes_text_frame.text)
        self.assertIn("historical dashboard capture", slide_text(deck.slides[32]))
        self.assertIn("historical dashboard capture", slide_text(deck.slides[33]))
        all_visible = "\n".join(slide_text(slide) for slide in deck.slides)
        self.assertNotIn("GPT-5.5", all_visible)
        self.assertNotIn("28 to 0", all_visible)
        self.assertNotIn("September", all_visible)
        self.assertNotIn("March", all_visible)

        times = []
        for index in range(22):
            slide = deck.slides[index]
            notes = slide.notes_slide.notes_text_frame.text
            self.assertNotIn("September", notes)
            self.assertNotIn("March", notes)
            self.assertIn("KEY MESSAGE:", notes)
            self.assertIn("SCRIPT:", notes)
            self.assertIn("SOURCE:", notes)
            match = re.search(r"TIMING: (\d+) seconds", notes)
            self.assertIsNotNone(match)
            times.append(int(match.group(1)))
        self.assertLessEqual(sum(times), 28 * 60)


if __name__ == "__main__":
    unittest.main()
