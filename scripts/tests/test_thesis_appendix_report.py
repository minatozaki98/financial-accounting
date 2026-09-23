from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document


ROOT = Path(__file__).resolve().parents[2]


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


appendix_builder = load_module(
    "appendix_builder", ROOT / "scripts" / "add_thesis_appendices_to_full_report.py"
)
report_formatter = load_module(
    "report_formatter", ROOT / "scripts" / "format_final_report_docx.py"
)
fresh_result_updater = load_module(
    "fresh_result_updater", ROOT / "scripts" / "update_full_report_fresh_results.py"
)
code_appendix_builder = load_module(
    "code_appendix_builder", ROOT / "scripts" / "rebuild_full_report_code_appendices.py"
)
committee_reviser = load_module(
    "committee_reviser", ROOT / "scripts" / "revise_full_report_committee_accessibility.py"
)
layout_normalizer = load_module(
    "layout_normalizer", ROOT / "scripts" / "normalize_full_report_layout.py"
)
signature_fixer = load_module(
    "signature_fixer", ROOT / "scripts" / "fix_thesis_signature_lines.py"
)
appendix_label_updater = load_module(
    "appendix_label_updater", ROOT / "scripts" / "standardize_appendix_labels.py"
)
figure_pairer = load_module(
    "figure_pairer", ROOT / "scripts" / "pair_chapter4_figures.py"
)
vu_label_updater = load_module(
    "vu_label_updater", ROOT / "scripts" / "clarify_jmeter_vu_labels.py"
)


class ThesisAppendixReportTests(unittest.TestCase):
    def test_fresh_report_discloses_device_and_service_tier_limits(self):
        doc = Document(vu_label_updater.DEFAULT_REPORT)
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        limitation = next(text for text in paragraphs if text.startswith("The findings are limited to one ASP.NET Core API"))
        appendix_limit = next(text for text in paragraphs if text.startswith("The fresh run records do not preserve"))
        resource_statement = next(text for text in paragraphs if text.startswith("Resource Utilization:"))
        self.assertIn("CPU model or core count", limitation)
        self.assertIn("installed RAM", limitation)
        self.assertIn("No managed-cloud service tier", limitation)
        self.assertIn("Docker resource limits", appendix_limit)
        self.assertIn("do not preserve per-run resource telemetry", resource_statement)
        self.assertNotIn("CPU and memory usage are tracked", resource_statement)
        environment = next(table for table in doc.tables if table.cell(0, 0).text == "Environment item")
        rows = {row.cells[0].text: [cell.text for cell in row.cells] for row in environment.rows[1:]}
        self.assertIn("CPU and RAM capacity not archived", rows["Host"][1])
        self.assertIn("no managed tier", rows["Database"][1])

    def test_jmeter_profile_labels_distinguish_virtual_users_from_percentiles(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(vu_label_updater.DEFAULT_REPORT.read_bytes())
            vu_label_updater.clarify_file(path)
            vu_label_updater.clarify_file(path)
            doc = Document(path)
            paragraphs = doc.paragraphs
            self.assertTrue(any(p.text == "p90 - 90th-percentile response latency" for p in paragraphs))
            self.assertTrue(any(p.text.startswith("VU - Virtual user") for p in paragraphs))
            self.assertTrue(any(p.text.startswith("Figure L.9. Apache JMeter baseline 50 VU") for p in paragraphs))
            self.assertTrue(any("report-baseline-20260922-192730-p50" in p.text for p in paragraphs))
            self.assertTrue(any("p50, p90, p95, and p99 denote response-time percentiles" in p.text for p in paragraphs))
            self.assertEqual(1, sum("Archived JMeter folder names retain" in p.text for p in paragraphs))
            self.assertEqual(1, sum("By contrast, p50, p90, p95, and p99" in p.text for p in paragraphs))
            profile_table = next(t for t in doc.tables if t.cell(0, 0).text == "Profile" and t.cell(0, 1).text == "Metric")
            self.assertEqual(["50 VU", "100 VU", "500 VU"], [profile_table.cell(i, 0).text for i in (1, 2, 3)])

    def test_chapter4_screenshots_are_paired_without_splitting_captions(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(figure_pairer.DEFAULT_REPORT.read_bytes())
            figure_pairer.pair_report_figures(path)
            figure_pairer.pair_report_figures(path)
            doc = Document(path)
            self.assertEqual(36, len(doc.inline_shapes))
            for number in range(2, 8):
                caption = figure_pairer.caption_for(doc, number)
                paragraphs = doc.paragraphs
                index = next(i for i, p in enumerate(paragraphs) if p._p is caption._p)
                picture = paragraphs[index - 1]
                shape = figure_pairer.InlineShape(picture._p.xpath(".//wp:inline")[0])
                preceding = paragraphs[index - 2]
                self.assertAlmostEqual(5.2, shape.width.inches, places=2)
                self.assertFalse(bool(caption.paragraph_format.keep_with_next))
                self.assertFalse(preceding._p.xpath('.//w:br[@w:type="page"]'))
                self.assertEqual(number in (2, 4, 6), bool(picture.paragraph_format.page_break_before))

    def test_appendix_field_labels_use_colons_without_losing_code_font(self):
        doc = Document()
        doc.add_paragraph("APPENDICES", style="Heading 1")
        purpose = doc.add_paragraph()
        purpose.add_run("Purpose. ").bold = True
        purpose.add_run("Explains the selected fix.")
        source = doc.add_paragraph()
        source.add_run("Source. ").bold = True
        path_run = source.add_run("API/Program.cs")
        path_run.font.name = "Courier New"
        verification = doc.add_paragraph()
        verification.add_run("Verification. ").bold = True
        verification.add_run("Related evidence: authenticated API scans.")

        changed = appendix_label_updater.standardize_appendix_labels(doc)
        repeated = appendix_label_updater.standardize_appendix_labels(doc)

        self.assertEqual({"Purpose:": 1, "Source:": 1, "Verification:": 1}, dict(changed))
        self.assertFalse(repeated)
        self.assertEqual("Purpose: Explains the selected fix.", purpose.text)
        self.assertEqual("Source: API/Program.cs", source.text)
        self.assertEqual("Courier New", path_run.font.name)
        self.assertEqual("LEFT (0)", str(source.alignment))
        self.assertEqual("Verification: Authenticated API scans.", verification.text)

    def test_signature_rules_do_not_use_wrapping_underscore_text(self):
        doc = Document()
        table = doc.add_table(rows=7, cols=3)
        for row_index in (0, 4):
            for column_index in (0, 2):
                table.cell(row_index, column_index).text = "_" * 29
        table.cell(1, 0).text = "            (Dr. Darun Kesrarat)"
        table.cell(1, 2).text = "     (                          )"
        table.cell(2, 0).text = "Advisor"
        table.cell(2, 2).text = "Committee"
        table.cell(5, 0).text = "     (                          )"
        table.cell(5, 2).text = "     (                          )"
        table.cell(6, 0).text = "Committee"
        table.cell(6, 2).text = "Committee"

        signature_fixer.fix_signature_table(doc)
        signature_fixer.fix_signature_table(doc)

        for row_index in (0, 4):
            for column_index in (0, 2):
                paragraph = table.cell(row_index, column_index).paragraphs[0]
                self.assertNotIn("_", paragraph.text)
                border = paragraph._p.xpath("./w:pPr/w:pBdr/w:bottom")
                self.assertEqual(1, len(border))
                self.assertEqual("single", border[0].get(signature_fixer.qn("w:val")))
        self.assertEqual("(Dr. Darun Kesrarat)", table.cell(1, 0).text)
        self.assertEqual("Advisor", table.cell(2, 0).text)

    def test_phase_methodology_becomes_idempotent_prose(self):
        doc = Document()
        doc.add_paragraph("3.1.1 Testing Environment and Reproducibility", style="Heading 3")
        old_paragraphs = [
            "Phase 1: Baseline Analysis",
            "Initial security and performance testing are conducted",
            "Tools: SonarQube",
            "Example: Financial Accounting API: A complex RESTful API",
            "The API features role-based access control",
            "Phase 2: ChatGPT-Generated Improvements",
            "Code recommendations are generated using ChatGPT",
            "Example Guidance from ChatGPT:",
            "Example: Security Enhancement",
            "Baseline Issue: Security and validation risk",
            "ChatGPT Recommendation: preserve Entity Framework",
            "Example: Performance Optimization",
            "Baseline Issue: Slow p95 latency",
            "ChatGPT Recommendation: Implement caching",
            "Example Adjustment: Use in-memory caching",
            "Phase 3: Post-Improvement Testing",
            "With ChatGPT-guided changes implemented",
            "Example: Security: Run ZAP",
            "Phase 4: Comparative Analysis and Reporting",
            "The data collected from the baseline",
        ]
        for text in old_paragraphs:
            doc.add_paragraph(text)
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "evidence preserved"
        doc.add_paragraph("3.1.2 Selected Source-Code Evidence", style="Heading 3")

        layout_normalizer.smooth_phase_narrative(doc)
        first_pass = [p.text for p in doc.paragraphs]
        layout_normalizer.smooth_phase_narrative(doc)

        self.assertEqual(first_pass, [p.text for p in doc.paragraphs])
        self.assertEqual("evidence preserved", table.cell(0, 0).text)
        self.assertEqual(4, len([p for p in doc.paragraphs if p.text.startswith("Phase ")]))
        self.assertFalse(any(p.text.startswith("Tools:") for p in doc.paragraphs))
        self.assertFalse(any(p._p.xpath("./w:pPr/w:numPr") for p in doc.paragraphs))
        prose = next(p for p in doc.paragraphs if p.text.startswith("Initial security, performance"))
        self.assertEqual("JUSTIFY (3)", str(prose.alignment))
        self.assertEqual(2.0, prose.paragraph_format.line_spacing)

    def test_example_labels_and_source_evidence_use_stable_alignment(self):
        doc = Document()
        example = doc.add_paragraph(" Example:")
        example.paragraph_format.left_indent = report_formatter.Inches(0.75)
        detail = doc.add_paragraph("Baseline Issue: A write endpoint needs validation.", style="List Paragraph")
        num_pr = report_formatter.OxmlElement("w:numPr")
        detail._p.get_or_add_pPr().append(num_pr)
        security_example = doc.add_paragraph("Example:")
        doc.add_paragraph("Security Enhancement:")
        performance_example = doc.add_paragraph("Performance Optimization:")
        evidence = doc.add_paragraph(
            "Source-code evidence A: SonarQube maintainability fix in API/Program.cs."
        )

        layout_normalizer.tidy_examples_and_evidence(doc)

        self.assertEqual("Example:", example.text)
        self.assertEqual("LEFT (0)", str(example.alignment))
        self.assertEqual(0.5, example.paragraph_format.left_indent.inches)
        self.assertTrue(bool(example.paragraph_format.keep_with_next))
        self.assertFalse(detail._p.xpath("./w:pPr/w:numPr"))
        self.assertEqual("JUSTIFY (3)", str(detail.alignment))
        self.assertEqual("LEFT (0)", str(evidence.alignment))
        self.assertEqual(0.0, evidence.paragraph_format.first_line_indent.inches)
        self.assertEqual("Example: Security Enhancement", security_example.text)
        self.assertEqual("Example: Performance Optimization", performance_example.text)
        self.assertNotIn("Security Enhancement:", [p.text for p in doc.paragraphs])

    def test_appendix_compaction_keeps_prose_justified_and_metadata_readable(self):
        doc = Document()
        doc.add_paragraph("APPENDICES", style="Heading 1")
        appendix_a = doc.add_paragraph("Appendix A - Environment", style="Normal")
        appendix_a.paragraph_format.page_break_before = True
        doc.add_paragraph("The study evaluates the API using branch-isolated evidence.")
        dataset = doc.add_paragraph(
            "Dataset: 120 accounts.\nFresh reproduction date: September 22, 2026.",
            style="List Paragraph",
        )
        dataset.paragraph_format.page_break_before = True
        appendix_b = doc.add_paragraph("Appendix B - Authentication", style="List Paragraph")
        appendix_b.paragraph_format.page_break_before = True
        source = doc.add_paragraph("Source. Controller and draft-creation authorization; file API/Program.cs")
        source.paragraph_format.page_break_before = True

        layout_normalizer.compact_appendices(doc)

        headings = [p for p in doc.paragraphs if p.text.startswith("Appendix ")]
        self.assertEqual(["Heading 1", "Heading 1"], [p.style.name for p in headings])
        self.assertTrue(all(not p.paragraph_format.page_break_before for p in doc.paragraphs[1:]))
        self.assertEqual("JUSTIFY (3)", str(doc.paragraphs[2].alignment))
        self.assertEqual("LEFT (0)", str(source.alignment))
        self.assertEqual(2.0, source.paragraph_format.line_spacing)
        self.assertEqual("Dataset: 120 accounts.", dataset.text)
        self.assertIn(
            "Fresh reproduction date: September 22, 2026.",
            [p.text for p in doc.paragraphs],
        )

    def test_layout_normalizer_standardizes_tables_bold_text_and_image_pagination(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(appendix_builder.DEFAULT_REPORT.read_bytes())

            layout_normalizer.normalize_report(path)

            doc = Document(path)
            data_tables = [
                table
                for table in doc.tables
                if not layout_normalizer.is_signature_table(table)
                and not layout_normalizer.is_code_table(table)
            ]
            self.assertTrue(data_tables)
            for table in data_tables:
                for row_index, row in enumerate(table.rows):
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            self.assertEqual(
                                "CENTER (1)" if row_index == 0 else "LEFT (0)",
                                str(paragraph.alignment),
                            )
                            for run in paragraph.runs:
                                if not run.text.strip():
                                    continue
                                self.assertEqual(10.0, run.font.size.pt)
                                self.assertEqual(row_index == 0, bool(run.bold))

            source_index = next(
                table
                for table in doc.tables
                if layout_normalizer.table_header(table)
                == ("Listing and purpose", "File", "Branch/ref", "Commit")
            )
            self.assertEqual(4, len(source_index.columns))
            self.assertRegex(source_index.cell(1, 0).text, r"^[A-Z]\.\d+ - ")

            code_tables = [table for table in doc.tables if layout_normalizer.is_code_table(table)]
            self.assertTrue(code_tables)
            for table in code_tables:
                for run in table.cell(0, 0).paragraphs[0].runs:
                    if run.text.strip():
                        self.assertEqual("Courier New", run.font.name)
                        self.assertEqual(9.5, run.font.size.pt)

            target = next(
                p
                for p in doc.paragraphs
                if p.text.startswith("The API features role-based access control")
            )
            self.assertFalse(any(run.bold for run in target.runs if run.text.strip()))

            paragraphs = doc.paragraphs
            for index, paragraph in enumerate(paragraphs[1:], start=1):
                if paragraph.paragraph_format.page_break_before:
                    self.assertFalse(layout_normalizer.is_empty_paragraph(paragraphs[index - 1]))

            caption_index = next(
                i
                for i, paragraph in enumerate(doc.paragraphs)
                if paragraph.text.startswith("Figure 3.2.")
                and paragraph.style.name == "Caption"
            )
            phase_image = next(
                paragraph
                for paragraph in reversed(doc.paragraphs[max(0, caption_index - 3) : caption_index])
                if layout_normalizer.has_drawing(paragraph)
            )
            self.assertTrue(layout_normalizer.has_drawing(phase_image))
            self.assertFalse(bool(phase_image.paragraph_format.page_break_before))

            section_3_1 = next(
                i for i, paragraph in enumerate(doc.paragraphs)
                if paragraph.text == "3.1 Research Design and Workflow"
            )
            section_3_1_1 = next(
                i for i, paragraph in enumerate(doc.paragraphs)
                if paragraph.text.startswith("3.1.1 Testing Environment")
                and paragraph.style.name == "Heading 3"
            )
            self.assertLess(section_3_1, section_3_1_1)
            self.assertEqual("Heading 2", doc.paragraphs[section_3_1].style.name)

            prose = next(
                paragraph for paragraph in doc.paragraphs
                if paragraph.text.startswith("The fresh ZAP evidence shows")
            )
            self.assertEqual(2.0, prose.paragraph_format.line_spacing)
            self.assertEqual("JUSTIFY (3)", str(prose.alignment))
            self.assertEqual(0.5, prose.paragraph_format.first_line_indent.inches)
            self.assertTrue(all(run.font.size.pt == 12 for run in prose.runs if run.text.strip()))

            numbered = next(
                paragraph for paragraph in doc.paragraphs
                if paragraph.text.startswith("Phase 3: Post-Improvement Testing")
            )
            self.assertIsNone(numbered.paragraph_format.first_line_indent)
            self.assertEqual("LEFT (0)", str(numbered.alignment))
            self.assertTrue(bool(numbered.paragraph_format.keep_with_next))

            report_formatter.rebuild_front_matter_lists(doc)
            report_formatter.enforce_document_typography(doc)
            for abstract in doc.part.numbering_part.element:
                if abstract.tag != report_formatter.qn("w:abstractNum"):
                    continue
                for level in abstract:
                    if level.tag != report_formatter.qn("w:lvl"):
                        continue
                    rpr = level.find(report_formatter.qn("w:rPr"))
                    fonts = rpr.find(report_formatter.qn("w:rFonts"))
                    self.assertEqual("Times New Roman", fonts.get(report_formatter.qn("w:ascii")))
            toc_entry = next(
                paragraph for paragraph in doc.paragraphs
                if paragraph.text.startswith("3.1 Research Design and Workflow\t")
            )
            self.assertEqual(2.0, toc_entry.paragraph_format.line_spacing)
            self.assertEqual(12.0, toc_entry.runs[0].font.size.pt)
            acronym = next(
                paragraph for paragraph in doc.paragraphs
                if paragraph.text.startswith("AI - Artificial Intelligence")
            )
            self.assertEqual(2.0, acronym.paragraph_format.line_spacing)
            self.assertEqual(12.0, acronym.runs[0].font.size.pt)

    def test_committee_revision_merges_tool_explanations_cites_appendices_and_styles_endpoints(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(appendix_builder.DEFAULT_REPORT.read_bytes())

            committee_reviser.revise_report(path)

            doc = Document(path)
            paragraphs = doc.paragraphs
            start = next(i for i, p in enumerate(paragraphs) if p.text == "3.3 Model and Tool Selection")
            end = next(i for i, p in enumerate(paragraphs[start + 1 :], start + 1) if p.text.startswith("3.4 Implementation"))
            section = "\n".join(p.text for p in paragraphs[start:end])
            self.assertNotIn("What is SonarQube", section)
            self.assertNotIn("SonarQube (Community Edition):", section)
            self.assertEqual(1, section.count("SonarQube works like an automated code reviewer"))
            self.assertEqual(1, section.count("OWASP ZAP works like an automated security tester"))
            self.assertEqual(1, section.count("Apache JMeter works like a controlled crowd"))
            intro = next(p for p in paragraphs[start:end] if p.text.startswith("This section explains"))
            self.assertTrue(bool(intro.paragraph_format.keep_together))
            self.assertIn("baseline-v0.1 is the frozen pre-remediation baseline", section)
            self.assertIn("baseline-sonarqube-v1", section)
            self.assertIn("Appendix A", section)

            appendix_index = next(i for i, p in enumerate(paragraphs) if p.text == "APPENDICES")
            body_text = "\n".join(p.text for p in paragraphs[:appendix_index])
            for reference in (
                "Appendix H",
                "Appendices B-D, I, and J",
                "Appendix E",
                "Appendix F",
                "Appendix G",
                "Appendix K",
                "Appendix L",
                "Appendix M",
                "Appendix N",
            ):
                self.assertIn(reference, body_text)

            endpoint_runs = [
                run
                for p in paragraphs[:appendix_index]
                for run in p.runs
                if "/auth" in run.text
            ]
            self.assertTrue(endpoint_runs)
            self.assertTrue(any(run.font.name == "Courier New" for run in endpoint_runs))
            self.assertEqual("Courier New", doc.styles["Appendix Code"].font.name)

            criteria = fresh_result_updater.find_table_by_header(
                doc,
                (
                    "Evaluation area",
                    "Measure",
                    "Acceptance target",
                    "Result",
                    "Supporting evidence",
                ),
            )
            criteria_rows = [[cell.text for cell in row.cells] for row in criteria.rows]
            self.assertEqual(
                [
                    "Static analysis",
                    "SonarQube fresh issues and Quality Gate",
                    "No retained blocker/critical issues after fixes; preserve Quality Gate OK",
                    "17 unique issues reduced to 0; Quality Gate OK; coverage 74.4% -> 74.3%",
                    "Appendix E; Figures L.1-L.4; Appendix M",
                ],
                criteria_rows[1],
            )
            self.assertNotIn("(Appendix A)", criteria_rows[1][3])

            branch_runs = [
                run
                for p in paragraphs[start:end]
                for run in p.runs
                if "baseline-v0.1" in run.text
            ]
            self.assertTrue(branch_runs)
            self.assertTrue(all(run.font.name == "Courier New" for run in branch_runs))
            path_runs = [
                run
                for p in paragraphs[appendix_index:]
                for run in p.runs
                if "API/Program.cs" in run.text
            ]
            self.assertTrue(path_runs)
            self.assertTrue(all(run.font.name == "Courier New" for run in path_runs))

    def test_direct_code_appendices_are_self_contained_and_cover_a_through_n(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(appendix_builder.DEFAULT_REPORT.read_bytes())

            count = code_appendix_builder.rebuild_code_appendices(path, ROOT)

            doc = Document(path)
            headings = [
                p.text
                for p in doc.paragraphs
                if p.style.name == "Heading 1" and p.text.startswith("Appendix ")
            ]
            self.assertEqual(14, len(headings))
            self.assertTrue(headings[0].startswith("Appendix A -"))
            self.assertTrue(headings[-1].startswith("Appendix N -"))
            self.assertGreaterEqual(count, 20)
            self.assertTrue(
                all(not p.paragraph_format.page_break_before for p in doc.paragraphs
                    if p.text.startswith("Appendix ") and p.style.name == "Heading 1")
            )
            self.assertEqual(
                18,
                len(
                    [
                        p
                        for p in doc.paragraphs
                        if p.style.name == "Caption" and p.text.startswith("Figure L.")
                    ]
                ),
            )
            full_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertNotIn("This Markdown file", full_text)
            self.assertNotIn("Admin@123", full_text)
            self.assertIn("B.1 JWT Authentication Configuration", full_text)
            self.assertIn("K.2 Environment and Tool Readiness Checks", full_text)
            self.assertIn("Appendix N - Limitations and Complete Source Index", full_text)
            self.assertIn("Purpose: Configures JWT bearer validation", full_text)
            self.assertIn("Source: Remediated authentication registration", full_text)
            self.assertIn("Verification: Authenticated API scans", full_text)
            self.assertNotIn("Purpose. ", full_text)
            first_sonar_listing = next(
                p for p in doc.paragraphs if p.text.startswith("E.1 S107")
            )
            self.assertFalse(bool(first_sonar_listing.paragraph_format.page_break_before))
            self.assertNotIn("`API/Program.cs`", full_text)

    def test_fresh_result_inputs_match_versioned_machine_readable_evidence(self):
        evidence = fresh_result_updater.load_fresh_evidence(ROOT)

        self.assertEqual(17, evidence["sonar"]["baseline_issues"])
        self.assertEqual(0, evidence["sonar"]["remediation_issues"])
        self.assertEqual("74.4", evidence["sonar"]["baseline_coverage"])
        self.assertEqual("74.3", evidence["sonar"]["remediation_coverage"])
        self.assertEqual(1, evidence["zap"]["remediation_api_medium"])
        self.assertEqual("FAIL", evidence["jmeter"]["remediation_gate"])
        self.assertEqual(1146.9, evidence["jmeter"]["remediation_profiles"]["p100"]["p95Ms"])

    def test_report_update_replaces_primary_historical_claims_but_preserves_appendix_history(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            path.write_bytes(appendix_builder.DEFAULT_REPORT.read_bytes())

            fresh_result_updater.update_report(path, ROOT)

            doc = Document(path)
            paragraphs = [p.text for p in doc.paragraphs]
            appendix_index = paragraphs.index("APPENDICES")
            primary_text = "\n".join(paragraphs[:appendix_index])
            appendix_text = "\n".join(paragraphs[appendix_index:])
            table_text = "\n".join(
                cell.text
                for table in doc.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("17 to 0", primary_text)
            self.assertIn("74.4% to 74.3%", primary_text)
            self.assertIn("remediation core gate failed", primary_text)
            self.assertNotIn("retained C# findings from 28 to 0", primary_text)
            self.assertIn("Historical measurements are retained", appendix_text)
            self.assertIn("28 retained findings to 0", table_text)

            table = fresh_result_updater.find_table_by_header(
                doc, ("Metric", "Baseline", "After fixes", "Interpretation")
            )
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            self.assertIn(["Total unique issues", "17", "0", "All fresh baseline issues closed"], rows)
            self.assertIn(["Coverage", "74.4%", "74.3%", "Decreased by 0.1 percentage points"], rows)

    def test_dashboard_page_breaks_keep_tool_heading_with_first_dashboard(self):
        self.assertFalse(
            appendix_builder.dashboard_heading_page_break(3, "SonarQube Dashboards", False)
        )
        self.assertTrue(
            appendix_builder.dashboard_heading_page_break(3, "OWASP ZAP Dashboards", False)
        )
        self.assertFalse(
            appendix_builder.dashboard_heading_page_break(4, "baseline-overview", True)
        )
        self.assertTrue(
            appendix_builder.dashboard_heading_page_break(4, "baseline-issues", False)
        )

    def test_builder_creates_template_independent_bullet_style(self):
        doc = Document(appendix_builder.DEFAULT_REPORT)
        style_name = appendix_builder.ensure_bullet_style(doc)
        paragraph = appendix_builder.add_body_paragraph(doc, "Repository entry", style=style_name)

        self.assertEqual("Appendix Bullet", paragraph.style.name)
        self.assertEqual(
            0.25,
            round(doc.styles[style_name].paragraph_format.left_indent.inches, 2),
        )

    def test_rebuild_references_preserves_appendices_and_places_references_before_them(self):
        doc = Document()
        doc.add_paragraph("REFERENCES", style="Heading 1")
        doc.add_paragraph("stale reference")
        doc.add_paragraph("APPENDICES", style="Heading 1")
        doc.add_paragraph("Appendix A - Repository and Environment", style="Heading 1")
        doc.add_paragraph("appendix body")

        report_formatter.rebuild_references(doc)

        texts = [report_formatter.normalized_text(p) for p in doc.paragraphs]
        self.assertNotIn("stale reference", texts)
        self.assertIn("appendix body", texts)
        reference_index = texts.index("REFERENCES")
        appendix_index = texts.index("APPENDICES")
        self.assertGreater(appendix_index, reference_index + 1)
        self.assertTrue(any("GPTutor" in text for text in texts[reference_index:appendix_index]))

    def test_lists_only_refresh_does_not_require_legacy_section_3_1_heading(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            doc = Document()
            doc.add_paragraph("Table of Contents")
            doc.add_paragraph("List of Acronyms", style="Heading 1")
            doc.add_paragraph("API - Application Programming Interface")
            doc.add_paragraph("Chapter 1 - Introduction", style="Heading 1")
            doc.add_paragraph("APPENDICES", style="Heading 1")
            doc.add_paragraph("Appendix A - Repository and Environment", style="Heading 1")
            doc.add_paragraph("appendix body")
            doc.save(path)

            with (
                patch.object(report_formatter, "export_pdf", return_value=path.with_suffix(".pdf")),
                patch.object(report_formatter, "rendered_page_numbers", return_value={}),
            ):
                report_formatter.refresh_generated_lists(path)

            texts = [p.text for p in Document(path).paragraphs]
            self.assertIn("APPENDICES", texts)
            self.assertIn("Appendix A - Repository and Environment", texts)
            self.assertIn("appendix body", texts)

    def test_word_page_selection_prefers_body_over_generated_list_entry(self):
        self.assertEqual(
            95,
            report_formatter.select_word_page(
                "Appendix A - Repository and Environment", [5, 95]
            ),
        )
        self.assertEqual(3, report_formatter.select_word_page("Abstract", [3, 12]))

    def test_cross_references_are_idempotent(self):
        doc = Document()
        for heading, sentence in appendix_builder.CROSS_REFERENCES.items():
            doc.add_paragraph(heading, style="Heading 2")
            doc.add_paragraph("Section body.")

        appendix_builder.add_cross_references(doc)
        appendix_builder.add_cross_references(doc)

        text = "\n".join(p.text for p in doc.paragraphs)
        for sentence in appendix_builder.CROSS_REFERENCES.values():
            self.assertEqual(1, text.count(sentence))

    def test_cross_reference_targets_normal_prose_not_a_trailing_caption(self):
        doc = Document()
        heading = "3.1.1 Testing Environment and Reproducibility"
        sentence = appendix_builder.CROSS_REFERENCES[heading]
        prose = None
        caption = None
        for current_heading in appendix_builder.CROSS_REFERENCES:
            doc.add_paragraph(current_heading, style="Heading 3")
            body = doc.add_paragraph("Methodology body.", style="Normal")
            if current_heading == heading:
                prose = body
                caption = doc.add_paragraph("Table 3.7. Evidence table.", style="Caption")

        appendix_builder.add_cross_references(doc)

        self.assertIsNotNone(prose)
        self.assertIsNotNone(caption)
        self.assertIn(sentence, prose.text)
        self.assertNotIn(sentence, caption.text)

    def test_existing_appendix_tail_is_replaced_without_removing_section_properties(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "report.docx"
            doc = Document()
            doc.add_paragraph("REFERENCES", style="Heading 1")
            doc.add_paragraph("reference body")
            doc.add_paragraph("APPENDICES", style="Heading 1")
            doc.add_paragraph("old appendix body")
            appendix_builder.remove_existing_appendices(doc)
            doc.save(path)

            reopened = Document(path)
            texts = [p.text for p in reopened.paragraphs]
            self.assertIn("reference body", texts)
            self.assertNotIn("APPENDICES", texts)
            self.assertNotIn("old appendix body", texts)
            self.assertIsNotNone(reopened._element.body.sectPr)

    def test_controlled_source_contains_explanations_without_embedded_demo_password(self):
        source = (ROOT / "docs" / "appendix" / "thesis-appendix-source-code-and-results.md").read_text(
            encoding="utf-8"
        )
        self.assertEqual(5, source.count("**Explanation.**"))
        self.assertIn("### C.1 Role-Based Authorization", source)
        self.assertIn("### C.5 Deterministic Research-Data Parameters", source)
        self.assertNotIn("Admin@123", source)


if __name__ == "__main__":
    unittest.main()
