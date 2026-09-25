from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "Document" / "outputs" / "final-report-23SEPv2-1.docx"
OUTPUT = (
    ROOT
    / "Document"
    / "outputs"
    / "final-defense"
    / "G6519692_Zaw_Ye_Htut_Ko_Final_Thesis_Defense.pptx"
)
DASHBOARD_ROOT = ROOT / "docs" / "appendix" / "dashboards"
WEB_ROOT = ROOT / "Document" / "outputs" / "web-app-evidence"
MANIFEST = ROOT / "docs" / "appendix" / "dashboard-capture-manifest.json"
EVIDENCE_ROOT = ROOT / "docs" / "appendix" / "verification-runs" / "research"

W = 13.333
H = 7.5
FONT = "Arial"
CODE_FONT = "Courier New"
PROFILE_DISPLAY = {"p50": "50 VU", "p100": "100 VU", "p500": "500 VU"}
INK = RGBColor(37, 43, 51)
MUTED = RGBColor(91, 104, 116)
LIGHT = RGBColor(247, 249, 252)
BORDER = RGBColor(216, 224, 233)
TEAL = RGBColor(15, 103, 109)
TEAL_LIGHT = RGBColor(231, 246, 246)
ORANGE = RGBColor(190, 116, 20)
ORANGE_LIGHT = RGBColor(255, 246, 231)
GREEN = RGBColor(36, 111, 70)
GREEN_LIGHT = RGBColor(232, 247, 238)
RED = RGBColor(179, 55, 61)
RED_LIGHT = RGBColor(253, 239, 240)
NAVY = RGBColor(36, 70, 108)
NAVY_LIGHT = RGBColor(235, 243, 251)
WHITE = RGBColor(255, 255, 255)


def read_facts() -> dict:
    report = Document(REPORT)
    title = " ".join(report.paragraphs[1].text.split()).replace("\xa0", "")
    tables = report.tables
    sonar_rows = {row.cells[0].text.strip(): (row.cells[1].text.strip(), row.cells[2].text.strip()) for row in tables[21].rows[1:]}
    closure = {row.cells[0].text.strip(): row.cells[3].text.strip() for row in tables[24].rows[1:]}
    if sonar_rows["Total unique issues"] != ("17", "0") or "FAIL" not in closure["JMeter core profiles"]:
        raise RuntimeError("The current report does not match the expected fresh-result baseline.")

    sonar = {}
    zap = {}
    jmeter = {}
    for role in ("baseline", "remediation"):
        sonar[role] = json.loads((EVIDENCE_ROOT / "sonar" / role / "summary.json").read_text(encoding="utf-8-sig"))
        zap[role] = json.loads((EVIDENCE_ROOT / "zap" / role / "summary.json").read_text(encoding="utf-8-sig"))
        jmeter[role] = json.loads((EVIDENCE_ROOT / "jmeter" / role / "summary.json").read_text(encoding="utf-8-sig"))
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8-sig"))
    captures = {capture["imagePath"]: capture for capture in manifest["captures"]}
    return {
        "title": title,
        "sonar": sonar,
        "zap": zap,
        "jmeter": jmeter,
        "sonar_rows": sonar_rows,
        "closure": closure,
        "captures": captures,
    }


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, margin=0.02, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = str(value)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def line(slide, x1, y1, x2, y2, color=BORDER, width=1):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def header(prs, section, title, source, accent=TEAL):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    box(slide, 0, 0, 0.16, H, accent)
    text(slide, section.upper(), 0.64, 0.22, 11.9, 0.24, 11, accent, bold=True)
    text(slide, title, 0.64, 0.55, 12.0, 0.58, 29, INK, bold=True)
    line(slide, 0.64, 1.18, 12.67, 1.18)
    line(slide, 0.64, 7.03, 12.67, 7.03)
    text(slide, source, 0.64, 7.08, 11.4, 0.20, 10, MUTED)
    text(slide, len(prs.slides), 12.16, 7.08, 0.48, 0.20, 11, accent, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def add_notes(slide, seconds, key, script, source, caution=""):
    slide.notes_slide.notes_text_frame.text = (
        f"TIMING: {seconds} seconds\nKEY MESSAGE: {key}\n\nSCRIPT:\n{script}\n\n"
        f"SOURCE: {source}"
        + (f"\nCAUTION: {caution}" if caution else "")
    )


def image_contain(slide, path: Path, x, y, w, h, border=True, visible_top_px: int | None = None):
    with Image.open(path) as image:
        px_w, px_h = image.size
    shown_h = min(px_h, visible_top_px) if visible_top_px else px_h
    scale = min(w / px_w, h / shown_h)
    placed_w = px_w * scale
    placed_h = shown_h * scale
    placed_x = x + (w - placed_w) / 2
    placed_y = y + (h - placed_h) / 2
    picture = slide.shapes.add_picture(str(path), Inches(placed_x), Inches(placed_y), width=Inches(placed_w), height=Inches(placed_h))
    if shown_h < px_h:
        picture.crop_bottom = (px_h - shown_h) / px_h
    if border:
        frame = box(slide, placed_x, placed_y, placed_w, placed_h, fill=WHITE, line=BORDER)
        frame.fill.background()
    return (placed_x, placed_y, placed_w, placed_h)


def metric(slide, x, y, w, label, value, sub="", fill=WHITE, color=INK):
    box(slide, x, y, w, 1.22, fill, BORDER, radius=True)
    text(slide, label.upper(), x+0.18, y+0.12, w-0.36, 0.22, 11, MUTED, bold=True)
    text(slide, value, x+0.18, y+0.39, w-0.36, 0.45, 25, color, bold=True)
    if sub:
        text(slide, sub, x+0.18, y+0.93, w-0.36, 0.18, 11, MUTED)


def table(slide, rows, x, y, w, h, widths=None, font_size=15, header_fill=INK):
    grid = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for column, fraction in zip(grid.columns, widths, strict=True):
            column.width = Inches(w * fraction)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = grid.cell(row_index, column_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if row_index == 0 else (WHITE if row_index % 2 else LIGHT)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = WHITE if row_index == 0 else INK
    return grid


def list_items(slide, items, x, y, w, row_h=0.68, size=19, accent=TEAL):
    for index, item in enumerate(items):
        top = y + index * row_h
        box(slide, x, top+0.12, 0.12, 0.12, accent, radius=True)
        text(slide, item, x+0.28, top, w-0.28, row_h-0.02, size, INK)


def source_line(slide, label, value, x, y, w, color=INK):
    text(slide, label, x, y, w, 0.28, 12, MUTED, bold=True)
    text(slide, value, x, y+0.27, w, 0.42, 18, color, bold=True)


def profile_map(summary):
    return {profile["profile"]: profile for profile in summary["profiles"]}


def percent_lower(before, after):
    return round((before - after) * 100.0 / before)


def build_main_slides(prs: Presentation, facts: dict) -> int:
    sonar = facts["sonar"]
    zap = facts["zap"]
    jmeter = facts["jmeter"]
    baseline_profiles = profile_map(jmeter["baseline"])
    after_profiles = profile_map(jmeter["remediation"])

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    box(slide, 0, 0, 0.20, H, TEAL)
    text(slide, "FINAL THESIS DEFENSE", 0.72, 0.55, 11.8, 0.32, 15, TEAL, bold=True)
    text(slide, facts["title"], 0.72, 1.22, 11.78, 2.10, 31, INK, bold=True, valign=MSO_ANCHOR.TOP)
    line(slide, 0.72, 3.55, 12.52, 3.55, TEAL, 2)
    text(slide, "Zaw Ye Htut Ko  |  G6519692", 0.72, 3.93, 9.5, 0.45, 23, INK, bold=True)
    text(slide, "Master of Science in Information Technology  |  Assumption University", 0.72, 4.52, 11.3, 0.35, 18, MUTED)
    text(slide, "Advisor: Asst. Prof. Dr. Darun Kesrarat", 0.72, 5.05, 10.9, 0.34, 18, MUTED)
    box(slide, 0.72, 6.20, 11.78, 0.66, TEAL_LIGHT, radius=True)
    text(slide, "One existing API · one shared baseline · three independent measurements", 0.95, 6.30, 11.30, 0.43, 19, TEAL, bold=True)
    add_notes(slide, 45, "This is an evidence-based evaluation of Codex 5.4 as an advisor.",
              "Good morning. I evaluate whether ChatGPT, used through Codex 5.4, can help improve an existing financial accounting API. The model proposes changes; people review them; SonarQube, OWASP ZAP, and JMeter decide what the measured results actually support.",
              "Current full report title page and Abstract.")

    # 2. Decision snapshot
    slide = header(prs, "Opening", "The fresh evidence is mixed", "Current report §4.5; Table 4.5; Appendix M", TEAL)
    metric(slide, 0.78, 1.56, 3.78, "Static quality", "17 to 0", "Unique SonarQube issues", GREEN_LIGHT, GREEN)
    metric(slide, 4.78, 1.56, 3.78, "Security", "Gate PASS", "One raw API Medium remains", ORANGE_LIGHT, ORANGE)
    metric(slide, 8.78, 1.56, 3.78, "Performance", "Target not met", "100 and 500 VU exceed limits", RED_LIGHT, RED)
    box(slide, 0.78, 3.28, 11.78, 2.64, WHITE, BORDER, radius=True)
    text(slide, "Defensible conclusion", 1.05, 3.62, 10.9, 0.36, 20, TEAL, bold=True)
    text(slide, "ChatGPT-guided remediation closed the fresh static issues and passed the configured ZAP gate. It did not meet the predefined core performance target, despite better soak and spike aggregates.",
         1.05, 4.10, 10.90, 1.30, 23, INK)
    add_notes(slide, 60, "Static quality succeeded; security has a disclosed residual; the core performance target was not met.",
              "I want to give the result before the details. The fresh reproduction reduced SonarQube issues from seventeen to zero. ZAP passed the configured rule gate, but one raw Medium HTTP-only observation remained in the authenticated scan. JMeter's primary core performance target was not met at the 100 and 500 VU workloads. Improvements in soak and spike do not change that conclusion.",
              "Current report Table 4.5 and Appendix M.",
              "Do not repeat the historical claim that all core JMeter profiles improved.")

    # 3. Practical problem
    slide = header(prs, "Problem", "One financial API has three different risks", "Current report §§1.1-1.4 and §3.1", NAVY)
    for x, title, body, color, pale in [
        (0.78, "Maintainability", "Code smells, complex controller signatures, and repeated logic", TEAL, TEAL_LIGHT),
        (4.78, "Security", "Headers, authorization boundaries, and scanner observations", RED, RED_LIGHT),
        (8.78, "Performance", "Tail latency, throughput, errors, and sustained load", ORANGE, ORANGE_LIGHT),
    ]:
        box(slide, x, 1.72, 3.78, 3.65, pale, BORDER, radius=True)
        box(slide, x, 1.72, 0.10, 3.65, color)
        text(slide, title, x+0.25, 2.10, 3.25, 0.50, 24, color, bold=True)
        text(slide, body, x+0.25, 2.94, 3.20, 1.72, 21, INK)
    text(slide, "No one scanner can stand in for all three dimensions.", 0.78, 5.84, 11.8, 0.55, 24, INK, bold=True)
    add_notes(slide, 60, "The three risks need different measurements.",
              "A financial API can have source-code warnings, runtime security observations, and slow responses at the same time. These failures have different causes. That is why the study uses three independent tools and checks that remediation does not break the accounting behavior.",
              "Current report Chapter 1 and §3.1.")

    # 4. Research questions
    slide = header(prs, "Problem", "Three questions, three independent answers", "Current report §1.5 and Table 4.1", NAVY)
    questions = [
        ("RQ1", "Can static issues be closed without material control-metric regression?", TEAL),
        ("RQ2", "Can configured security alerts be cleared and residual raw alerts disclosed?", RED),
        ("RQ3", "Can latency targets pass under core load as well as stress profiles?", ORANGE),
    ]
    for i, (tag, question, color) in enumerate(questions):
        y = 1.54 + i*1.62
        box(slide, 0.80, y, 11.73, 1.33, WHITE, BORDER, radius=True)
        box(slide, 0.80, y, 1.30, 1.33, color)
        text(slide, tag, 0.96, y+0.40, 0.96, 0.44, 24, WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(slide, question, 2.38, y+0.24, 9.80, 0.84, 22, INK)
    add_notes(slide, 65, "Each research question has an acceptance rule.",
              "The first question is about code quality: can the issue count fall without overlooking coverage and complexity? The second is about security: can configured scanner rules pass while raw observations remain visible? The third is about performance: do the primary load gates pass, not just selected stress profiles? These rules were used to judge the results in Chapter Four.",
              "Current report §1.5 and Table 4.1.")

    # 5. System under study
    slide = header(prs, "Method", "The unit of analysis is an existing accounting API", "Current report §§3.1-3.2; Table 3.3", TEAL)
    box(slide, 0.80, 1.55, 4.10, 4.83, NAVY_LIGHT, BORDER, radius=True)
    text(slide, "ASP.NET Core 8", 1.09, 1.88, 3.52, 0.50, 27, NAVY, bold=True)
    text(slide, "SQL Server financial database\nJWT authentication · RBAC\nDouble-entry posting · audit log", 1.09, 2.61, 3.52, 2.64, 22, INK)
    text(slide, "Four roles: Admin, FinanceManager, User, Auditor", 1.09, 5.61, 3.48, 0.48, 15, MUTED)
    box(slide, 5.12, 1.55, 7.40, 4.83, WHITE, BORDER, radius=True)
    text(slide, "Endpoint groups exercised", 5.41, 1.86, 6.83, 0.42, 22, TEAL, bold=True)
    list_items(slide, ["Authentication and users", "Accounts and periods", "Journal entries and bulk posting", "Financial reports", "Audit logs"],
               5.42, 2.48, 6.65, row_h=0.70, size=20)
    add_notes(slide, 70, "This is one nontrivial API, not a set of unrelated projects.",
              "The evaluated system has accounting workflows, a SQL Server database, JWT authentication, four exact role names, protected write operations, report generation, and audit logging. The same API is measured before and after each tool-specific branch. The dashboard later in this talk is a demonstration client, not the benchmark dataset itself.",
              "Current report §§3.1-3.2, Table 3.3, Figures 4.2-4.7.")

    # 6. Branch isolation
    slide = header(prs, "Method", "Every remediation branch starts from one baseline", "Current report Table 3.1 and Appendix A", TEAL)
    metric(slide, 4.65, 1.45, 4.03, "Common starting point", "baseline-v0.1", "commit 7a0469d9", NAVY_LIGHT, NAVY)
    paths = [
        ("SonarQube", "baseline-sonarqube-v1", "a2279bcb", TEAL, 0.80),
        ("OWASP ZAP", "baseline-zap-v1", "5f3bd3cc", RED, 4.78),
        ("JMeter", "baseline-jmeter-v1", "78b08b62", ORANGE, 8.76),
    ]
    for label, branch, commit, color, x in paths:
        line(slide, 6.67, 2.68, x+1.90, 3.30, color, 2)
        box(slide, x, 3.34, 3.80, 2.10, WHITE, BORDER, radius=True)
        box(slide, x, 3.34, 0.10, 2.10, color)
        text(slide, label, x+0.22, 3.63, 3.32, 0.42, 22, color, bold=True)
        text(slide, branch, x+0.22, 4.18, 3.32, 0.38, 16, INK, bold=True)
        text(slide, f"Tested commit {commit}", x+0.22, 4.74, 3.32, 0.28, 13, MUTED)
    text(slide, "Compare one tool branch at a time; do not combine their deltas into one causal claim.", 0.80, 5.96, 11.8, 0.50, 19, INK)
    add_notes(slide, 70, "Branch isolation makes each before-and-after comparison traceable.",
              "We froze a common baseline at commit 7a0469d9. The SonarQube, ZAP, and JMeter changes were assessed on separate refs. The exact tested commits are recorded, so a later change to a branch name cannot silently change the result. I compare each remediation branch to that same baseline, one tool at a time.",
              "Current report Appendix A and Table 3.1.")

    # 7. Workflow
    slide = header(prs, "Method", "Four phases turn suggestions into evidence", "Current report §3.1.1; Figures 3.1-3.4", TEAL)
    phases = [
        ("01", "Capture baseline", "Freeze code, dataset, and raw tool output"),
        ("02", "Propose bounded fix", "Give Codex the finding, source context, and constraints"),
        ("03", "Review and re-test", "Human review, focused tests, same tool rerun"),
        ("04", "Compare and report", "Apply gates; disclose residuals and unmet targets"),
    ]
    for i, (number, title, body) in enumerate(phases):
        x = 0.76 + i*3.13
        box(slide, x, 1.83, 2.91, 3.65, WHITE, BORDER, radius=True)
        text(slide, number, x+0.20, 2.08, 0.80, 0.57, 30, TEAL, bold=True)
        text(slide, title, x+0.20, 2.91, 2.48, 0.70, 20, INK, bold=True)
        text(slide, body, x+0.20, 3.88, 2.48, 1.16, 17, MUTED)
    text(slide, "The model does not approve its own output.", 0.78, 5.98, 11.8, 0.50, 23, TEAL, bold=True)
    add_notes(slide, 70, "The model's suggestion is a candidate, not the result.",
              "Phase one captures the baseline. Phase two supplies a tool finding, relevant code, and a narrow instruction to the model. Phase three is human review, tests, and a fresh run of the same measuring tool. Phase four compares the branch result with the baseline using the original decision rule. This prevents an attractive patch from being called successful just because it looks plausible.",
              "Current report §3.1.1 and Figures 3.1-3.4.")

    # 8. Dataset and environment
    slide = header(prs, "Method", "One controlled environment and a seeded dataset", "Current report Tables 3.2 and 3.11; Appendix H", TEAL)
    metric(slide, 0.78, 1.58, 3.77, "Accounts", "120", "Deterministic seed", NAVY_LIGHT, NAVY)
    metric(slide, 4.78, 1.58, 3.77, "Journal entries", "30,000", "At least 5,000 posted", NAVY_LIGHT, NAVY)
    metric(slide, 8.78, 1.58, 3.77, "Test host", "Windows 11", ".NET 8 · local SQL Server", NAVY_LIGHT, NAVY)
    box(slide, 0.78, 3.27, 11.77, 2.45, WHITE, BORDER, radius=True)
    text(slide, "Tool versions in the fresh reproduction", 1.06, 3.57, 11.05, 0.42, 21, TEAL, bold=True)
    text(slide, "SonarQube 26.7     |     OWASP ZAP 2.17     |     Apache JMeter 5.5", 1.06, 4.22, 11.04, 0.60, 24, INK, bold=True)
    text(slide, "A single host and one fresh run per branch limit generalization.", 1.06, 5.03, 11.04, 0.35, 17, MUTED)
    add_notes(slide, 60, "The comparison uses a reproducible local setup but only one host.",
              "The seeded database contains 120 accounts and thirty thousand journal entries, with at least five thousand posted entries. The benchmark used one Windows 11 host, ASP.NET Core 8, local SQL Server, and pinned tool versions. These controls help repeat the run, but one host and one fresh run per branch are still important limitations.",
              "Current report Tables 3.2 and 3.11, Appendix H and verification summaries.")

    # 9. Acceptance rules
    slide = header(prs, "Method", "Success was decided by tool-specific gates", "Current report Table 4.1 and §4.1", TEAL)
    table(slide, [
        ["Area", "Primary decision", "Control metric"],
        ["SonarQube", "Close fresh issues; Quality Gate OK", "Coverage and complexity disclosed"],
        ["OWASP ZAP", "Configured-rule gate; report raw alerts", "Residual HTTP-only observation disclosed"],
        ["JMeter", "50/100/500 VU core gate", "100 VU ≤500 ms; 500 VU ≤1200 ms"],
        ["Regression", "Focused tests and role checks", "Accounting and authorization preserved"],
    ], 0.80, 1.58, 11.72, 3.80, widths=[0.21,0.45,0.34], font_size=17)
    box(slide, 0.80, 5.62, 11.72, 0.96, ORANGE_LIGHT, radius=True)
    text(slide, "100 VU is peak use—keep p95 within 0.5 s; 500 VU is stress—allow up to 1.2 s. Study-defined targets.",
         1.04, 5.76, 11.18, 0.70, 18, ORANGE, bold=True)
    add_notes(slide, 70, "A result is not called successful because one chart looks better.",
              "Static analysis has an issue count and a Quality Gate, but coverage and complexity also matter. ZAP has a configured gate, while raw observations still have to be disclosed. JMeter has primary 50, 100, and 500 VU core checks. The 100 VU profile represents peak use, so the study sets a p95 target of five hundred milliseconds: 95 percent of requests should complete within half a second. The 500 VU profile is stress beyond normal use, so the study permits a longer p95 of twelve hundred milliseconds while still requiring bounded degradation. These are engineering acceptance choices set in the test plan before the retest, not universal industry standards or a validated production SLA. We report the raw response times as well as the gate judgment.",
              "Current report Table 4.1 and §4.1; Document/JMETER_ZAP_TEST_PLAN.md §4.6.")

    # 10. Human-in-the-loop boundary
    slide = header(prs, "Method", "Codex proposed; humans and tools decided", "Current report §§3.1.3-3.1.4; Appendices I and K", TEAL)
    stages = ["Native tool finding", "Bounded code context", "Candidate patch", "Human review", "Tests", "Tool rerun"]
    for i, stage in enumerate(stages):
        x = 0.78 + i*2.08
        box(slide, x, 2.35, 1.87, 1.24, TEAL_LIGHT if i<3 else GREEN_LIGHT, BORDER, radius=True)
        text(slide, stage, x+0.13, 2.53, 1.61, 0.89, 17, TEAL if i<3 else GREEN, bold=True, align=PP_ALIGN.CENTER)
        if i<5:text(slide, "→", x+1.88, 2.66, 0.20, 0.30, 21, MUTED, bold=True)
    box(slide, 1.20, 4.18, 10.92, 1.34, WHITE, BORDER, radius=True)
    text(slide, "Every accepted interpretation is linked to a branch, a tested commit, and a machine-readable run.",
         1.54, 4.47, 10.24, 0.77, 22, INK, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, 70, "The model is an advisor under a human-controlled test loop.",
              "The prompt contains the scanner or benchmark finding, a small relevant source excerpt, and constraints such as keeping tests and contracts intact. Codex suggests a patch. A human reviews it, the focused test suite runs, and the original tool is rerun. A patch can still miss its measured target; the JMeter result is the clearest example of that boundary.",
              "Current report §§3.1.3-3.1.4, Appendices I and K.")

    # 11. What changed in code
    slide = header(prs, "Method", "The implemented fixes had different mechanisms", "Current report Table 4.3; Appendices E-G", TEAL)
    table(slide, [
        ["Track", "Representative change", "What it was meant to address"],
        ["SonarQube", "Eight query parameters → JournalEntryQueryDto", "Controller maintainability finding S107"],
        ["ZAP", "SecurityHeadersMiddleware before Swagger", "Browser-facing security headers"],
        ["JMeter", "Ledger caching; avoid report write amplification", "Repeated report calculation and database pressure"],
    ], 0.80, 1.65, 11.72, 3.58, widths=[0.20,0.43,0.37], font_size=18)
    text(slide, "Mechanism and measured outcome are shown separately on the next slides.", 0.83, 5.65, 11.72, 0.50, 20, TEAL, bold=True)
    add_notes(slide, 70, "Similar-looking code changes can have very different measured outcomes.",
              "For maintainability, a request DTO replaced a controller action with eight query parameters. For security, header middleware moved ahead of Swagger and UI exposure was constrained. For performance, the JMeter branch added ledger caching and reduced report persistence work. These are mechanisms, not proof of success by themselves; the tool results on the next slides show what actually happened.",
              "Current report Table 4.3 and Appendices E-G.")

    # 12. SonarQube result
    slide = header(prs, "Results", "SonarQube closed all 17 fresh issues", "Current report §4.2; Table 4.2; Figures L.1-L.4", GREEN)
    metric(slide, 0.80, 1.50, 3.76, "Unique issues", "17 → 0", "Current run", GREEN_LIGHT, GREEN)
    metric(slide, 4.78, 1.50, 3.76, "Quality Gate", "OK → OK", "Both branches pass", GREEN_LIGHT, GREEN)
    metric(slide, 8.76, 1.50, 3.76, "Coverage", "74.4 → 74.3%", "Down 0.1 percentage point", ORANGE_LIGHT, ORANGE)
    box(slide, 0.80, 3.14, 11.72, 2.30, WHITE, BORDER, radius=True)
    text(slide, "Do not add impact counts together", 1.09, 3.45, 10.90, 0.40, 22, TEAL, bold=True)
    text(slide, "The 17 distinct findings carried overlapping Reliability 2, Security 1, and Maintainability 16 impacts. Cyclomatic complexity rose 886 → 889; duplication stayed 0.0%.",
         1.09, 4.03, 10.88, 0.99, 20, INK)
    add_notes(slide, 80, "The fresh issue count reached zero, with small control-metric regressions disclosed.",
              "SonarQube found seventeen distinct issues on the shared baseline and none on the SonarQube remediation branch. Both Quality Gates were OK. Coverage slipped from 74.4 to 74.3 percent and cyclomatic complexity rose by three, so I do not call those measures unchanged. The newer Reliability, Security, and Maintainability impact counts overlap; they are not nineteen separate issues.",
              "Current report Table 4.2, Appendix E, Figures L.1-L.4.",
              "Legacy Bugs and Vulnerabilities are both zero; the MQR Security impact is a different classification.")

    # 13. Native Sonar evidence
    slide = header(prs, "Results", "Native SonarQube views support the metric table", "Fresh Sep 22 capture · Appendix L, Figures L.1 and L.3", GREEN)
    text(slide, "Baseline · commit 7a0469d9", 0.88, 1.42, 5.62, 0.38, 17, NAVY, bold=True)
    text(slide, "Remediation · commit a2279bcb", 6.82, 1.42, 5.62, 0.38, 17, GREEN, bold=True)
    image_contain(slide, DASHBOARD_ROOT / "sonarqube" / "sonarqube-baseline-overview.png", 0.86, 1.88, 5.63, 3.85)
    image_contain(slide, DASHBOARD_ROOT / "sonarqube" / "sonarqube-remediation-overview.png", 6.82, 1.88, 5.63, 3.85)
    box(slide, 0.84, 5.99, 11.64, 0.59, GREEN_LIGHT, radius=True)
    text(slide, "Fresh source screenshots; use Appendix L for the larger overview and issue views.", 1.07, 6.07, 11.19, 0.41, 16, GREEN, bold=True)
    add_notes(slide, 60, "The tool views correspond to the pinned baseline and remediation commits.",
              "These are native SonarQube captures from the fresh reproduction, not a chart invented for the presentation. The baseline overview shows Security, Reliability, and Maintainability impacts plus 74.4 percent coverage. The remediation overview shows zero open issues and 74.3 percent coverage. Appendix L also includes the issue-list views if the committee asks for rule-level detail.",
              "Appendix L Figures L.1-L.4 and dashboard-capture manifest.")

    # 14. Fresh ZAP results
    passive_before = next(x for x in zap["baseline"]["scans"] if x["scan"] == "baseline")
    passive_after = next(x for x in zap["remediation"]["scans"] if x["scan"] == "baseline")
    api_before = next(x for x in zap["baseline"]["scans"] if x["scan"] == "api")
    api_after = next(x for x in zap["remediation"]["scans"] if x["scan"] == "api")
    slide = header(prs, "Results", "ZAP passed the configured gate; raw API Medium remains", "Current report §4.3; fresh ZAP summaries; Appendix F", RED)
    table(slide, [
        ["Fresh scan", "Baseline Medium / Low", "After Medium / Low", "Interpretation"],
        ["Passive", f"{passive_before['medium']} / {passive_before['low']}", f"{passive_after['medium']} / {passive_after['low']}", "Both severities cleared"],
        ["Authenticated API", f"{api_before['medium']} / {api_before['low']}", f"{api_after['medium']} / {api_after['low']}", "One raw Medium disclosed"],
    ], 0.80, 1.63, 11.72, 2.60, widths=[0.20,0.23,0.23,0.34], font_size=18)
    metric(slide, 0.80, 4.66, 5.66, "Configured rules", "PASS", "Both branches", GREEN_LIGHT, GREEN)
    metric(slide, 6.85, 4.66, 5.66, "Residual raw alert", "HTTP Only Site", "Local HTTP transport · CWE-311", ORANGE_LIGHT, ORANGE)
    add_notes(slide, 80, "The configured ZAP rule gate passed, but one raw Medium alert remains.",
              "In the fresh passive scan, Medium went from three to zero and Low from six to zero. In the authenticated API scan, Low went from three to zero, but Medium remained one. That observation is HTTP Only Site on the local HTTP transport. The configured business-endpoint rule set passed on both branches, but I do not say every raw Medium alert was eliminated.",
              "Current report §4.3, fresh ZAP baseline/remediation summary.json, Appendix F.",
              "Do not use the historical dashboard screenshot as proof of zero fresh API Medium alerts.")

    # 15. ZAP capture provenance
    slide = header(prs, "Results", "ZAP dashboard captures are historical visuals", "Appendix L, Figures L.5-L.8; fresh counts on prior slide", RED)
    image_contain(slide, DASHBOARD_ROOT / "zap" / "zap-api-before-summary.png", 0.85, 1.62, 5.65, 3.87)
    image_contain(slide, DASHBOARD_ROOT / "zap" / "zap-api-after-summary.png", 6.82, 1.62, 5.65, 3.87)
    text(slide, "Historical API before", 0.91, 5.61, 5.50, 0.32, 17, MUTED, bold=True)
    text(slide, "Historical API after", 6.90, 5.61, 5.50, 0.32, 17, MUTED, bold=True)
    box(slide, 0.84, 6.10, 11.65, 0.53, ORANGE_LIGHT, radius=True)
    text(slide, "Current scan JSON is authoritative: the API still has 1 raw Medium after remediation.", 1.05, 6.15, 11.20, 0.39, 16, ORANGE, bold=True)
    add_notes(slide, 70, "A dashboard screenshot and a fresh measurement can have different provenance.",
              "The side-by-side ZAP screenshots were rendered from earlier scan reports, and I label them historical on purpose. They show the tool's native alert summary and evidence layout. The current numerical claim comes from the fresh machine-readable scan JSON: one raw authenticated API Medium remains. This is why the thesis separates visual captures from current gate judgments.",
              "Dashboard-capture manifest, Appendix L Figures L.7-L.8, fresh ZAP summary.json.")

    # 16. JMeter workload definitions
    slide = header(prs, "Results", "JMeter measured three core profiles and two stress profiles", "Current report §§3.8 and 4.4; Appendix K", ORANGE)
    for i, (label, value, detail) in enumerate([
        ("50 VU", "50 users", "Normal load"), ("100 VU", "100 users", "Higher concurrency"),
        ("500 VU", "500 users", "Heavy concurrency"), ("soak", "sustained", "Long-running behavior"),
        ("spike", "burst", "Surge and recovery"),
    ]):
        x = 0.78 + i*2.49
        box(slide, x, 2.00, 2.30, 2.57, WHITE, BORDER, radius=True)
        text(slide, label.upper(), x+0.18, 2.27, 1.94, 0.45, 25, ORANGE, bold=True, align=PP_ALIGN.CENTER)
        text(slide, value, x+0.18, 3.04, 1.94, 0.45, 20, INK, bold=True, align=PP_ALIGN.CENTER)
        text(slide, detail, x+0.16, 3.79, 1.98, 0.48, 14, MUTED, align=PP_ALIGN.CENTER)
    text(slide, "50/100/500 VU are workload sizes; p95 is a latency percentile.", 0.82, 5.35, 11.78, 0.65, 24, ORANGE, bold=True)
    add_notes(slide, 60, "Stress improvements cannot stand in for the core gate.",
              "The 50, 100, and 500 virtual-user profiles exercise increasingly heavy core load. These are workload sizes, whereas p50, p90, p95, and p99 are response-time percentiles measured in milliseconds. Archived JMeter folder IDs retain p50, p100, and p500 for traceability. Soak and spike add sustained and sudden-load evidence; they are additional interpretation, not a substitute for the primary core gate.",
              "Current report §§3.8 and 4.4, Appendix K.")

    # 17. JMeter core target not met
    slide = header(prs, "Results", "Core performance target was not met", "Current report §4.4; Table 4.5; fresh JMeter summary.json", RED)
    rows = [["Profile", "Baseline p95", "After p95", "Fresh judgment"]]
    for profile in ("p50", "p100", "p500"):
        before = baseline_profiles[profile]["p95Ms"]
        after = after_profiles[profile]["p95Ms"]
        rows.append([PROFILE_DISPLAY[profile], f"{before:g} ms", f"{after:g} ms", "Slower after fixes"])
    table(slide, rows, 0.80, 1.55, 11.72, 3.40, widths=[0.19,0.24,0.24,0.33], font_size=20)
    box(slide, 0.80, 5.26, 11.72, 1.07, RED_LIGHT, radius=True)
    text(slide, "NOT MET: 100 VU 1,146.9 > 500 ms; 500 VU 1,547.95 > 1,200 ms", 1.04, 5.43, 11.21, 0.67, 24, RED, bold=True)
    text(slide, "100 VU = peak responsiveness; 500 VU = stress tolerance. Study-defined, not universal.",
         1.04, 6.52, 11.15, 0.51, 16, MUTED)
    add_notes(slide, 90, "The primary performance criterion was not achieved.",
              "All three fresh core total-p95 values increased. At 50 VU, p95 rose from thirty-three to about two hundred milliseconds; at 100 VU, from about two hundred fifty-six to one thousand one hundred forty-seven; at 500 VU, from one hundred six to one thousand five hundred forty-eight. The different limits are engineering acceptance choices: 100 VU represents peak use, where the study asks for p95 within half a second; 500 VU is stress, where up to 1.2 seconds is allowed. They are not universal standards or a validated production SLA. Errors stayed at zero, but the measured p95 values exceeded both limits, so the predefined core performance target was not met. The raw machine-readable gate status is FAIL.",
              "Document/JMETER_ZAP_TEST_PLAN.md §4.6; fresh JMeter baseline/remediation summary.json; current report §4.4 and Table 4.5.",
              "Do not carry the historical core-gate PASS into this fresh result.")

    # 18. Stress-profile improvements
    soak_before, soak_after = baseline_profiles["soak"], after_profiles["soak"]
    spike_before, spike_after = baseline_profiles["spike"], after_profiles["spike"]
    slide = header(prs, "Results", "Soak and spike improved, but they are secondary", "Current report §4.4; fresh JMeter summary.json", ORANGE)
    box(slide, 0.80, 1.70, 5.66, 3.72, GREEN_LIGHT, BORDER, radius=True)
    text(slide, "SOAK TOTAL p95", 1.10, 2.04, 5.06, 0.32, 16, GREEN, bold=True)
    text(slide, f"{soak_before['p95Ms']:g} → {soak_after['p95Ms']:g} ms", 1.10, 2.68, 5.05, 0.62, 30, GREEN, bold=True)
    text(slide, f"{percent_lower(soak_before['p95Ms'],soak_after['p95Ms'])}% lower on the fresh run", 1.10, 3.65, 5.05, 0.52, 21, INK)
    box(slide, 6.84, 1.70, 5.66, 3.72, GREEN_LIGHT, BORDER, radius=True)
    text(slide, "SPIKE TOTAL p95", 7.14, 2.04, 5.05, 0.32, 16, GREEN, bold=True)
    text(slide, f"{spike_before['p95Ms']:,.0f} → {spike_after['p95Ms']:,.0f} ms", 7.14, 2.68, 5.05, 0.62, 29, GREEN, bold=True)
    text(slide, f"{percent_lower(spike_before['p95Ms'],spike_after['p95Ms'])}% lower; errors {spike_before['errorPct']}% → {spike_after['errorPct']}%", 7.14, 3.65, 5.05, 0.52, 19, INK)
    text(slide, "Core 100 and 500 VU limits were exceeded; stress gains do not change the decision.", 0.84, 5.79, 11.70, 0.63, 20, RED, bold=True)
    add_notes(slide, 75, "Stress gains are real but do not meet the primary core target.",
              "The soak aggregate p95 fell from 443.95 to 171 milliseconds. Spike aggregate p95 fell from about 33.4 seconds to 4.73 seconds, and spike errors fell to zero. These are meaningful fresh stress-profile gains, but the core 100 VU and 500 VU target remains unmet. I therefore call performance mixed, not generally improved.",
              "Fresh JMeter baseline/remediation summary.json; current report §4.4.")

    # 19. Endpoint-level tradeoff
    slide = header(prs, "Results", "One endpoint improved under spike but regressed under soak", "Current report Table 4.4; Appendix G", ORANGE)
    table(slide, [
        ["Account-ledger profile", "Baseline p95", "After p95", "Interpretation"],
        ["Soak", "741.95 ms", "3,825.70 ms", "Regression"],
        ["Spike", "50,288.30 ms", "12,910.20 ms", "Improvement"],
    ], 0.80, 1.62, 11.72, 2.54, widths=[0.29,0.22,0.22,0.27], font_size=19)
    box(slide, 0.80, 4.49, 11.72, 1.72, WHITE, BORDER, radius=True)
    text(slide, "Why this matters", 1.07, 4.73, 10.90, 0.38, 19, ORANGE, bold=True)
    text(slide, "Caching and query changes can help one workload while making another worse. Endpoint results and core aggregate gates must both be read.",
         1.07, 5.14, 10.88, 0.77, 20, INK)
    add_notes(slide, 70, "The account-ledger endpoint is a concrete counterexample to a universal speedup claim.",
              "Under the fresh soak run, account-ledger p95 rose from about 742 to 3,826 milliseconds. Under spike it fell from about 50.3 seconds to 12.9 seconds. The summary report endpoints also improved under spike, but this one endpoint shows why a code mechanism or one workload cannot be generalized to all load conditions.",
              "Current report Table 4.4 and Appendix G.")

    # 20. User-visible application evidence
    slide = header(prs, "Results", "The API was exercised through a role-aware client", "Current report §4.6, Figures 4.2 and 4.7", NAVY)
    image_contain(slide, WEB_ROOT / "02-dashboard-admin.png", 0.82, 1.63, 5.69, 3.79, visible_top_px=1000)
    image_contain(slide, WEB_ROOT / "07-journal-entry-details.png", 6.83, 1.63, 5.69, 3.79, visible_top_px=1000)
    text(slide, "Admin dashboard", 0.94, 5.56, 5.49, 0.33, 18, NAVY, bold=True)
    text(slide, "Balanced journal-entry lines", 6.96, 5.56, 5.47, 0.33, 18, NAVY, bold=True)
    box(slide, 0.84, 6.03, 11.65, 0.58, NAVY_LIGHT, radius=True)
    text(slide, "These are browser workflow screenshots; their displayed counts are not benchmark metrics.", 1.04, 6.10, 11.19, 0.39, 16, NAVY, bold=True)
    add_notes(slide, 70, "The application evidence demonstrates the user-visible API path.",
              "A React client signs in with JWT, shows roles, calls the accounting API, and presents reports and journal-entry lines. The dashboard and expanded journal detail shown here are captured from that workflow. They are not performance measurements, and their live displayed counts should not be confused with the fixed JMeter seed dataset. Other browser views, including auditor restrictions, are in the appendix.",
              "Current report §4.6 and Figures 4.2-4.7.")

    # 21. Decision matrix and limits
    slide = header(prs, "Synthesis", "The decision is deliberately asymmetric", "Current report Table 4.5; Appendix N", TEAL)
    table(slide, [
        ["Area", "Decision", "Why"],
        ["Code quality", "Achieved", "17 issues → 0; gate OK; small coverage/complexity regressions"],
        ["Security", "Configured gate passed", "Passive cleared; raw API Medium transport alert remains"],
        ["Core performance", "Not achieved", "100 VU and 500 VU p95 exceeded thresholds"],
        ["Stress behavior", "Partially achieved", "Soak and spike aggregates improved"],
    ], 0.80, 1.49, 11.72, 4.12, widths=[0.20,0.28,0.52], font_size=17)
    text(slide, "Limits: one API · one host · one run per branch · CPU/RAM and service tiers not controlled", 0.83, 5.92, 11.68, 0.45, 19, MUTED)
    add_notes(slide, 80, "The thesis supports some improvements, not a blanket success claim.",
              "Code quality met its primary criterion while coverage and cyclomatic complexity had small regressions. ZAP's configured rule gate passed, but the raw HTTP-only Medium remained. JMeter core performance did not meet its gate, though stress aggregates improved. This is one API on one host with one fresh run per branch. The dated evidence does not preserve CPU and RAM specifications or utilization, application and database resource limits, or a managed service tier. These results should be replicated under recorded and matched capacity before broader generalization.",
              "Current report Table 4.5 and Appendix N.")

    # 22. Conclusion
    slide = header(prs, "Closing", "The strongest result is the controlled audit-to-fix method", "Current report Chapter 5; Appendices A-N", TEAL)
    for index, (number, title, body, color) in enumerate([
        ("01", "The model can help", "Seventeen fresh static issues closed; security rule gate passed.", GREEN),
        ("02", "Measurement can reject it", "The primary JMeter target was not met despite stress gains.", RED),
        ("03", "Provenance makes it auditable", "Pinned branches, tests, dashboard captures, and raw records remain available.", TEAL),
    ]):
        y = 1.51 + index*1.58
        box(slide, 0.80, y, 11.72, 1.35, WHITE, BORDER, radius=True)
        text(slide, number, 1.05, y+0.31, 0.70, 0.51, 26, color, bold=True)
        text(slide, title, 1.95, y+0.15, 10.20, 0.43, 22, color, bold=True)
        text(slide, body, 1.95, y+0.68, 10.20, 0.45, 18, INK)
    text(slide, "Thank you. Questions?", 0.82, 6.36, 11.72, 0.39, 25, TEAL, bold=True)
    add_notes(slide, 80, "Human review and repeatable tools decide which AI suggestions count.",
              "My answer is qualified. ChatGPT through Codex 5.4 helped close fresh static issues and pass the configured security gate. The same structured workflow also showed that the primary performance target was not met, rather than hiding it behind better stress charts. The contribution is a traceable method for asking the model to improve existing code and then allowing independent evidence to accept, limit, or reject the claim. Thank you; I welcome questions.",
              "Current report Chapter 5 and Appendices A-N.")

    return len(prs.slides)


def backup_notes(slide, title, source, caveat=""):
    add_notes(slide, 0, f"Backup evidence: {title}.",
              "Open this slide only if the committee asks for the underlying records. Read the classification and source line before interpreting the screenshot or table.",
              source, caveat)


def dashboard_pair(prs, title, before_path: str, after_path: str, summary: str, source: str,
                   classification: str, accent=TEAL):
    slide = header(prs, "Backup · native dashboards", title, source, accent)
    text(slide, "BASELINE", 0.88, 1.40, 5.62, 0.34, 17, NAVY, bold=True)
    text(slide, "REMEDIATION", 6.83, 1.40, 5.62, 0.34, 17, accent, bold=True)
    image_contain(slide, DASHBOARD_ROOT / before_path, 0.86, 1.84, 5.64, 3.97)
    image_contain(slide, DASHBOARD_ROOT / after_path, 6.82, 1.84, 5.64, 3.97)
    box(slide, 0.84, 5.99, 11.64, 0.68, ORANGE_LIGHT if classification == "rendered-historical" else TEAL_LIGHT, radius=True)
    text(slide, summary, 1.02, 6.06, 11.28, 0.49, 17,
         ORANGE if classification == "rendered-historical" else accent, bold=True)
    before = facts_capture(before_path)
    after = facts_capture(after_path)
    backup_notes(slide, title, source,
                 f"Capture classification: {classification}. Baseline {before['branch']} @ {before['commit'][:8]}, run {before['runId']}; remediation {after['branch']} @ {after['commit'][:8]}, run {after['runId']}.")


CAPTURES: dict[str, dict] = {}


def facts_capture(relative_path: str) -> dict:
    if relative_path not in CAPTURES:
        raise RuntimeError(f"Dashboard image is not registered in the capture manifest: {relative_path}")
    return CAPTURES[relative_path]


def app_pair(prs, title, left_file: str, left_label: str, right_file: str | None,
             right_label: str | None, source: str):
    slide = header(prs, "Backup · application evidence", title, source, NAVY)
    if right_file:
        image_contain(slide, WEB_ROOT / left_file, 0.85, 1.82, 5.65, 4.28, visible_top_px=1000)
        image_contain(slide, WEB_ROOT / right_file, 6.82, 1.82, 5.65, 4.28, visible_top_px=1000)
        text(slide, left_label, 0.89, 6.20, 5.57, 0.39, 16, NAVY, bold=True)
        text(slide, right_label or "", 6.87, 6.20, 5.57, 0.39, 16, NAVY, bold=True)
    else:
        image_contain(slide, WEB_ROOT / left_file, 1.20, 1.63, 10.85, 4.79, visible_top_px=1000)
        text(slide, left_label, 1.24, 6.39, 10.80, 0.28, 16, NAVY, bold=True)
    backup_notes(slide, title, source,
                 "These are browser workflow captures; visible live application counts are not JMeter benchmark metrics.")


def build_appendix_slides(prs: Presentation, facts: dict) -> None:
    sonar = facts["sonar"]
    zap = facts["zap"]
    jmeter = facts["jmeter"]
    sb, sa = profile_map(jmeter["baseline"]), profile_map(jmeter["remediation"])

    # 23. Evidence map
    slide = header(prs, "Backup evidence", "Where each defense claim can be checked", "Current report Appendices A-N and dashboard-capture manifest", NAVY)
    table(slide, [
        ["Claim", "Machine-readable record", "Report evidence"],
        ["Static issues", "sonar/*/summary.json", "Table 4.2 · Appendix E · Figures L.1-L.4"],
        ["ZAP raw/gate", "zap/*/summary.json", "§4.3 · Appendix F · Figures L.5-L.8"],
        ["JMeter p95/gate", "jmeter/*/summary.json", "§4.4 · Appendix G · Figures L.9-L.18"],
        ["Working API", "role/browser run records", "§4.6 · Figures 4.2-4.7"],
    ], 0.80, 1.52, 11.72, 4.26, widths=[0.23,0.30,0.47], font_size=16)
    backup_notes(slide, "evidence map", "Current report Appendices A-N and dashboard-capture manifest.")

    # 24. Provenance
    slide = header(prs, "Backup evidence", "Exact baseline and remediation commits", "Current report Appendix A; Table 3.1", NAVY)
    table(slide, [
        ["Role", "Branch", "Tested commit"],
        ["Shared baseline", "baseline-v0.1", "7a0469d9"],
        ["Static quality", "baseline-sonarqube-v1", "a2279bcb"],
        ["Security", "baseline-zap-v1", "5f3bd3cc"],
        ["Performance", "baseline-jmeter-v1", "78b08b62"],
        ["Reproduction", "codex/thesis-reproducibility-v1", "efa19cea"],
    ], 0.80, 1.51, 11.72, 4.85, widths=[0.25,0.47,0.28], font_size=17)
    backup_notes(slide, "branch provenance", "Current report Appendix A and Table 3.1.")

    # 25. Full Sonar metrics
    slide = header(prs, "Backup evidence", "Fresh SonarQube metric detail", "Current report Table 4.2; fresh SonarQube summaries", GREEN)
    rows = [["Metric", "Baseline", "After", "Meaning"]]
    important = [
        "Total unique issues", "Code smells (legacy type)", "Bugs (legacy type)",
        "Vulnerabilities (legacy type)", "Reliability impacts (MQR)*", "Security impacts (MQR)*",
        "Maintainability impacts (MQR)*", "Security hotspots", "Coverage",
        "Duplicated lines", "Cyclomatic complexity", "Cognitive complexity", "Quality Gate",
    ]
    interpretation = {
        "Total unique issues": "Distinct issue count", "Code smells (legacy type)": "Legacy type",
        "Bugs (legacy type)": "Legacy type", "Vulnerabilities (legacy type)": "Legacy type",
        "Reliability impacts (MQR)*": "Overlapping impact", "Security impacts (MQR)*": "Overlapping impact",
        "Maintainability impacts (MQR)*": "Overlapping impact", "Security hotspots": "Scanner finding",
        "Coverage": "Down 0.1 point", "Duplicated lines": "Unchanged", "Cyclomatic complexity": "Up 3",
        "Cognitive complexity": "Unchanged", "Quality Gate": "OK on both",
    }
    for name in important:
        before, after = facts["sonar_rows"][name]
        rows.append([name.replace(" (legacy type)", "").replace(" impacts (MQR)*", " impact*"), before, after, interpretation[name]])
    table(slide, rows, 0.79, 1.40, 11.75, 5.39, widths=[0.36,0.15,0.15,0.34], font_size=12)
    backup_notes(slide, "full fresh SonarQube metrics", "Current report Table 4.2; fresh SonarQube summary.json.",
                 "* MQR impacts overlap; do not sum them as unique issues.")

    # 26. Full ZAP raw outcomes
    slide = header(prs, "Backup evidence", "Fresh ZAP raw alert counts and gate", "Current report §4.3; fresh ZAP summaries; Appendix F", RED)
    rows = [["Scan", "Role", "High", "Medium", "Low", "Info", "Gate"]]
    for role in ("baseline", "remediation"):
        for scan in zap[role]["scans"]:
            rows.append(["Passive" if scan["scan"] == "baseline" else "Auth API", role,
                         str(scan["high"]), str(scan["medium"]), str(scan["low"]),
                         str(scan["informational"]), "PASS"])
    table(slide, rows, 0.79, 1.48, 11.75, 4.18, widths=[0.19,0.20,0.10,0.12,0.10,0.10,0.19], font_size=16)
    box(slide, 0.83, 5.91, 11.66, 0.63, ORANGE_LIGHT, radius=True)
    text(slide, "Raw API Medium after remediation = 1 (HTTP Only Site); configured rule gate = PASS.",
         1.03, 5.99, 11.20, 0.42, 17, ORANGE, bold=True)
    backup_notes(slide, "full fresh ZAP outcomes", "Current report §4.3 and fresh ZAP summary.json.",
                 "Do not infer fresh zero-Medium API outcome from historical dashboards.")

    # 27. Full JMeter profiles
    slide = header(prs, "Backup evidence", "Fresh JMeter profile matrix", "Current report §4.4; fresh JMeter summaries", ORANGE)
    rows = [["Profile", "Baseline p95", "After p95", "Before err", "After err", "After tx/s"]]
    for profile in ("p50", "p100", "p500", "soak", "spike"):
        before, after = sb[profile], sa[profile]
        rows.append([PROFILE_DISPLAY.get(profile, profile), f"{before['p95Ms']:,.2f}", f"{after['p95Ms']:,.2f}",
                     f"{before['errorPct']:.4f}%", f"{after['errorPct']:.4f}%", f"{after['throughput']:.2f}"])
    table(slide, rows, 0.79, 1.53, 11.75, 4.53, widths=[0.14,0.19,0.19,0.17,0.17,0.14], font_size=16)
    text(slide, "Baseline core gate PASS; remediation core gate FAIL (100 VU and 500 VU thresholds).", 0.83, 6.29, 11.67, 0.41, 17, RED, bold=True)
    backup_notes(slide, "full fresh JMeter matrix", "Fresh JMeter baseline/remediation summary.json and current report §4.4.")

    # 28. Source-code example: S107
    slide = header(prs, "Backup evidence", "S107 source: eight parameters to one DTO", "Current report Listing E.1; Tables 3.4 and 4.3", TEAL)
    box(slide, 0.80, 1.55, 5.70, 3.97, WHITE, BORDER, radius=True)
    box(slide, 6.82, 1.55, 5.70, 3.97, WHITE, BORDER, radius=True)
    text(slide, "BEFORE · baseline-v0.1", 1.03, 1.78, 5.22, 0.34, 17, RED, bold=True)
    text(slide, "AFTER · baseline-sonarqube-v1", 7.05, 1.78, 5.22, 0.34, 17, GREEN, bold=True)
    text(slide, "GetPaged(\n  from, to, status, accountId,\n  search, page, pageSize, sort\n)\n\n// 8 query parameters", 1.03, 2.31, 5.16, 2.72, 18, INK, font=CODE_FONT, valign=MSO_ANCHOR.TOP)
    text(slide, "GetPaged(\n  JournalEntryQueryDto query\n)\n\n// DTO carries the query inputs", 7.05, 2.31, 5.15, 2.72, 18, INK, font=CODE_FONT, valign=MSO_ANCHOR.TOP)
    text(slide, "Measured result: fresh SonarQube issues 17 → 0; the route remains GET /journal-entries.",
         0.86, 5.83, 11.58, 0.58, 18, TEAL, bold=True)
    backup_notes(slide, "S107 controller DTO change", "Current report Listing E.1, Table 4.3, Appendix E.")

    # 29. Source-code paths for security and performance
    slide = header(prs, "Backup evidence", "Source examples: security order and report persistence", "Current report Listings F.1-F.2 and G.1-G.3", TEAL)
    box(slide, 0.80, 1.55, 5.70, 4.55, WHITE, BORDER, radius=True)
    box(slide, 6.82, 1.55, 5.70, 4.55, WHITE, BORDER, radius=True)
    text(slide, "SECURITY · API/Program.cs", 1.04, 1.82, 5.12, 0.36, 17, RED, bold=True)
    text(slide, "SecurityHeadersMiddleware\n→ before Swagger/OpenAPI\n\nSwagger UI exposed only when the separate UI flag allows it.",
         1.04, 2.37, 5.13, 2.70, 20, INK, valign=MSO_ANCHOR.TOP)
    text(slide, "PERFORMANCE · FinancialReportService.cs", 7.04, 1.82, 5.17, 0.36, 17, ORANGE, bold=True)
    text(slide, "Account-ledger payload cache\n+ avoid snapshot write pressure\n\nFresh load behavior remains mixed; core gate FAIL.",
         7.04, 2.37, 5.15, 2.70, 20, INK, valign=MSO_ANCHOR.TOP)
    text(slide, "A code-level mechanism is evidence of implementation, not evidence that a load gate passed.",
         0.85, 6.24, 11.70, 0.42, 16, TEAL, bold=True)
    backup_notes(slide, "security and performance source paths", "Current report Appendices F and G.")

    # 30. Role/invariant checks
    slide = header(prs, "Backup evidence", "Authorization and accounting controls remained testable", "Current report Appendices C, D, I and §4.6", NAVY)
    table(slide, [
        ["Control", "Evidence path", "Expected boundary"],
        ["JWT / roles", "Controller attributes + RBAC matrix", "Protected routes reject unauthorized roles"],
        ["Journal posting", "Balanced debit-credit source", "Closed periods and unequal totals rejected"],
        ["Audit trail", "Service + browser log view", "Actor and sensitive action recorded"],
        ["Browser roles", "Admin / Auditor screenshots", "Navigation and read-only behavior visible"],
    ], 0.79, 1.50, 11.75, 4.69, widths=[0.23,0.35,0.42], font_size=17)
    backup_notes(slide, "authorization and accounting evidence", "Current report Appendices C, D, I and Figures 4.2-4.7.")

    # 31-39. Every tool dashboard capture appears once, paired by profile/role.
    dashboard_pair(prs, "SonarQube overview · baseline vs remediation",
                   "sonarqube/sonarqube-baseline-overview.png", "sonarqube/sonarqube-remediation-overview.png",
                   "Fresh: 17 unique issues → 0; Quality Gate OK/OK; coverage 74.4% → 74.3%.",
                   "Appendix L Figures L.1 and L.3", "fresh-reproduction", GREEN)
    dashboard_pair(prs, "SonarQube issues · baseline vs remediation",
                   "sonarqube/sonarqube-baseline-issues.png", "sonarqube/sonarqube-remediation-issues.png",
                   "Fresh issue-list views: 17 baseline findings and 0 after remediation.",
                   "Appendix L Figures L.2 and L.4", "fresh-reproduction", GREEN)
    dashboard_pair(prs, "ZAP passive scan · historical dashboard capture",
                   "zap/zap-baseline-before-summary.png", "zap/zap-baseline-after-summary.png",
                   "Historical screenshots; fresh passive result: Medium 3 / Low 6 → 0 / 0.",
                   "Appendix L Figures L.5 and L.6; fresh ZAP summary.json", "rendered-historical", RED)
    dashboard_pair(prs, "ZAP authenticated API · historical dashboard capture",
                   "zap/zap-api-before-summary.png", "zap/zap-api-after-summary.png",
                   "Historical screenshots; fresh API result: Medium 1 / Low 3 → Medium 1 / Low 0.",
                   "Appendix L Figures L.7 and L.8; fresh ZAP summary.json", "rendered-historical", RED)
    for profile, number in [("p50",9),("p100",10),("p500",11),("soak",12),("spike",13)]:
        before, after = sb[profile], sa[profile]
        judgment = "CORE FAIL" if profile in ("p100", "p500") else ("STRESS GAIN" if profile in ("soak", "spike") else "CORE SLOWER")
        dashboard_pair(prs, f"JMeter {PROFILE_DISPLAY.get(profile, profile)} · baseline vs remediation",
                       f"jmeter/jmeter-baseline-{profile}-dashboard.png",
                       f"jmeter/jmeter-remediation-{profile}-dashboard.png",
                       f"Fresh total p95 {before['p95Ms']:g} → {after['p95Ms']:g} ms · {judgment}.",
                       f"Appendix L Figures L.{number} and L.{number+5}; fresh JMeter summary.json",
                       "fresh-reproduction", ORANGE)

    # 40-42. Browser evidence without the credential-bearing login screen.
    app_pair(prs, "Application: administration and reporting", "02-dashboard-admin.png", "Admin dashboard",
             "03-reports-trial-balance.png", "Trial-balance report", "Current report Figures 4.2 and 4.3")
    app_pair(prs, "Application: audit and role restrictions", "04-audit-logs.png", "Audit-log view",
             "05-auditor-role-view.png", "Auditor navigation", "Current report Figures 4.4 and 4.5")
    app_pair(prs, "Application: evidence and balanced journal lines", "06-research-evidence.png", "Research-evidence map",
             "07-journal-entry-details.png", "Expanded debit and credit lines", "Current report Figures 4.6 and 4.7")


def build() -> Path:
    facts = read_facts()
    CAPTURES.clear()
    CAPTURES.update(facts["captures"])
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Final Thesis Defense - Current Fresh Results"
    prs.core_properties.subject = "Branch-isolated evidence for financial accounting API remediation"
    prs.core_properties.author = "Zaw Ye Htut Ko"
    prs.core_properties.keywords = "SonarQube, OWASP ZAP, JMeter, Codex 5.4, financial API"
    main_count = build_main_slides(prs, facts)
    if main_count != 22:
        raise RuntimeError(f"Expected 22 timed slides, built {main_count}")
    build_appendix_slides(prs, facts)
    if len(prs.slides) != 42:
        raise RuntimeError(f"Expected 42 total slides, built {len(prs.slides)}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
