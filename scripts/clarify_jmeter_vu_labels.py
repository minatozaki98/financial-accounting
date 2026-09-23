from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-23SEPv2-1.docx"
PROFILE_NAMES = {"p50": "50 VU", "p100": "100 VU", "p500": "500 VU"}
PROFILE_PATTERN = re.compile(r"(?<![A-Za-z0-9_])(?:p50|p100|p500)(?![A-Za-z0-9_])")
PERCENTILE_NOTE = (
    "In this report, 50 VU, 100 VU, and 500 VU name concurrent-user workloads. "
    "By contrast, p50, p90, p95, and p99 denote response-time percentiles measured in milliseconds, not percentages."
)
ARTIFACT_NOTE = (
    "Archived JMeter folder names retain the historical run IDs p50, p100, and p500; "
    "those IDs refer to 50 VU, 100 VU, and 500 VU runs and are not latency percentiles."
)


def normalized(text: str) -> str:
    return " ".join(text.split())


def replace_profiles(text: str) -> str:
    text = text.replace("p50/p100/p500 pass", "50, 100, and 500 VU profiles pass")
    text = text.replace("p50/p100/p500", "50, 100, and 500 VU")
    text = text.replace("50, 100, and 500 VU pass", "50, 100, and 500 VU profiles pass")
    text = text.replace("p50/p500", "50 and 500 VU")
    return PROFILE_PATTERN.sub(lambda match: PROFILE_NAMES[match.group(0)], text)


def replace_in_runs(paragraph: Paragraph) -> int:
    original = paragraph.text
    if not PROFILE_PATTERN.search(original) and "50, 100, and 500 VU pass" not in original:
        return 0
    for run in paragraph.runs:
        replacement = replace_profiles(run.text)
        if replacement != run.text:
            run.text = replacement
    if PROFILE_PATTERN.search(paragraph.text):
        raise RuntimeError(f"Could not safely rewrite a split profile label: {original[:120]}")
    return 1


def set_acronym(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    paragraph.style = "Normal"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def clarify_report(doc: DocumentObject) -> Counter[str]:
    paragraphs = doc.paragraphs
    acronym_heading = next(i for i, p in enumerate(paragraphs) if p.text == "List of Acronyms")
    body_start = next(i for i, p in enumerate(paragraphs) if p.text == "Chapter 1 - Introduction")
    appendices = next(i for i, p in enumerate(paragraphs) if p.text == "APPENDICES")
    glossary = paragraphs[acronym_heading + 1 : body_start]
    by_term = {p.text.split(" - ", 1)[0]: p for p in glossary if " - " in p.text}
    for term in ("p50", "p95", "p99", "VU"):
        if term not in by_term:
            raise RuntimeError(f"Missing existing acronym entry: {term}")
    set_acronym(by_term["p50"], "p50 - 50th-percentile response latency (median)")
    set_acronym(by_term["p95"], "p95 - 95th-percentile response latency")
    set_acronym(by_term["p99"], "p99 - 99th-percentile response latency")
    set_acronym(by_term["VU"], "VU - Virtual user (simulated concurrent JMeter user)")
    if "p90" not in by_term:
        new_p = OxmlElement("w:p")
        by_term["p50"]._p.addnext(new_p)
        set_acronym(Paragraph(new_p, by_term["p50"]._parent), "p90 - 90th-percentile response latency")

    changed: Counter[str] = Counter()
    # Abstract and main-body prose. Generated front-matter lists are refreshed from captions later.
    for paragraph in paragraphs[:acronym_heading] + paragraphs[body_start:appendices]:
        if "Archived JMeter folder names retain" in paragraph.text or "By contrast, p50, p90, p95, and p99" in paragraph.text:
            continue
        if replace_in_runs(paragraph):
            changed["prose"] += 1

    # Only reader-facing appendix labels; machine-readable paths and code remain byte-for-byte.
    for paragraph in paragraphs[appendices + 1 :]:
        text = normalized(paragraph.text)
        if paragraph.style.name == "Appendix Code" or text.startswith("Classification:"):
            continue
        if text.startswith(("Purpose:", "Verification:", "L.", "Figure L.")):
            if replace_in_runs(paragraph):
                changed["appendix_labels"] += 1

    for table in doc.tables:
        if len(table.rows) == 1 and len(table.columns) == 1:
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_in_runs(paragraph):
                        changed["table_cells"] += 1

    performance_explanation = next(
        p for p in paragraphs if normalized(p.text).startswith("How to read the performance evidence.")
    )
    if PERCENTILE_NOTE not in performance_explanation.text:
        performance_explanation.add_run(" " + PERCENTILE_NOTE)
        changed["explanations"] += 1
    tool_explanation = next(
        p for p in paragraphs if normalized(p.text).startswith("Apache JMeter works like a controlled crowd")
    )
    if ARTIFACT_NOTE not in tool_explanation.text:
        tool_explanation.add_run(" " + ARTIFACT_NOTE)
        changed["explanations"] += 1
    return changed


def clarify_file(path: Path) -> Counter[str]:
    doc = Document(path)
    changed = clarify_report(doc)
    doc.save(path)
    return changed


def main() -> None:
    parser = argparse.ArgumentParser(description="Distinguish JMeter virtual-user profiles from latency percentiles.")
    parser.add_argument("target", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    path = args.target.resolve()
    print(path)
    print(dict(clarify_file(path)))


if __name__ == "__main__":
    main()
