import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

src = r'c:\Users\Asus\Documents\GitHub\school\financial-accounting\Document\(Paitoon)Proposal_latest_6519692_ZAWYEHTUTKO (1).docx'
dst = r'c:\Users\Asus\Documents\GitHub\school\financial-accounting\Document\(Paitoon)Proposal_latest_6519692_ZAWYEHTUTKO_UPDATED.docx'

doc = Document(src)
changes = []

def replace_paragraph_text(para, new_text, label=""):
    if para.runs:
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""
        changes.append(f"[P] {label}")
    else:
        para.text = new_text
        changes.append(f"[P-norun] {label}")

def replace_in_runs(para, old, new, label=""):
    full = para.text
    if old in full:
        for run in para.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                changes.append(f"[R] {label}: '{old[:40]}' -> '{new[:40]}'")
                return True
        new_full = full.replace(old, new)
        if para.runs:
            para.runs[0].text = new_full
            for run in para.runs[1:]:
                run.text = ""
        changes.append(f"[R-multi] {label}: '{old[:40]}' -> '{new[:40]}'")
        return True
    return False

# ─── TABLE 0: Academic Year ───
for row in doc.tables[0].rows:
    for cell in row.cells:
        if "2/2024" in cell.text:
            for p in cell.paragraphs:
                replace_in_runs(p, "2/2024", "2/2025", "Academic Year")

# ─── Cover page date ───
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "January 2025":
        replace_paragraph_text(p, "March 2026", f"P{i} Cover date")

# ─── "ChatGPT 4" -> "ChatGPT" everywhere ───
for i, p in enumerate(doc.paragraphs):
    replace_in_runs(p, "ChatGPT 4", "ChatGPT", f"P{i}")
for row in doc.tables[0].rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            replace_in_runs(p, "ChatGPT 4", "ChatGPT", "Table0")

# ─── Abstract ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("In today") and "two API projects" in t:
        replace_paragraph_text(p, (
            "In today\u2019s digital landscape, RESTful APIs have become indispensable for enabling "
            "communication across diverse software systems, from e-commerce to financial applications. "
            "Despite their widespread use, securing and optimizing these APIs remains a complex challenge, "
            "with common vulnerabilities exposing them to threats such as SQL injection and Cross-Site "
            "Scripting (XSS). This thesis explores the potential of ChatGPT, a large language model by "
            "OpenAI, to assist in improving the security, performance, and maintainability of ASP.NET Core "
            "RESTful APIs. Using a financial accounting API as the study subject\u2014a complex project with "
            "endpoints for journal entries, chart of accounts, financial reporting, user authentication, "
            "and audit logging\u2014this study investigates the impact of ChatGPT\u2019s code recommendations on "
            "key metrics, including code security, performance, and code quality. Key tools employed include "
            "SonarQube for static code analysis, Apache JMeter for performance testing, and OWASP ZAP for "
            "vulnerability assessment. The study follows a branch-based testing methodology: baseline metrics "
            "are captured first, then ChatGPT-guided fixes are applied on separate branches (one per tool), "
            "and re-testing is performed to measure improvements. The outcomes of this study demonstrate "
            "ChatGPT\u2019s capability in addressing security and optimization issues in RESTful APIs, providing "
            "valuable insights into the role of AI-driven assistance in modern software development."
        ), f"P{i} Abstract")

# ─── Problem Statement: "two API projects" ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if "two API projects" in t and "one simple and one complex" in t:
        new_text = t.replace(
            "two API projects\u2014one simple and one complex\u2014showing",
            "a financial accounting API project, demonstrating"
        )
        replace_paragraph_text(p, new_text, f"P{i} Problem Statement")

# ─── Section 3: "What is my work?" ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("My scope focuses on") and "two projects" in t:
        replace_paragraph_text(p, (
            "My scope focuses on improving the security, performance, and code quality of ASP.NET Core "
            "RESTful APIs by using ChatGPT to generate recommendations for code enhancements. The study "
            "involves evaluating a financial accounting API project by comparing baseline results with the "
            "outcomes after applying ChatGPT-guided improvements. The API includes endpoints for user "
            "authentication (JWT-based), chart of accounts management, journal entry processing (including "
            "bulk operations, posting, and reversal), accounting period management, financial reporting "
            "(trial balance, profit & loss, balance sheet, account ledger), and audit logging. Key tools "
            "employed include SonarQube for static code analysis, Apache JMeter for performance testing, "
            "and OWASP ZAP for vulnerability assessment."
        ), f"P{i} Scope")

# ─── Phase 1 description ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Initial security and performance testing are conducted on two"):
        replace_paragraph_text(p, (
            "Initial security and performance testing are conducted on the financial accounting "
            "ASP.NET Core RESTful API project. The API includes endpoints for authentication, chart of "
            "accounts, journal entries, accounting periods, financial reports, and audit logs. These "
            "baseline metrics are captured on the baseline-v0.1 branch and provide reference points for "
            "evaluating the effectiveness of improvements."
        ), f"P{i} Phase1")

# ─── Tools line ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t == "Tools: SonarQube for code quality, Postman and OWASP ZAP for security vulnerabilities, and Apache JMeter for performance.":
        replace_paragraph_text(p, (
            "Tools: SonarQube for code quality and static analysis, OWASP ZAP for security vulnerability "
            "scanning (baseline scan and authenticated API scan), and Apache JMeter for performance testing "
            "(load, soak, and spike test profiles at 50, 100, and 500 concurrent users)."
        ), f"P{i} Tools")

# ─── Example endpoints ───
ENDPOINT_REPLACEMENTS = {
    "Simple API Project: A basic CRUD (Create, Read, Update, Delete) API for managing employee records.":
        "Financial Accounting API: A complex RESTful API for financial accounting with multi-tier architecture (API, BAL/Services, MODEL/Entities, Tests).",
    "Example Endpoint: GET /employees/{id} to retrieve an employee record.":
        "Example Endpoints: GET /accounts to list chart of accounts, POST /journal-entries to create journal entries, GET /reports/trial-balance to generate trial balance reports.",
    "Complex API Project: A financial management API with endpoints for transaction handling and reporting.":
        "The API features role-based access control (Admin, FinanceManager, User, Auditor), JWT authentication, and comprehensive audit logging.",
    "Example Endpoint: POST /transactions to add a transaction, requiring input validation for amounts and currencies.":
        "Example Endpoint: POST /journal-entries/bulk to create bulk journal entries with debit/credit line validation, requiring authenticated access and finance role authorization.",
}
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t in ENDPOINT_REPLACEMENTS:
        replace_paragraph_text(p, ENDPOINT_REPLACEMENTS[t], f"P{i} endpoint")

# ─── Phase 2: ChatGPT improvements ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Code recommendations are sought using ChatGPT for addressing"):
        replace_paragraph_text(p, (
            "Code recommendations are generated using ChatGPT for addressing detected vulnerabilities and "
            "performance bottlenecks. Each tool\u2019s fixes are applied on a dedicated branch: "
            "baseline-sonarqube-v1 for code quality issues, baseline-zap-v1 for security vulnerabilities, "
            "and baseline-jmeter-v1 for performance optimizations. This branch-based approach isolates "
            "changes per tool for clear before/after comparison."
        ), f"P{i} Phase2")

# ─── Phase 3: Post-improvement testing ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("With ChatGPT-guided changes implemented, the projects undergo"):
        replace_paragraph_text(p, (
            "With ChatGPT-guided changes implemented on each branch, the project undergoes a new round of "
            "testing using the same tools and configurations as the baseline. Security, performance, and "
            "code quality metrics are collected again and compared to the baseline-v0.1 results."
        ), f"P{i} Phase3")
    elif t == "Security: Test POST /transactions for SQL injection attacks using OWASP ZAP.":
        replace_paragraph_text(p, (
            "Security: Run ZAP baseline scan and authenticated API scan against all endpoints (auth, "
            "accounts, journal-entries, periods, reports, audit-logs) on the baseline-zap-v1 branch."
        ), f"P{i} Phase3-sec")
    elif t.startswith("Performance: Simulate 500 concurrent users accessing GET /reports"):
        replace_paragraph_text(p, (
            "Performance: Run JMeter test matrix (50, 100, 500 concurrent users + soak + spike profiles) "
            "against all API endpoints on the baseline-jmeter-v1 branch."
        ), f"P{i} Phase3-perf")

# ─── Phase 4: Comparative analysis ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("The data collected from both phases (baseline and post-improvement)"):
        replace_paragraph_text(p, (
            "The data collected from the baseline (baseline-v0.1) and post-improvement branches "
            "(baseline-sonarqube-v1, baseline-zap-v1, baseline-jmeter-v1) are analyzed and compared to "
            "draw conclusions on ChatGPT\u2019s effectiveness in enhancing the security, performance, and "
            "code quality of the financial accounting RESTful API. Results are compiled into a "
            "baseline-comparison-report documenting improvements across all three dimensions."
        ), f"P{i} Phase4")

# ─── Section 3.3 Tool descriptions ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("SonarQube: Used for assessing code quality"):
        replace_paragraph_text(p, (
            "SonarQube (Community Edition): Used for assessing code quality and security vulnerabilities by "
            "scanning the .NET codebase for bugs, code smells, and security vulnerabilities. Runs via "
            "dotnet-sonarscanner with unit and integration test coverage analysis."
        ), f"P{i} SQ-tool")
    elif t.startswith("Apache JMeter: Measures performance metrics"):
        replace_paragraph_text(p, (
            "Apache JMeter: Measures performance metrics such as response times, throughput, and error "
            "rates across multiple load profiles (50, 100, 500 concurrent users), soak testing, and spike "
            "testing. Test plans target all API endpoints including authentication, CRUD operations, and "
            "reporting."
        ), f"P{i} JM-tool")
    elif t.startswith("Postman and OWASP ZAP:"):
        replace_paragraph_text(p, (
            "OWASP ZAP: Provides comprehensive vulnerability scanning through baseline scans (passive "
            "analysis of all endpoints) and authenticated API scans (active testing with JWT tokens and "
            "role-based access). Custom rule configurations are maintained in TSV files for reproducible "
            "scan profiles."
        ), f"P{i} ZAP-tool")

# ─── Section 3.5 Testing Phases ───
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t == "Baseline Testing: The initial state of each project is tested for security, performance, and code quality.":
        replace_paragraph_text(p, (
            "Baseline Testing: The initial state of the financial accounting API is tested for security "
            "(ZAP baseline + API scan), performance (JMeter load/soak/spike), and code quality (SonarQube). "
            "Results are captured on the baseline-v0.1 branch."
        ), f"P{i} Baseline-testing")
    elif t.startswith("Post-Improvement Testing: After implementing ChatGPT"):
        replace_paragraph_text(p, (
            "Post-Improvement Testing: After implementing ChatGPT\u2019s improvements on dedicated branches "
            "(baseline-sonarqube-v1, baseline-zap-v1, baseline-jmeter-v1), each branch is re-tested with "
            "the same tool and configuration to assess the efficacy of the changes. All issues and fixes "
            "are documented in per-tool fix tracking documents."
        ), f"P{i} Post-testing")

# ─── Chapter 4 sections ───
CHAPTER4_REPLACEMENTS = {
    "Vulnerability Scanning: OWASP ZAP will scan all API endpoints for known vulnerabilities. The focus will be on high-risk issues like injection flaws and broken authentication mechanisms.":
        "Vulnerability Scanning: OWASP ZAP performs both baseline (passive) and authenticated API (active) scans across all endpoints: /auth, /accounts, /journal-entries, /periods, /reports, and /audit-logs. Focus areas include injection flaws, broken authentication, missing security headers (CSP, X-Content-Type-Options, CORS), and vulnerable dependencies.",
    "Authentication and Authorization Testing: Postman is used to simulate user roles and permissions, testing the authorization mechanisms of the API. Unauthorized access attempts will be executed to verify the robustness of the API's access control.":
        "Authentication and Authorization Testing: OWASP ZAP authenticated API scan is used to test role-based access control (Admin, FinanceManager, User, Auditor roles) and JWT authentication. Unauthorized access attempts against protected endpoints (e.g., POST /accounts, POST /periods/{id}/close, DELETE /journal-entries/{id}) verify the robustness of the API\u2019s access control.",
    "Response Time and Throughput: Apache JMeter will measure the response times and throughput for each API endpoint under different load conditions (e.g., 50, 100, 500 concurrent users).":
        "Response Time and Throughput: Apache JMeter measures response times (average, p95, p99) and throughput for each API endpoint under load profiles of 50, 100, and 500 concurrent users, plus soak (sustained load) and spike (burst) test profiles.",
    "Resource Utilization: CPU and memory usage are tracked during load tests to assess the resource efficiency of the API.":
        "Resource Utilization: CPU and memory usage are tracked during load tests. Error rates are monitored to identify endpoints that fail under stress.",
    "Baseline Analysis: The initial codebase is scanned to identify code smells, potential bugs, and security issues.":
        "Baseline Analysis: The initial codebase on baseline-v0.1 is scanned with SonarQube (via dotnet-sonarscanner) to identify code smells, potential bugs, security vulnerabilities, and security hotspots across all five modules (API, BAL, MODEL, REPOSITORY, Tests).",
    "Security: Evaluate the percentage reduction in vulnerabilities and their severity post-improvement.":
        "Security: Evaluate the reduction in ZAP alerts by severity level (High, Medium, Low, Informational) between baseline-v0.1 and baseline-zap-v1. Document specific alerts resolved and any remaining issues.",
    "Code Quality: Analyze improvements in code smells, cyclomatic complexity, and code readability.":
        "Code Quality: Compare SonarQube results between baseline-v0.1 and baseline-sonarqube-v1. Key metrics: bugs, vulnerabilities, code smells, security hotspots, cyclomatic complexity, maintainability index, and quality gate status.",
}

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t in CHAPTER4_REPLACEMENTS:
        replace_paragraph_text(p, CHAPTER4_REPLACEMENTS[t], f"P{i}")

# Partial matches for Chapter 4
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Comparison of Pre- and Post-Improvement Results: Security test results"):
        replace_paragraph_text(p, (
            "Comparison of Pre- and Post-Improvement Results: Security test results from baseline-v0.1 and "
            "baseline-zap-v1 branches are compared, analyzing how ChatGPT\u2019s recommendations mitigated "
            "identified vulnerabilities. Alert counts are compared by severity level (High, Medium, Low, "
            "Informational)."
        ), f"P{i} Sec-compare")
    elif t.startswith("Comparison of Performance Metrics: Baseline and post-improvement"):
        replace_paragraph_text(p, (
            "Comparison of Performance Metrics: Results from baseline-v0.1 and baseline-jmeter-v1 branches "
            "are compared to determine the impact of ChatGPT\u2019s suggestions on performance, focusing on "
            "improvements in response times, error rate reduction, throughput gains, and stability under "
            "sustained and spike load conditions."
        ), f"P{i} Perf-compare")
    elif t.startswith("Post-Improvement Analysis: After applying ChatGPT"):
        replace_paragraph_text(p, (
            "Post-Improvement Analysis: After applying ChatGPT\u2019s recommendations on baseline-sonarqube-v1, "
            "SonarQube is run again with the same configuration to evaluate improvements. This second "
            "analysis identifies reductions in bugs, code smells, security vulnerabilities, and security "
            "hotspots. Quality gate status is compared."
        ), f"P{i} CQ-post")
    elif t.startswith("Performance: Compare response times, throughput"):
        replace_paragraph_text(p, (
            "Performance: Compare JMeter results between baseline-v0.1 and baseline-jmeter-v1 across all "
            "load profiles (p50, p100, p500, soak, spike). Key metrics: average response time, p95/p99 "
            "latency, throughput (req/s), and error rate."
        ), f"P{i} CA-perf")

# ─── Remaining "Postman" references ───
for i, p in enumerate(doc.paragraphs):
    replace_in_runs(p, "Postman and OWASP ZAP", "OWASP ZAP", f"P{i}")
    replace_in_runs(p, "Postman, OWASP ZAP", "OWASP ZAP", f"P{i}")
    replace_in_runs(p, "(OWASP ZAP, Postman)", "(OWASP ZAP)", f"P{i}")
    replace_in_runs(p, "OWASP ZAP, Postman", "OWASP ZAP", f"P{i}")
    # Standalone Postman in tool context
    if "Postman" in p.text and "OWASP" not in p.text:
        replace_in_runs(p, "Postman", "OWASP ZAP", f"P{i}")

# ─── Save ───
doc.save(dst)

print(f"\nSaved to: {os.path.basename(dst)}")
print(f"Total changes: {len(changes)}\n")
print("--- Change Log ---")
for c in changes:
    print(c)
