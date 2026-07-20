from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "Document" / "outputs"

CAMERA_READY_SOURCE = OUTPUTS / "camera-ready-manuscript.docx"
FINAL_REPORT_SOURCE = OUTPUTS / "final-report.docx"

CAMERA_READY_DOCX = OUTPUTS / "camera-ready-manuscript-gpt55-comparison.docx"
CAMERA_READY_PDF = OUTPUTS / "camera-ready-manuscript-gpt55-comparison.pdf"
FINAL_REPORT_DOCX = OUTPUTS / "final-report-gpt55-comparison.docx"
FINAL_REPORT_PDF = OUTPUTS / "final-report-gpt55-comparison.pdf"
NOTES_MD = OUTPUTS / "gpt55-comparison-deliverables.md"

SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.com")


def normalized_text(paragraph: Paragraph) -> str:
    return " ".join(paragraph.text.split())


def find_paragraph(doc: Document, startswith: str, *, last: bool = False) -> Paragraph:
    matches = [
        paragraph
        for paragraph in doc.paragraphs
        if normalized_text(paragraph).startswith(startswith)
    ]
    if not matches:
        raise RuntimeError(f"Could not find paragraph starting with {startswith!r}")
    return matches[-1] if last else matches[0]


def insert_paragraph_before(
    paragraph: Paragraph,
    text: str = "",
    style: str | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def append_paragraph_after(
    paragraph: Paragraph,
    text: str = "",
    style: str | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def append_table_after(paragraph: Paragraph, rows: list[list[str]], *, style: str = "Table Grid") -> Table:
    doc = paragraph.part.document
    table = doc.add_table(rows=0, cols=len(rows[0]))
    try:
        table.style = style
    except KeyError:
        pass

    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value

    paragraph._p.addnext(table._tbl)
    return table


def append_paragraph_after_table(
    table: Table,
    text: str = "",
    style: str | None = None,
) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    new_para = Paragraph(new_p, table._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def make_header_row_bold(table: Table) -> None:
    if not table.rows:
        return
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def set_table_font(table: Table, size: int = 8) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(size)


def set_table_widths(table: Table, widths: list[float]) -> None:
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
        for cell in column.cells:
            cell.width = Inches(width)


def add_note(anchor: Paragraph, text: str, *, italic: bool = False) -> Paragraph:
    paragraph = append_paragraph_after(anchor, text)
    for run in paragraph.runs:
        run.italic = italic
    return paragraph


def add_labeled_paragraph(anchor: Paragraph, label: str, body: str) -> Paragraph:
    paragraph = append_paragraph_after(anchor)
    label_run = paragraph.add_run(label)
    label_run.bold = True
    paragraph.add_run(body)
    return paragraph


def add_camera_ready_addendum() -> None:
    shutil.copy2(CAMERA_READY_SOURCE, CAMERA_READY_DOCX)
    doc = Document(CAMERA_READY_DOCX)

    target = find_paragraph(doc, "Threats To Validity")
    anchor = insert_paragraph_before(target, "Model Upgrade Benchmark: GPT-5.4 vs GPT-5.5", "Heading 2")
    anchor = append_paragraph_after(
        anchor,
        "After the accepted ChatGPT Codex 5.4 study, a second blinded benchmark was run with "
        "GPT-5.5 on the same Financial Accounting API pre-fix source lineage. The GPT-5.5 "
        "experiment started from commit 7a0469d, used the same SonarQube, OWASP ZAP, and "
        "Apache JMeter 5.5 measurement workflow, and kept one isolated remediation branch per "
        "tool before comparing against normalized GPT-5.4 reruns.",
    )

    anchor = add_labeled_paragraph(
        anchor,
        "SonarQube. ",
        "Both models reached quality gate OK with code smells reduced from 21 to 0 and "
        "coverage at 74.3%; this is a tie on final static-quality outcome, with GPT-5.5 "
        "using a smaller source patch.",
    )
    anchor = add_labeled_paragraph(
        anchor,
        "OWASP ZAP. ",
        "GPT-5.4 was cleaner in the normalized rerun: baseline/API residuals were 2/4 "
        "alert categories and 5/110 instances, compared with GPT-5.5 residuals of 5/7 "
        "categories and 12/330 instances. Both passed business-endpoint high/medium gates.",
    )
    anchor = add_labeled_paragraph(
        anchor,
        "Apache JMeter. ",
        "Both models completed all five profiles with 0% errors and passed drift gates. "
        "GPT-5.5 was more balanced overall: p50 p99 was 48 ms and spike p99 was 4226 ms, "
        "while GPT-5.4 had a p50 p99 regression to 1086 ms but remained competitive on "
        "selected p100, p500, and soak tail metrics.",
    )

    anchor = append_paragraph_after(
        anchor,
        "The model-upgrade result does not replace the original 5.4 contribution. It adds a "
        "controlled follow-up: GPT-5.5 matched SonarQube quality with less code, was weaker "
        "than GPT-5.4 on ZAP residual cleanliness, and produced the more balanced JMeter "
        "performance remediation. The strongest validity constraint is process asymmetry: "
        "GPT-5.4 was retested from historical fix commits, while GPT-5.5 was evaluated live "
        "with complete interaction logs.",
    )

    add_note(
        anchor,
        "Supplemental provenance: GPT-5.5 branches were codex/gpt-5.5-baseline, "
        "codex/gpt-5.5-sonarqube-v1, codex/gpt-5.5-zap-v1, codex/gpt-5.5-jmeter-v1, "
        "and codex/gpt-5.5-comparison. GPT-5.4 branch sources were baseline-sonarqube-v1, "
        "baseline-zap-v1, and baseline-jmeter-v1; the JMeter tested code commit was 78b08b6 "
        "inside that branch because the branch head later received a documentation commit.",
        italic=True,
    )

    doc.save(CAMERA_READY_DOCX)


def add_full_report_addendum() -> None:
    shutil.copy2(FINAL_REPORT_SOURCE, FINAL_REPORT_DOCX)
    doc = Document(FINAL_REPORT_DOCX)

    chapter_5 = find_paragraph(doc, "Chapter 5: Conclusion", last=True)
    chapter_5.paragraph_format.page_break_before = True

    anchor = insert_paragraph_before(chapter_5, "")
    anchor.add_run().add_break(WD_BREAK.PAGE)
    anchor = append_paragraph_after(anchor, "4.7 GPT-5.4 vs GPT-5.5 Remediation Benchmark", "Heading 2")
    anchor = append_paragraph_after(
        anchor,
        "This section adds the model-upgrade benchmark requested after the camera-ready "
        "deliverables. It compares the earlier ChatGPT Codex 5.4 remediation branches against "
        "new GPT-5.5 remediation branches on the same Financial Accounting API pre-fix source "
        "lineage. The comparison uses normalized reruns so the older GPT-5.4 branch code is "
        "measured with the same local machine, database restore, Docker image identities, and "
        "benchmark scripts used for GPT-5.5.",
    )

    anchor = append_paragraph_after(anchor, "4.7.1 Controls and Branch Provenance", "Heading 3")
    anchor = append_paragraph_after(
        anchor,
        "The common pre-fix source commit was 7a0469d9401a3061941b44fcd46b3beca1c9c729. "
        "GPT-5.5 was blinded from GPT-5.4 branches, commits, reflogs, reports, prompts, and "
        "patches until all GPT-5.5 fix branches were committed and pushed. The normalized rerun "
        "date was 2026-07-19. The clean SQL Server restore point was financial-gpt55-clean.bak.",
    )

    provenance_rows = [
        ["Track", "Branch or source", "Commit"],
        ["Common pre-fix code", "historical pre-fix parent", "7a0469d9401a3061941b44fcd46b3beca1c9c729"],
        ["GPT-5.5 baseline evidence", "codex/gpt-5.5-baseline", "ec034038c284e77efa4ed2310b13fc9662259286"],
        ["GPT-5.5 SonarQube", "codex/gpt-5.5-sonarqube-v1", "19763f15cbc8412eb74ff5b9d706be69456e62ff"],
        ["GPT-5.5 OWASP ZAP", "codex/gpt-5.5-zap-v1", "aa1a9bd7cefe4f3a07c3bd10d0bde7d5409e3f4b"],
        ["GPT-5.5 JMeter", "codex/gpt-5.5-jmeter-v1", "8d0be3b918008f1cf38a7c61669f8d7609df0963"],
        ["GPT-5.5 comparison report", "codex/gpt-5.5-comparison", "d1baa7afc4c79a0177a7534235c55cf3c97d8a91"],
        ["GPT-5.4 SonarQube normalized code", "baseline-sonarqube-v1", "a2279bcbdc20751529335cf72f8baac1bc6f994b"],
        ["GPT-5.4 OWASP ZAP normalized code", "baseline-zap-v1", "5f3bd3ccd9e0d460f52cac6a8a0d5fdceff1ce3d"],
        ["GPT-5.4 JMeter normalized code", "baseline-jmeter-v1 tested code commit; branch head is docs commit 3c9392d", "78b08b62570dcf6fe436354e8e6b770ab2074445"],
    ]
    table = append_table_after(anchor, provenance_rows)
    make_header_row_bold(table)
    set_table_font(table, 7)
    set_table_widths(table, [1.55, 1.75, 3.0])

    anchor = append_paragraph_after_table(table, "4.7.2 SonarQube Outcome Comparison", "Heading 3")
    anchor = append_paragraph_after(
        anchor,
        "Both models reduced SonarQube code smells from 21 to 0 and kept the quality gate OK. "
        "Coverage was effectively unchanged: 74.4% before fixes and 74.3% after each model's "
        "remediation. The practical result is a tie on measured static quality, with GPT-5.5 "
        "reaching the same gate using a smaller source/test patch.",
    )
    sonar_rows = [
        ["Model", "Quality gate", "Coverage", "Code smells", "Bugs", "Vulnerabilities", "NCLOC"],
        ["Before fixes", "OK", "74.4", "21", "0", "0", "3211"],
        ["GPT-5.4 normalized", "OK", "74.3", "0", "0", "0", "3198"],
        ["GPT-5.5 final", "OK", "74.3", "0", "0", "0", "3199"],
    ]
    table = append_table_after(anchor, sonar_rows)
    make_header_row_bold(table)
    set_table_font(table, 8)
    set_table_widths(table, [1.55, 1.0, 0.75, 0.8, 0.55, 1.0, 0.65])

    anchor = append_paragraph_after_table(table, "4.7.3 OWASP ZAP Outcome Comparison", "Heading 3")
    anchor = append_paragraph_after(
        anchor,
        "ZAP was run in both unauthenticated baseline mode and authenticated OpenAPI/API mode. "
        "The counts below include informational categories, so pass/fail is interpreted "
        "separately from raw alert volume. GPT-5.4 produced cleaner normalized output. GPT-5.5 "
        "removed the major business-endpoint security-header problems, but retained more "
        "non-business residual alerts such as Swagger/static-asset and local HTTP transport warnings.",
    )
    zap_rows = [
        ["Scan", "Before fixes", "GPT-5.4 normalized", "GPT-5.5 final"],
        ["Baseline alert categories", "11", "2", "5"],
        ["Baseline alert instances", "35", "5", "12"],
        ["API alert categories", "9", "4", "7"],
        ["API alert instances", "332", "110", "330"],
        ["Gate interpretation", "Security headers exposed", "Pass for business endpoints", "Pass for business endpoints; more residual alerts"],
    ]
    table = append_table_after(anchor, zap_rows)
    make_header_row_bold(table)
    set_table_font(table, 8)
    set_table_widths(table, [1.55, 1.2, 1.45, 1.75])

    anchor = append_paragraph_after_table(table, "4.7.4 Apache JMeter Outcome Comparison", "Heading 3")
    anchor = append_paragraph_after(
        anchor,
        "All compared JMeter runs completed with 0.00% error rate in all five workload profiles. "
        "GPT-5.4 was slightly faster than GPT-5.5 on selected p100, p500, and soak tail metrics, "
        "but it introduced a large p50 tail-latency regression. GPT-5.5 produced the more "
        "balanced performance result and the stronger spike p99/throughput outcome.",
    )
    jmeter_rows = [
        ["Profile", "Metric", "Before fixes", "GPT-5.4 normalized", "GPT-5.5 final"],
        ["p50", "p95 ms", "24", "229", "23"],
        ["p50", "p99 ms", "62", "1086", "48"],
        ["p50", "throughput/s", "53.34", "53.95", "55.02"],
        ["p100", "p95 ms", "114", "51", "52"],
        ["p100", "p99 ms", "585", "127", "141"],
        ["p500", "p95 ms", "88", "40", "40"],
        ["p500", "p99 ms", "170", "92", "99"],
        ["soak", "p95 ms", "409", "117", "143"],
        ["soak", "p99 ms", "892", "284", "287"],
        ["spike", "p95 ms", "4509", "2594", "2482"],
        ["spike", "p99 ms", "6431", "4932", "4226"],
    ]
    table = append_table_after(anchor, jmeter_rows)
    make_header_row_bold(table)
    set_table_font(table, 7)
    set_table_widths(table, [0.75, 1.0, 1.1, 1.45, 1.25])

    drift_anchor = append_paragraph_after_table(table, "Drift gate results:", "Normal")
    drift_rows = [
        ["Run", "Soak drift", "Spike drift", "Gate result"],
        ["Before fixes", "-14.44%", "44.16%", "Fail: spike recovery drift"],
        ["GPT-5.4 normalized", "-64.99%", "1.01%", "Pass"],
        ["GPT-5.5 final", "-29.70%", "-23.69%", "Pass"],
    ]
    table = append_table_after(drift_anchor, drift_rows)
    make_header_row_bold(table)
    set_table_font(table, 8)
    set_table_widths(table, [1.5, 1.0, 1.0, 2.35])

    anchor = append_paragraph_after_table(table, "4.7.5 Patch Size, Tests, and Process Evidence", "Heading 3")
    anchor = append_paragraph_after(
        anchor,
        "Source/test-only diff statistics exclude generated evidence under Document/**. "
        "GPT-5.5 generally used smaller patches than GPT-5.4 for the same measured tool area. "
        "The strongest process evidence exists for GPT-5.5 because those branches were generated "
        "live with complete Codex task logs; the historical GPT-5.4 prompt timing and attempt "
        "metadata are incomplete.",
    )
    patch_rows = [
        ["Track", "GPT-5.4 source/test patch", "GPT-5.5 source/test patch", "Verification evidence"],
        ["SonarQube", "13 files, 148 insertions, 160 deletions", "13 files, 86 insertions, 99 deletions", "Both final scans gate OK with 0 code smells."],
        ["OWASP ZAP", "7 files, 104 insertions, 24 deletions", "4 files, 38 insertions, 6 deletions", "Both passed business-endpoint high/medium gates."],
        ["JMeter", "13 files, 783 insertions, 56 deletions", "8 files, 316 insertions, 79 deletions", "Both passed 0% error and drift gates after normalized rerun."],
    ]
    table = append_table_after(anchor, patch_rows)
    make_header_row_bold(table)
    set_table_font(table, 7)
    set_table_widths(table, [1.0, 1.45, 1.45, 2.1])

    anchor = append_paragraph_after_table(table, "4.7.6 Interpretation and Limitations", "Heading 3")
    append_paragraph_after(
        anchor,
        "The model comparison does not make GPT-5.5 a universal winner. SonarQube is a tie; "
        "ZAP residual cleanliness favors GPT-5.4; JMeter favors GPT-5.5 when stable behavior "
        "across all profiles and smaller implementation risk are weighted more heavily than "
        "selected p100/p500/soak p99 wins. The most defensible conclusion is that GPT-5.5 "
        "matched or improved the maintainability/performance workflow with smaller patches, "
        "but the security outcome still requires tighter Swagger/static-asset hardening.",
    )

    doc.save(FINAL_REPORT_DOCX)


def write_notes() -> None:
    NOTES_MD.write_text(
        """# GPT-5.5 Comparison Deliverables

Generated new deliverables without overwriting the accepted camera-ready manuscript or existing final report.

## New files

- Document/outputs/camera-ready-manuscript-gpt55-comparison.docx
- Document/outputs/camera-ready-manuscript-gpt55-comparison.pdf
- Document/outputs/final-report-gpt55-comparison.docx
- Document/outputs/final-report-gpt55-comparison.pdf

## Source evidence

- GPT-5.5 comparison branch: codex/gpt-5.5-comparison
- GPT-5.5 comparison commit: d1baa7afc4c79a0177a7534235c55cf3c97d8a91
- GPT-5.4 SonarQube branch source: baseline-sonarqube-v1 at a2279bcbdc20751529335cf72f8baac1bc6f994b
- GPT-5.4 OWASP ZAP branch source: baseline-zap-v1 at 5f3bd3ccd9e0d460f52cac6a8a0d5fdceff1ce3d
- GPT-5.4 JMeter branch source: baseline-jmeter-v1, tested code commit 78b08b62570dcf6fe436354e8e6b770ab2074445. Branch head is 3c9392d8dcecf986da5438962e20c5d862e8be08 because a documentation commit was added after the performance fix.
- Source comparison report: Document/experiments/gpt-5.5/gpt-5.4-vs-gpt-5.5.md
- Pre-fix source commit: 7a0469d9401a3061941b44fcd46b3beca1c9c729

## Implementation policy

The existing files `camera-ready-manuscript.docx`, `camera-ready-manuscript.pdf`,
`final-report.docx`, and `final-report.pdf` were used as sources only. They were not
overwritten.
""",
        encoding="utf-8",
    )


def export_pdf(docx_path: Path, expected_pdf: Path) -> None:
    if not SOFFICE.exists():
        raise RuntimeError(f"LibreOffice not found at {SOFFICE}")
    profile_parent = ROOT / "tmp" / "docs"
    profile_parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=f"lo-profile-{docx_path.stem}-",
        dir=profile_parent,
        ignore_cleanup_errors=True,
    ) as profile_name:
        profile = Path(profile_name)
        subprocess.run(
            [
                str(SOFFICE),
                f"-env:UserInstallation=file:///{profile.resolve().as_posix()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(docx_path),
            ],
            check=True,
            cwd=ROOT,
        )
        generated = docx_path.with_suffix(".pdf")
        if generated != expected_pdf and generated.exists():
            if expected_pdf.exists():
                expected_pdf.unlink()
            generated.rename(expected_pdf)


def main() -> None:
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    add_camera_ready_addendum()
    add_full_report_addendum()
    write_notes()
    export_pdf(CAMERA_READY_DOCX, CAMERA_READY_PDF)
    export_pdf(FINAL_REPORT_DOCX, FINAL_REPORT_PDF)
    print(CAMERA_READY_DOCX)
    print(CAMERA_READY_PDF)
    print(FINAL_REPORT_DOCX)
    print(FINAL_REPORT_PDF)
    print(NOTES_MD)


if __name__ == "__main__":
    main()
