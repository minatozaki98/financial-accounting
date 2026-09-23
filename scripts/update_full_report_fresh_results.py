from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table

try:
    from scripts.format_final_report_docx import format_data_table, normalized_text
except ModuleNotFoundError:
    from format_final_report_docx import format_data_table, normalized_text


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-gpt55-comparison.docx"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def keyed(items: list[dict], key: str) -> dict[str, dict]:
    return {str(item[key]): item for item in items}


def load_fresh_evidence(root: Path) -> dict:
    evidence_root = root / "docs" / "appendix" / "verification-runs" / "research"

    sonar_baseline = read_json(evidence_root / "sonar" / "baseline" / "summary.json")
    sonar_remediation = read_json(evidence_root / "sonar" / "remediation" / "summary.json")
    sonar_issues = read_json(evidence_root / "sonar" / "baseline" / "issues.json")["issues"]
    impact_counts = Counter(
        impact["softwareQuality"]
        for issue in sonar_issues
        for impact in issue.get("impacts", [])
    )

    zap_baseline = read_json(evidence_root / "zap" / "baseline" / "summary.json")
    zap_remediation = read_json(evidence_root / "zap" / "remediation" / "summary.json")
    zap_before = keyed(zap_baseline["scans"], "scan")
    zap_after = keyed(zap_remediation["scans"], "scan")

    jmeter_baseline = read_json(evidence_root / "jmeter" / "baseline" / "summary.json")
    jmeter_remediation = read_json(evidence_root / "jmeter" / "remediation" / "summary.json")
    baseline_profiles = keyed(jmeter_baseline["profiles"], "profile")
    remediation_profiles = keyed(jmeter_remediation["profiles"], "profile")

    endpoint_statistics: dict[str, dict[str, dict[str, dict]]] = {
        "baseline": {},
        "remediation": {},
    }
    for role, profiles in (
        ("baseline", baseline_profiles),
        ("remediation", remediation_profiles),
    ):
        for profile in ("soak", "spike"):
            stats_path = root / profiles[profile]["statisticsPath"]
            endpoint_statistics[role][profile] = read_json(stats_path)

    evidence = {
        "sonar": {
            "baseline_issues": int(sonar_baseline["totalIssues"]),
            "remediation_issues": int(sonar_remediation["totalIssues"]),
            "baseline_coverage": sonar_baseline["metrics"]["coverage"],
            "remediation_coverage": sonar_remediation["metrics"]["coverage"],
            "baseline_code_smells": int(sonar_baseline["metrics"]["code_smells"]),
            "remediation_code_smells": int(sonar_remediation["metrics"]["code_smells"]),
            "baseline_bugs": int(sonar_baseline["metrics"]["bugs"]),
            "remediation_bugs": int(sonar_remediation["metrics"]["bugs"]),
            "baseline_vulnerabilities": int(sonar_baseline["metrics"]["vulnerabilities"]),
            "remediation_vulnerabilities": int(sonar_remediation["metrics"]["vulnerabilities"]),
            "baseline_hotspots": int(sonar_baseline["metrics"]["security_hotspots"]),
            "remediation_hotspots": int(sonar_remediation["metrics"]["security_hotspots"]),
            "baseline_duplication": sonar_baseline["metrics"]["duplicated_lines_density"],
            "remediation_duplication": sonar_remediation["metrics"]["duplicated_lines_density"],
            "baseline_complexity": int(sonar_baseline["metrics"]["complexity"]),
            "remediation_complexity": int(sonar_remediation["metrics"]["complexity"]),
            "baseline_cognitive_complexity": int(sonar_baseline["metrics"]["cognitive_complexity"]),
            "remediation_cognitive_complexity": int(sonar_remediation["metrics"]["cognitive_complexity"]),
            "baseline_gate": sonar_baseline["qualityGate"],
            "remediation_gate": sonar_remediation["qualityGate"],
            "impact_counts": dict(impact_counts),
            "tool_version": sonar_baseline["toolVersion"],
            "run_id": sonar_baseline["runId"],
        },
        "zap": {
            "baseline_passive": zap_before["baseline"],
            "remediation_passive": zap_after["baseline"],
            "baseline_api": zap_before["api"],
            "remediation_api": zap_after["api"],
            "remediation_api_medium": int(zap_after["api"]["medium"]),
            "baseline_gate": zap_baseline["gateStatus"],
            "remediation_gate": zap_remediation["gateStatus"],
            "tool_version": zap_baseline["toolVersion"],
        },
        "jmeter": {
            "baseline_profiles": baseline_profiles,
            "remediation_profiles": remediation_profiles,
            "baseline_gate": jmeter_baseline["gateStatus"],
            "remediation_gate": jmeter_remediation["gateStatus"],
            "gate_failures": jmeter_remediation["gateFailures"],
            "endpoint_statistics": endpoint_statistics,
            "tool_version": jmeter_baseline["toolVersion"],
            "baseline_run_id": jmeter_baseline["runId"],
            "remediation_run_id": jmeter_remediation["runId"],
        },
    }

    if evidence["sonar"]["baseline_issues"] != len(sonar_issues):
        raise RuntimeError("SonarQube summary and issue artifact counts disagree.")
    if evidence["sonar"]["baseline_code_smells"] != len(sonar_issues):
        raise RuntimeError("Fresh SonarQube legacy issue types are not all CODE_SMELL as expected.")
    return evidence


def find_paragraph_starting_with(doc: DocumentObject, prefix: str):
    matches = [p for p in doc.paragraphs if normalized_text(p).startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph starting with {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_paragraph(doc: DocumentObject, prefix: str, text: str) -> None:
    paragraph = find_paragraph_starting_with(doc, prefix)
    paragraph.clear()
    paragraph.add_run(text)


def table_header(table: Table) -> tuple[str, ...]:
    if not table.rows:
        return ()
    return tuple(" ".join(cell.text.split()) for cell in table.rows[0].cells)


def find_table_by_header(doc: DocumentObject, header: tuple[str, ...]) -> Table:
    matches = [table for table in doc.tables if table_header(table) == header]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one table with header {header!r}; found {len(matches)}")
    return matches[0]


def replace_table_rows(table: Table, rows: list[list[str]]) -> None:
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("All replacement rows must have the same number of cells.")
    if len(table.columns) != len(rows[0]):
        raise ValueError("Replacement column count does not match the existing table.")

    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)

    for row, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    format_data_table(table)


def percent_change(before: float, after: float) -> float:
    return ((after - before) / before) * 100.0


def fmt_ms(value: float) -> str:
    return f"{value:.2f} ms"


def update_report(report_path: Path, root: Path = ROOT) -> None:
    evidence = load_fresh_evidence(root)
    sonar = evidence["sonar"]
    zap = evidence["zap"]
    jmeter = evidence["jmeter"]
    before = jmeter["baseline_profiles"]
    after = jmeter["remediation_profiles"]
    endpoints = jmeter["endpoint_statistics"]

    doc = Document(report_path)
    abstract = find_paragraph_starting_with(doc, "This full report evaluates how ChatGPT")
    if "Fresh reproduction runs executed on September 22, 2026 provide the primary results." in abstract.text:
        return

    replace_paragraph(
        doc,
        "This full report evaluates how ChatGPT",
        "This full report evaluates how ChatGPT (Codex 5.4), used as an AI-assisted code improvement advisor, can improve the security, performance, and code quality of a financial accounting RESTful API implemented in ASP.NET Core. The study uses one existing API codebase and compares a common baseline branch with isolated SonarQube, OWASP ZAP, and JMeter remediation branches. Fresh reproduction runs executed on September 22, 2026 provide the primary results. SonarQube 26.7 reduced 17 unique baseline issues to 0 while the Quality Gate remained OK; legacy Code Smells fell from 17 to 0, and coverage changed slightly from 74.4% to 74.3%. OWASP ZAP 2.17 passed the configured rules: the passive scan changed from 3 Medium and 6 Low alerts to 0 Medium and 0 Low, while the authenticated API scan retained one Medium HTTP Only Site observation and reduced Low alerts from 3 to 0. Apache JMeter 5.5 produced mixed results. Soak and spike aggregate p95 improved, but the remediation core gate failed because 100 VU p95 reached 1146.9 ms and 500 VU p95 reached 1547.95 ms. The evidence therefore supports a strong code-quality result, a configured-scope security improvement with a disclosed transport observation, and an unsuccessful core-performance outcome under the fresh host run. Historical measurements are retained separately in the appendices for provenance rather than used as the current result.",
    )

    replace_table_rows(
        find_table_by_header(doc, ("Rule or area", "File", "Before", "After", "Status")),
        [
            ["Rule or area", "File", "Before", "After", "Status"],
            ["S107", "JournalEntriesController.cs", "Controller action used eight query parameters.", "Introduced JournalEntryQueryDto and passed one query model to the service.", "Code smell removed"],
            ["S2325", "ServiceManager.cs", "Private helper did not access instance state.", "Converted helper to static.", "Code smell removed"],
            ["S3267/S6602", "JournalEntryService.cs", "Loops and LINQ operations were less maintainable.", "Rewrote selected loops and lookup calls with clearer collection operations.", "Code smell removed"],
            ["S6966", "Program.cs", "Application startup used blocking Run.", "Changed startup to await app.RunAsync().", "Code smell removed"],
            ["Fresh summary", "C# project", f"{sonar['baseline_issues']} unique issues; {sonar['baseline_code_smells']} legacy Code Smells.", f"{sonar['remediation_issues']} issues; Quality Gate {sonar['remediation_gate']}.", "Closed"],
        ],
    )

    passive_before = zap["baseline_passive"]
    passive_after = zap["remediation_passive"]
    api_before = zap["baseline_api"]
    api_after = zap["remediation_api"]
    replace_table_rows(
        find_table_by_header(doc, ("Scan", "Severity", "Before -> After", "Vulnerability class", "Fix applied")),
        [
            ["Scan", "Severity", "Before -> After", "Vulnerability class", "Fix applied"],
            ["Passive baseline scan", "Medium", f"{passive_before['medium']} -> {passive_after['medium']}", "Security-header and browser-facing configuration alerts.", "Header middleware and Swagger hardening cleared the raw Medium alerts."],
            ["Passive baseline scan", "Low", f"{passive_before['low']} -> {passive_after['low']}", "Missing hardening and isolation headers.", "Added and consistently applied response-security headers."],
            ["Authenticated API scan", "Medium", f"{api_before['medium']} -> {api_after['medium']}", "HTTP Only Site (CWE-311) on the local HTTP test transport.", "Disclosed as a residual environment observation; not reported as cleared."],
            ["Authenticated API scan", "Low", f"{api_before['low']} -> {api_after['low']}", "Authenticated endpoint hardening findings.", "Same middleware path protects authenticated API responses."],
            ["Configured gate", "Result", f"{zap['baseline_gate']} -> {zap['remediation_gate']}", "Configured business-endpoint rule set.", "Gate passed, but raw Medium transport observation remains visible."],
        ],
    )

    profile_rows = [["Profile", "Metric", "Baseline", "After fixes", "Change", "Status"]]
    for profile in ("p50", "p100", "p500", "soak", "spike"):
        baseline_value = float(before[profile]["p95Ms"])
        remediation_value = float(after[profile]["p95Ms"])
        change = percent_change(baseline_value, remediation_value)
        status = "FAIL" if profile in {"p100", "p500"} else ("Regression" if change > 0 else "Improved")
        profile_rows.append(
            [
                {"p50": "50 VU", "p100": "100 VU", "p500": "500 VU"}.get(profile, profile),
                "Total p95 response time",
                fmt_ms(baseline_value),
                fmt_ms(remediation_value),
                f"{change:+.1f}%",
                status,
            ]
        )
    profile_rows.append(
        [
            "Core gate",
            "Threshold result",
            jmeter["baseline_gate"],
            jmeter["remediation_gate"],
            "100 VU and 500 VU exceeded thresholds",
            "NOT ACHIEVED",
        ]
    )
    replace_table_rows(
        find_table_by_header(doc, ("Profile", "Metric", "Baseline", "After fixes", "Change", "Status")),
        profile_rows,
    )

    replace_table_rows(
        find_table_by_header(doc, ("Fix class", "Evidence branch", "LLM rounds", "Human intervention", "Acceptance evidence")),
        [
            ["Fix class", "Evidence branch", "LLM rounds", "Human intervention", "Acceptance evidence"],
            ["SonarQube controller DTO refactor", "SonarQube remediation branch", "1", "Reviewed route compatibility and service-layer signature.", "Fresh issues reduced from 17 to 0; Quality Gate remained OK."],
            ["Security headers and Swagger hardening", "OWASP ZAP remediation branch", "1-2", "Moved middleware order, split Swagger document/UI switches, reviewed package upgrade.", "Configured gate passed; raw API HTTP Only Site Medium remained."],
            ["ARM/camelCase analyzer cleanup", "SonarQube remediation branch", "2-3", "Manually corrected provider-specific conventions after initial model output was insufficient.", "Fresh remediation scan reported 0 issues."],
            ["Report-performance remediation", "JMeter remediation branch", "1-2", "Reviewed cache/audit tradeoff and removed read-path report snapshot writes.", "Fresh 100 VU and 500 VU p95 exceeded thresholds; core gate failed."],
            ["Stress-profile interpretation", "JMeter remediation branch", "Human review", "Separated aggregate totals from endpoint-level and core-threshold behavior.", "Soak and spike aggregate p95 improved, but this did not offset the failed core gate."],
        ],
    )

    replace_table_rows(
        find_table_by_header(doc, ("Evaluation area", "Measure", "Acceptance target", "Result")),
        [
            ["Evaluation area", "Measure", "Acceptance target", "Result"],
            ["Static analysis", "SonarQube fresh issues and Quality Gate", "No retained blocker/critical issues after fixes; preserve Quality Gate OK", "17 unique issues reduced to 0; Quality Gate OK; coverage 74.4% -> 74.3%"],
            ["Security", "OWASP ZAP raw alerts and configured gate", "Clear actionable configured-rule alerts and disclose residual observations", "Configured gate passed; passive M3/L6 -> M0/L0; API M1/L3 -> M1/L0"],
            ["Performance", "JMeter p95, throughput, error rate, and gate", "50, 100, and 500 VU profiles pass with no material latency regression", "Remediation core gate FAIL; 100 VU and 500 VU exceeded thresholds"],
            ["Maintainability", "Representative before/after code evidence", "Selected code smells resolved with traceable branch evidence", "DTO, middleware-order, and report-persistence examples documented"],
        ],
    )

    replace_paragraph(
        doc,
        "Code-quality comparison units:",
        "Code-quality comparison units: fresh unique issue count, legacy issue type, SonarQube MQR software-quality impacts, rule ID, affected file, Quality Gate status, coverage percentage, duplicated-line density, and complexity. Legacy Bugs/Vulnerabilities/Code Smells and MQR Reliability/Security/Maintainability impacts are reported separately because the MQR impacts are non-exclusive and must not be summed as unique issues.",
    )
    replace_paragraph(
        doc,
        "Observed SonarQube result:",
        f"Observed fresh SonarQube result: {sonar['baseline_issues']} unique baseline issues decreased to {sonar['remediation_issues']}. Under the legacy type metrics, Bugs stayed at 0, Vulnerabilities stayed at 0, Code Smells decreased from {sonar['baseline_code_smells']} to {sonar['remediation_code_smells']}, and Security Hotspots stayed at 0. Under MQR impact classification, the baseline issues carried {sonar['impact_counts'].get('RELIABILITY', 0)} Reliability, {sonar['impact_counts'].get('SECURITY', 0)} Security, and {sonar['impact_counts'].get('MAINTAINABILITY', 0)} Maintainability impacts; these overlapping impact counts all fell to 0 after remediation. The Quality Gate remained {sonar['remediation_gate']}, coverage changed from {sonar['baseline_coverage']}% to {sonar['remediation_coverage']}%, and duplicated-line density remained {sonar['remediation_duplication']}%.",
    )

    sonar_rows = [
        ["Metric", "Baseline", "After fixes", "Interpretation"],
        ["Total unique issues", str(sonar["baseline_issues"]), str(sonar["remediation_issues"]), "All fresh baseline issues closed"],
        ["Code smells (legacy type)", str(sonar["baseline_code_smells"]), str(sonar["remediation_code_smells"]), "Legacy issue-type count"],
        ["Bugs (legacy type)", str(sonar["baseline_bugs"]), str(sonar["remediation_bugs"]), "No legacy BUG-type issues"],
        ["Vulnerabilities (legacy type)", str(sonar["baseline_vulnerabilities"]), str(sonar["remediation_vulnerabilities"]), "No legacy VULNERABILITY-type issues"],
        ["Reliability impacts (MQR)*", str(sonar["impact_counts"].get("RELIABILITY", 0)), "0", "Overlapping software-quality impact"],
        ["Security impacts (MQR)*", str(sonar["impact_counts"].get("SECURITY", 0)), "0", "Overlapping software-quality impact"],
        ["Maintainability impacts (MQR)*", str(sonar["impact_counts"].get("MAINTAINABILITY", 0)), "0", "Overlapping software-quality impact"],
        ["Security hotspots", str(sonar["baseline_hotspots"]), str(sonar["remediation_hotspots"]), "No hotspots reported"],
        ["Coverage", f"{sonar['baseline_coverage']}%", f"{sonar['remediation_coverage']}%", "Decreased by 0.1 percentage points"],
        ["Duplicated lines", f"{sonar['baseline_duplication']}%", f"{sonar['remediation_duplication']}%", "No change"],
        ["Cyclomatic complexity", str(sonar["baseline_complexity"]), str(sonar["remediation_complexity"]), "Increased by 3"],
        ["Cognitive complexity", str(sonar["baseline_cognitive_complexity"]), str(sonar["remediation_cognitive_complexity"]), "No change"],
        ["Quality Gate", sonar["baseline_gate"], sonar["remediation_gate"], "Remained OK"],
    ]
    replace_table_rows(
        find_table_by_header(doc, ("Metric", "Baseline", "After fixes", "Interpretation")),
        sonar_rows,
    )
    replace_paragraph(
        doc,
        "As shown in Table 4.2",
        "As shown in Table 4.2, all 17 unique fresh baseline issues were closed. Coverage decreased slightly by 0.1 percentage points and cyclomatic complexity increased by 3, so those controls are reported as small regressions rather than described as unchanged. The MQR Reliability, Security, and Maintainability impact counts overlap and therefore do not sum to the 17 unique issues.",
    )
    replace_paragraph(
        doc,
        "The SonarQube result is the cleanest closure case",
        "The fresh SonarQube result remains the cleanest closure case because all 17 unique issues moved to 0 and the Quality Gate remained OK. However, the evidence does not support saying every control metric was unchanged: coverage moved from 74.4% to 74.3% and cyclomatic complexity moved from 886 to 889. The most important maintainability change remains the replacement of the long controller parameter list with a request DTO.",
    )
    replace_paragraph(
        doc,
        "Outcome judgment: Achieved. The code-quality result",
        "Outcome judgment: Achieved with minor control-metric regressions. The fresh scan reduced 17 unique issues and 17 legacy Code Smells to 0, while the Quality Gate remained OK and duplicated-line density remained 0.0%. Legacy Bugs, Vulnerabilities, and Security Hotspots remained 0. Coverage decreased from 74.4% to 74.3%, and cyclomatic complexity increased from 886 to 889; these changes are small but are not described as unchanged. The implementation excerpts and fresh SonarQube dashboard evidence are reported in Appendices C and F.",
    )

    replace_paragraph(
        doc,
        "Observed authenticated API scan result:",
        f"Observed fresh authenticated API scan result: High alerts stayed at 0, Medium alerts stayed at {api_after['medium']}, Low alerts decreased from {api_before['low']} to {api_after['low']}, and Informational alerts stayed at {api_after['informational']}. The remaining Medium alert is HTTP Only Site (CWE-311) on the local HTTP test transport. It is disclosed as a raw scanner observation and was not counted as cleared.",
    )
    replace_paragraph(
        doc,
        "Observed ZAP baseline scan result:",
        f"Observed fresh passive-scan result: High alerts stayed at 0, Medium alerts decreased from {passive_before['medium']} to {passive_after['medium']}, Low alerts decreased from {passive_before['low']} to {passive_after['low']}, and Informational alerts changed from {passive_before['informational']} to {passive_after['informational']}. The configured-rule gate passed on both branches.",
    )
    replace_paragraph(
        doc,
        "The ZAP result shows that the vulnerable surface",
        "The fresh ZAP evidence shows effective header and configuration hardening, but it does not support a claim that every raw Medium alert was eliminated. The passive scan cleared all Medium and Low alerts, and the authenticated API scan cleared all Low alerts. One HTTP Only Site Medium observation remained because the local reproduction used HTTP rather than a production TLS endpoint.",
    )
    replace_paragraph(
        doc,
        "Outcome judgment: Achieved within the defined OWASP ZAP scope.",
        f"Outcome judgment: Partially Achieved in raw-alert terms and Achieved for the configured gate. Both branches returned {zap['remediation_gate']} for the configured business-endpoint rules. The passive scan improved from Medium {passive_before['medium']} / Low {passive_before['low']} to Medium {passive_after['medium']} / Low {passive_after['low']}; the authenticated API scan improved from Medium {api_before['medium']} / Low {api_before['low']} to Medium {api_after['medium']} / Low {api_after['low']}. Because one raw Medium HTTP Only Site observation remained, the report does not claim zero raw Medium alerts. The security-header implementation, scan provenance, and dashboards are provided in Appendices C, E, and F.",
    )

    replace_paragraph(
        doc,
        "Observed stress-profile result:",
        f"Observed fresh stress-profile result: soak total p95 decreased from {before['soak']['p95Ms']} ms to {after['soak']['p95Ms']} ms and spike total p95 decreased from {before['spike']['p95Ms']} ms to {after['spike']['p95Ms']} ms. Spike error rate also decreased from {before['spike']['errorPct']:.4f}% to {after['spike']['errorPct']:.2f}%. These aggregate improvements are reported separately from the failed core gate and do not override it.",
    )
    replace_paragraph(
        doc,
        "Observed JMeter result:",
        f"Observed fresh JMeter core result: 50 VU total p95 increased from {before['p50']['p95Ms']} ms to {after['p50']['p95Ms']} ms, 100 VU increased from {before['p100']['p95Ms']} ms to {after['p100']['p95Ms']} ms, and 500 VU increased from {before['p500']['p95Ms']} ms to {after['p500']['p95Ms']} ms. Error rate remained 0.00% for the three core profiles, but the remediation gate failed because 100 VU exceeded 500 ms and 500 VU exceeded 1200 ms.",
    )

    def endpoint(role: str, profile: str, label: str) -> float:
        return float(endpoints[role][profile][label]["pct2ResTime"])

    replace_table_rows(
        find_table_by_header(doc, ("Endpoint", "Profile metric", "Baseline", "After fixes", "Main remediation")),
        [
            ["Endpoint", "Profile metric", "Baseline", "After fixes", "Fresh interpretation"],
            ["GET /reports/account-ledger", "Soak p95", fmt_ms(endpoint("baseline", "soak", "GET /reports/account-ledger")), fmt_ms(endpoint("remediation", "soak", "GET /reports/account-ledger")), "Regression under the fresh sustained-load run"],
            ["GET /reports/account-ledger", "Spike p95", fmt_ms(endpoint("baseline", "spike", "GET /reports/account-ledger")), fmt_ms(endpoint("remediation", "spike", "GET /reports/account-ledger")), "Improved under spike, but remained the slowest report endpoint"],
            ["GET /reports/trial-balance", "Spike p95", fmt_ms(endpoint("baseline", "spike", "GET /reports/trial-balance")), fmt_ms(endpoint("remediation", "spike", "GET /reports/trial-balance")), "Large spike-profile improvement"],
            ["GET /reports/profit-loss", "Spike p95", fmt_ms(endpoint("baseline", "spike", "GET /reports/profit-loss")), fmt_ms(endpoint("remediation", "spike", "GET /reports/profit-loss")), "Large spike-profile improvement"],
            ["GET /reports/balance-sheet", "Spike p95", fmt_ms(endpoint("baseline", "spike", "GET /reports/balance-sheet")), fmt_ms(endpoint("remediation", "spike", "GET /reports/balance-sheet")), "Large spike-profile improvement"],
            ["POST /journal-entries/bulk", "Spike p95", fmt_ms(endpoint("baseline", "spike", "POST /journal-entries/bulk")), fmt_ms(endpoint("remediation", "spike", "POST /journal-entries/bulk")), "Large spike-profile improvement"],
        ],
    )
    replace_paragraph(
        doc,
        "The endpoint evidence in Table 4.4",
        "The fresh endpoint evidence in Table 4.4 shows that spike-profile report endpoints improved substantially, but account-ledger soak p95 regressed. Endpoint improvements under one workload therefore cannot be generalized to the failed 100 VU and 500 VU core profiles.",
    )
    replace_paragraph(
        doc,
        "The JMeter results show major p95 reductions",
        "The fresh JMeter results are mixed. Soak and spike aggregate p95 improved, and spike errors fell to 0%, but 50 VU, 100 VU, and 500 VU aggregate p95 all worsened. Because the pre-registered core gate failed, the report treats the current performance remediation as unsuccessful for the core criterion rather than carrying forward the earlier historical success claim.",
    )
    replace_paragraph(
        doc,
        "Outcome judgment: Achieved for core profiles",
        "Outcome judgment: Not Achieved for the core performance criterion and Partially Achieved for stress-profile behavior. The remediation core gate failed because 100 VU p95 was 1146.9 ms against a 500 ms threshold and 500 VU p95 was 1547.95 ms against a 1200 ms threshold. Soak total p95 improved from 443.95 ms to 171.0 ms, spike total p95 improved from 33401.75 ms to 4730.85 ms, and spike error rate fell from 0.4688% to 0.00%. Those stress improvements are meaningful but do not compensate for the failed core gate. Appendices D-F preserve the profiles, machine-readable statistics, and dashboards.",
    )

    replace_table_rows(
        find_table_by_header(doc, ("Tool", "Baseline evidence", "After-fix evidence", "Closure status")),
        [
            ["Tool", "Fresh baseline evidence", "Fresh after-fix evidence", "Closure status"],
            ["SonarQube", "17 unique issues; coverage 74.4%", "0 issues; coverage 74.3%; gate OK", "Achieved with minor metric regressions"],
            ["ZAP passive scan", "3 Medium, 6 Low", "0 Medium, 0 Low", "Achieved"],
            ["ZAP authenticated API", "1 Medium, 3 Low", "1 Medium, 0 Low", "Configured gate passed; raw Medium remains"],
            ["JMeter core profiles", "p95: 33 / 255.9 / 106 ms", "p95: 199.9 / 1146.9 / 1547.95 ms", "Not Achieved; remediation gate FAIL"],
            ["JMeter stress profiles", "Soak 443.95 ms; spike 33401.75 ms", "Soak 171 ms; spike 4730.85 ms", "Partially Achieved"],
        ],
    )
    replace_paragraph(
        doc,
        "Overall, the code-quality criterion",
        "Overall, the fresh evidence is mixed. The code-quality criterion was achieved with small coverage and complexity regressions. The ZAP configured gate was achieved, but one raw Medium HTTP-only transport observation remained. The JMeter core performance criterion was not achieved, although soak and spike aggregate behavior improved. The defensible conclusion is therefore that ChatGPT-guided remediation was effective for the measured static-analysis findings and selected security hardening, but performance benefits were not reproducible across the fresh core profiles.",
    )

    replace_paragraph(
        doc,
        "The SonarQube branch reduced retained C# findings",
        "Fresh reproduction showed that the SonarQube branch reduced 17 unique issues to 0 while the Quality Gate remained OK; coverage changed from 74.4% to 74.3%. The OWASP ZAP configured gate passed: the passive scan cleared 3 Medium and 6 Low alerts, while the authenticated API scan retained one Medium HTTP Only Site observation and cleared all 3 Low alerts. The JMeter remediation branch did not reproduce the earlier core-profile improvements: 50 VU p95 increased from 33.0 ms to 199.9 ms, 100 VU from 255.9 ms to 1146.9 ms, and 500 VU from 106.0 ms to 1547.95 ms, causing the core gate to fail. Soak and spike aggregate p95 improved, but the performance conclusion remains Not Achieved for the core criterion.",
    )
    replace_paragraph(
        doc,
        "The findings are limited to one ASP.NET Core API",
        "The findings are limited to one ASP.NET Core API, one language (C#), one database-backed financial accounting domain, one AI assistant configuration, and one fresh reproduction per evidence branch. Tool versions, host load, database state, and workload composition can materially change measured values. The fresh JMeter run demonstrates this limitation directly: stress aggregates improved while core p95 thresholds failed. Repeated runs with controlled host utilization and confidence intervals are required before claiming general performance improvement. Historical measurements remain in the appendices for provenance, but Chapters 4 and 5 use the September 22, 2026 fresh reproduction as the current result. Future work should compare multiple LLMs, repeat each measurement, and test production-like HTTPS deployment so the remaining ZAP HTTP Only Site observation can be evaluated under the intended transport.",
    )

    doc.save(report_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Update the latest full report from versioned fresh SonarQube, ZAP, and JMeter evidence."
    )
    parser.add_argument("report", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    report_path = args.report.resolve()
    update_report(report_path, ROOT)
    print(report_path)


if __name__ == "__main__":
    main()
