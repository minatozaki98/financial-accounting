from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-gpt55-comparison.docx"
DEFAULT_SOURCE = ROOT / "docs" / "appendix" / "thesis-appendix-source-code-and-results.md"
APPENDIX_MARKER = "APPENDICES"
IMAGE_WIDTH = Inches(6.1)

CROSS_REFERENCES = {
    "3.1.1 Testing Environment and Reproducibility": (
        "Appendices A-E provide the exact environment, branch and commit provenance, "
        "selected code listings, reproduction commands, and claim-to-artifact traceability."
    ),
    "4.2 Code-Quality Outcome": (
        "The implementation excerpts and the fresh SonarQube dashboard evidence are "
        "reported separately in Appendices C and F."
    ),
    "4.3 Security Outcome": (
        "The security-header implementation, scan provenance, and OWASP ZAP dashboards "
        "are provided in Appendices C, E, and F."
    ),
    "4.4 Performance Outcome": (
        "The caching implementation, JMeter profiles, and fresh dashboard results are "
        "provided in Appendices C, D, E, and F."
    ),
    "5.4 Limitations and Future Work": (
        "Appendix H records the interpretation limits for the fresh reproduction runs; "
        "these runs supplement rather than replace the historical Chapter 4 measurements."
    ),
}


def normalized_text(paragraph) -> str:
    return " ".join(paragraph.text.split())


def remove_existing_appendices(doc: DocumentObject) -> None:
    body = doc._element.body
    marker = next(
        (p for p in doc.paragraphs if normalized_text(p) == APPENDIX_MARKER),
        None,
    )
    if marker is None:
        return
    start = body.index(marker._p)
    for child in list(body)[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_cross_references(doc: DocumentObject) -> None:
    paragraphs = doc.paragraphs
    for heading_text, sentence in CROSS_REFERENCES.items():
        heading_index = next(
            (
                index
                for index, paragraph in enumerate(paragraphs)
                if normalized_text(paragraph) == heading_text
            ),
            None,
        )
        if heading_index is None:
            raise RuntimeError(f"Could not locate report section: {heading_text}")

        end = len(paragraphs)
        for index in range(heading_index + 1, len(paragraphs)):
            if paragraphs[index].style.name.startswith("Heading"):
                end = index
                break
        candidates = [
            paragraph
            for paragraph in paragraphs[heading_index + 1 : end]
            if paragraph.style.name == "Normal"
            and normalized_text(paragraph)
            and not normalized_text(paragraph).startswith(("Figure ", "Table "))
        ]
        target = candidates[-1] if candidates else paragraphs[heading_index]
        for paragraph in paragraphs:
            if paragraph._p is target._p or sentence not in paragraph.text:
                continue
            paragraph.text = " ".join(paragraph.text.replace(sentence, "").split())
        if sentence not in target.text:
            target.add_run(" " + sentence)


def ensure_code_style(doc: DocumentObject) -> str:
    style_name = "Appendix Code"
    if style_name not in [style.name for style in doc.styles]:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles[style_name]
    style.font.name = "Consolas"
    style.font.size = Pt(8)
    style.font.color.rgb = RGBColor(31, 41, 55)
    style.paragraph_format.left_indent = Inches(0.2)
    style.paragraph_format.right_indent = Inches(0.1)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    return style_name


def ensure_bullet_style(doc: DocumentObject) -> str:
    style_name = "Appendix Bullet"
    if style_name not in [style.name for style in doc.styles]:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles[style_name]
    style.base_style = doc.styles["List Paragraph"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(0.25)
    style.paragraph_format.first_line_indent = Inches(-0.15)
    style.paragraph_format.space_after = Pt(3)

    numbering = doc.part.numbering_part.element
    bullet_abstract_ids: set[str] = set()
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            number_format = level.find(qn("w:numFmt"))
            if (
                level.get(qn("w:ilvl")) == "0"
                and number_format is not None
                and number_format.get(qn("w:val")) == "bullet"
            ):
                bullet_abstract_ids.add(abstract.get(qn("w:abstractNumId")))
                break
    num_id = next(
        (
            num.get(qn("w:numId"))
            for num in numbering.findall(qn("w:num"))
            if num.find(qn("w:abstractNumId")).get(qn("w:val")) in bullet_abstract_ids
        ),
        None,
    )
    if num_id is None:
        raise RuntimeError("The report package does not contain a reusable bullet numbering definition.")

    properties = style.element.get_or_add_pPr()
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), num_id)
    num_properties.extend([level, identifier])
    properties.insert(0, num_properties)
    return style_name


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def add_inline_markdown(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def add_body_paragraph(doc: DocumentObject, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    add_inline_markdown(paragraph, text)
    return paragraph


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_table(doc: DocumentObject, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index in range(width):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            add_inline_markdown(paragraph, text)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(7.5)
                run.bold = row_index == 0
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph()


def dashboard_caption(image_path: Path, image_number: int) -> str:
    tool_names = {
        "sonarqube": "SonarQube",
        "zap": "OWASP ZAP",
        "jmeter": "Apache JMeter",
    }
    tool = tool_names.get(image_path.parent.name.lower(), image_path.parent.name)
    stem = image_path.stem
    for prefix in ("sonarqube-", "zap-", "jmeter-"):
        if stem.startswith(prefix):
            stem = stem[len(prefix) :]
            break
    description = stem.replace("-", " ")
    return f"Figure F.{image_number}. {tool} {description} dashboard."


def dashboard_heading_page_break(level: int, text: str, first_in_group: bool) -> bool:
    if level == 2:
        return True
    if level == 3 and text.endswith("Dashboards"):
        return text != "SonarQube Dashboards"
    if level == 4:
        return not first_in_group
    return False


def add_dashboard(doc: DocumentObject, source_dir: Path, target: str, number: int) -> None:
    image_path = (source_dir / target).resolve()
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(image_path), width=IMAGE_WIDTH)

    caption = doc.add_paragraph(dashboard_caption(image_path, number), style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_after = Pt(6)
    for run in caption.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def append_markdown(doc: DocumentObject, source_path: Path) -> int:
    lines = source_path.read_text(encoding="utf-8").splitlines()
    code_style = ensure_code_style(doc)
    bullet_style = ensure_bullet_style(doc)
    dashboard_count = 0
    index = 0
    in_code = False
    first_dashboard_in_group = False

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue

        if in_code:
            paragraph = doc.add_paragraph(style=code_style)
            paragraph.paragraph_format.keep_together = True
            set_cell_like = OxmlElement("w:shd")
            set_cell_like.set(qn("w:fill"), "F3F4F6")
            set_cell_like.set(qn("w:val"), "clear")
            paragraph._p.get_or_add_pPr().append(set_cell_like)
            run = paragraph.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            dashboard_count += 1
            add_dashboard(doc, source_path.parent, image_match.group(2), dashboard_count)
            index += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            if level == 1:
                index += 1
                continue
            style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}[level]
            paragraph = doc.add_paragraph(text, style=style)
            if dashboard_heading_page_break(level, text, first_dashboard_in_group):
                paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.keep_with_next = True
            if level == 3 and text.endswith("Dashboards"):
                first_dashboard_in_group = True
            elif level == 4:
                first_dashboard_in_group = False
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            rows = [parse_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(parse_table_row(lines[index]))
                index += 1
            add_markdown_table(doc, rows)
            continue

        if stripped.startswith("- "):
            add_body_paragraph(doc, stripped[2:], style=bullet_style)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "- ", "```", "|", "![")):
                break
            paragraph_lines.append(candidate)
            index += 1
        add_body_paragraph(doc, " ".join(paragraph_lines))

    return dashboard_count


def build_appendix(report_path: Path, source_path: Path) -> int:
    if not report_path.exists():
        raise FileNotFoundError(report_path)
    if not source_path.exists():
        raise FileNotFoundError(source_path)

    doc = Document(report_path)
    remove_existing_appendices(doc)
    add_cross_references(doc)

    heading = doc.add_paragraph(APPENDIX_MARKER, style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    dashboard_count = append_markdown(doc, source_path)
    if dashboard_count != 18:
        raise RuntimeError(f"Expected 18 dashboard images, found {dashboard_count}")

    doc.save(report_path)
    return dashboard_count


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Append the controlled thesis evidence source to the latest full-report DOCX."
    )
    parser.add_argument("report", nargs="?", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()

    report_path = args.report.resolve()
    source_path = args.source.resolve()
    count = build_appendix(report_path, source_path)
    print(f"Updated {report_path}")
    print(f"Embedded dashboard images: {count}")


if __name__ == "__main__":
    main()
