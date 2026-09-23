from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shape import InlineShape
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-23SEPv2-1.docx"
FIGURE_WIDTH = Inches(5.2)
PAIR_STARTS = {2, 4, 6}


def caption_for(doc: DocumentObject, figure_number: int):
    prefix = f"Figure 4.{figure_number}. "
    matches = [
        paragraph for paragraph in doc.paragraphs
        if paragraph.style.name == "Caption" and paragraph.text.startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one Figure 4.{figure_number} caption; found {len(matches)}")
    return matches[0]


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def pair_chapter4_figures(doc: DocumentObject) -> None:
    caption_for(doc, 1).paragraph_format.keep_with_next = False
    for figure_number in range(2, 8):
        caption = caption_for(doc, figure_number)
        paragraphs = doc.paragraphs
        caption_index = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is caption._p)
        if caption_index < 2:
            raise RuntimeError(f"Figure 4.{figure_number} has no preceding image and break")
        picture = paragraphs[caption_index - 1]
        inline_elements = picture._p.xpath(".//wp:inline")
        if len(inline_elements) != 1:
            raise RuntimeError(f"Figure 4.{figure_number} does not have exactly one inline image")
        preceding = paragraphs[caption_index - 2]
        page_breaks = preceding._p.xpath('.//w:br[@w:type="page"]')
        if page_breaks and (preceding.text.strip() or preceding._p.xpath(".//w:drawing")):
            raise RuntimeError(f"Figure 4.{figure_number} has a mixed-content page break")
        if figure_number in PAIR_STARTS:
            if len(page_breaks) != 1:
                raise RuntimeError(f"Figure 4.{figure_number} must start a new paired page")
            preceding.paragraph_format.space_before = Pt(0)
            preceding.paragraph_format.space_after = Pt(0)
        elif page_breaks:
            remove_paragraph(preceding)

        shape = InlineShape(inline_elements[0])
        original_width = int(shape.width)
        original_height = int(shape.height)
        shape.width = FIGURE_WIDTH
        shape.height = round(original_height * int(FIGURE_WIDTH) / original_width)
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        picture.paragraph_format.space_before = Pt(0)
        picture.paragraph_format.space_after = Pt(0)
        picture.paragraph_format.line_spacing = None
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = False
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(8)


def pair_report_figures(path: Path) -> None:
    doc = Document(path)
    pair_chapter4_figures(doc)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Lay out Chapter 4 Figures 4.2-4.7 two per page.")
    parser.add_argument("target", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    target = args.target.resolve()
    pair_report_figures(target)
    print(target)


if __name__ == "__main__":
    main()
