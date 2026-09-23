from __future__ import annotations

import argparse
import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "Document" / "outputs" / "final-report-gpt55-comparison.docx"
BODY_STYLES = {"Normal", "List Paragraph"}
CHAPTER_PATTERN = re.compile(r"^Chapter\s+\d+", re.IGNORECASE)
SECTION_PATTERN = re.compile(r"^\d+\.\d+(?:\.\d+)?\s")


def load_report_formatter():
    try:
        import format_final_report_docx

        return format_final_report_docx
    except ModuleNotFoundError:
        path = Path(__file__).with_name("format_final_report_docx.py")
        spec = importlib.util.spec_from_file_location("format_final_report_docx", path)
        if spec is None or spec.loader is None:
            raise RuntimeError(f"Could not load {path}")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module


report_formatter = load_report_formatter()


def normalized(text: str) -> str:
    return " ".join(text.split())


def has_page_break(paragraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def is_empty_paragraph(paragraph) -> bool:
    return (
        not paragraph.text.strip()
        and not has_drawing(paragraph)
        and not has_page_break(paragraph)
        and not paragraph._p.xpath(".//w:fldChar")
    )


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def remove_page_breaks(paragraph) -> None:
    for page_break in list(paragraph._p.xpath('.//w:br[@w:type="page"]')):
        parent = page_break.getparent()
        if parent is not None:
            parent.remove(page_break)


def set_cell_margins(cell, *, top: int = 72, start: int = 90, bottom: int = 72, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        margin = tc_mar.find(qn(f"w:{margin_name}"))
        if margin is None:
            margin = OxmlElement(f"w:{margin_name}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def is_signature_table(table) -> bool:
    if not table.rows:
        return False
    first_row = [normalized(cell.text) for cell in table.rows[0].cells]
    return not any(first_row) or all(not text or set(text) <= {"_"} for text in first_row)


def is_code_table(table) -> bool:
    return len(table.rows) == 1 and len(table.columns) == 1


def table_header(table) -> tuple[str, ...]:
    if not table.rows:
        return ()
    return tuple(normalized(cell.text) for cell in table.rows[0].cells)


def set_fixed_table_widths(table, widths) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    for grid_column, width in zip(grid_columns, widths, strict=True):
        grid_column.set(qn("w:w"), str(width.twips))

    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(width.twips))
            tc_width.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def normalize_source_index_table(table) -> None:
    header = table_header(table)
    if header == ("Listing", "Purpose", "File", "Branch/ref", "Commit"):
        for row_index, row in enumerate(table.rows):
            cells = row.cells
            cells[0].text = (
                "Listing and purpose"
                if row_index == 0
                else f"{normalized(cells[0].text)} - {normalized(cells[1].text)}"
            )
            row._tr.remove(cells[1]._tc)
        grid_columns = table._tbl.tblGrid.gridCol_lst
        table._tbl.tblGrid.remove(grid_columns[1])

    if table_header(table) == ("Listing and purpose", "File", "Branch/ref", "Commit"):
        set_fixed_table_widths(
            table,
            (Inches(1.8), Inches(2.65), Inches(1.3), Inches(0.75)),
        )
        for row in table.rows:
            prevent_row_split(row)


def format_code_table(table) -> None:
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=72, start=108, bottom=72, end=108)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
            run.bold = False


def format_data_tables(doc: DocumentObject) -> None:
    section = doc.sections[0]
    usable_twips = round(
        (section.page_width - section.left_margin - section.right_margin) / 635
    )
    for table in doc.tables:
        normalize_source_index_table(table)
        if is_signature_table(table):
            continue
        if is_code_table(table):
            format_code_table(table)
            continue

        report_formatter.format_data_table(table)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.first_line_indent = Inches(0)
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        technical = run.font.name == "Courier New"
                        run.font.name = "Courier New" if technical else "Times New Roman"
                        run.font.size = Pt(10)
                        run.bold = row_index == 0
        if table_header(table) == ("Listing and purpose", "File", "Branch/ref", "Commit"):
            normalize_source_index_table(table)
        else:
            grid = list(table._tbl.tblGrid.gridCol_lst)
            widths = [int(column.get(qn("w:w"))) for column in grid]
            if sum(widths) > usable_twips:
                scale = usable_twips / sum(widths)
                widths = [round(width * scale) for width in widths]
                widths[-1] += usable_twips - sum(widths)
            set_fixed_table_widths(table, tuple(Twips(width) for width in widths))


def body_bounds(doc: DocumentObject) -> tuple[int, int]:
    paragraphs = doc.paragraphs
    chapter_one = next(
        i for i, p in enumerate(paragraphs)
        if normalized(p.text) == "Chapter 1 - Introduction"
        and p.style.name == "Heading 1"
    )
    appendices = next(
        i for i, p in enumerate(paragraphs)
        if normalized(p.text) == "APPENDICES"
        and p.style.name == "Heading 1"
    )
    return chapter_one, appendices


def ensure_section_3_1(doc: DocumentObject) -> None:
    if any(normalized(p.text) == "3.1 Research Design and Workflow" for p in doc.paragraphs):
        return
    chapter_three = next(
        p for p in doc.paragraphs
        if normalized(p.text) == "Chapter 3: Proposed Methodology"
        and p.style.name == "Heading 1"
    )
    caption = next(
        p for p in doc.paragraphs
        if normalized(p.text).startswith("Table 3.1.")
        and p.style.name == "Caption"
    )
    if chapter_three._p.getparent() is not caption._p.getparent():
        raise RuntimeError("Chapter 3 and Table 3.1 are not in the same document body")
    new_p = OxmlElement("w:p")
    caption._p.addprevious(new_p)
    heading = Paragraph(new_p, caption._parent)
    heading.style = "Heading 2"
    heading.add_run("3.1 Research Design and Workflow")
    heading.paragraph_format.keep_with_next = True


def normalize_heading_pagination(doc: DocumentObject) -> None:
    start, _ = body_bounds(doc)
    for paragraph in list(doc.paragraphs[start:]):
        text = normalized(paragraph.text)
        if CHAPTER_PATTERN.match(text):
            paragraph.text = text
            paragraph.style = "Heading 1"
            paragraph.paragraph_format.page_break_before = True
        elif text in {"REFERENCES", "References", "APPENDICES"}:
            paragraph.style = "Heading 1"
            paragraph.paragraph_format.page_break_before = True
        elif SECTION_PATTERN.match(text):
            remove_page_breaks(paragraph)

    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs[:-1]):
        if not has_page_break(paragraph) or paragraph.text.strip() or has_drawing(paragraph):
            continue
        next_text = normalized(paragraphs[index + 1].text)
        if CHAPTER_PATTERN.match(next_text) or next_text in {"REFERENCES", "References", "APPENDICES"}:
            remove_paragraph(paragraph)
        elif SECTION_PATTERN.match(next_text):
            remove_paragraph(paragraph)

    paragraphs = list(doc.paragraphs)
    for index, paragraph in enumerate(paragraphs[1:], start=1):
        if not paragraph.paragraph_format.page_break_before:
            continue
        previous = paragraphs[index - 1]
        if is_empty_paragraph(previous):
            remove_paragraph(previous)


def normalize_body_paragraphs(doc: DocumentObject) -> None:
    start, end = body_bounds(doc)
    for paragraph in doc.paragraphs[start:end]:
        text = normalized(paragraph.text)
        if not text or paragraph.style.name not in BODY_STYLES:
            continue

        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            if run.font.name != "Courier New":
                run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        visible_runs = [run for run in paragraph.runs if run.text.strip()]
        if len(text) >= 55 and visible_runs and all(run.bold is True for run in visible_runs):
            for run in visible_runs:
                run.bold = False


def release_unnecessary_image_page_breaks(doc: DocumentObject) -> None:
    start, end = body_bounds(doc)
    for paragraph in doc.paragraphs[start:end]:
        if has_drawing(paragraph) and paragraph.paragraph_format.page_break_before:
            paragraph.paragraph_format.page_break_before = False


def collapse_repeated_blank_paragraphs(doc: DocumentObject) -> None:
    start, end = body_bounds(doc)
    body = list(doc.paragraphs[start:end])
    previous_empty = False
    for paragraph in body:
        empty = is_empty_paragraph(paragraph)
        if empty and previous_empty:
            remove_paragraph(paragraph)
        previous_empty = empty


def normalize_report(path: Path) -> None:
    doc = Document(path)
    report_formatter.configure_caption_style(doc)
    report_formatter.configure_heading_styles(doc)
    report_formatter.apply_heading_styles(doc)
    ensure_section_3_1(doc)
    normalize_heading_pagination(doc)
    report_formatter.normalize_heading_runs(doc)
    normalize_body_paragraphs(doc)
    release_unnecessary_image_page_breaks(doc)
    collapse_repeated_blank_paragraphs(doc)
    format_data_tables(doc)
    report_formatter.enforce_document_typography(doc)
    report_formatter.enable_field_update_on_open(doc)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize tables, paragraph emphasis, and pagination in the full thesis report.")
    parser.add_argument("target", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    target = args.target.resolve()
    normalize_report(target)
    print(target)


if __name__ == "__main__":
    main()
