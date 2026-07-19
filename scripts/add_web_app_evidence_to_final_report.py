from __future__ import annotations

import subprocess
import shutil
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = ROOT / "Document" / "outputs"
FINAL_REPORT_DOCX = OUTPUTS / "final-report.docx"
FINAL_REPORT_PDF = OUTPUTS / "final-report.pdf"
EVIDENCE_DIR = OUTPUTS / "web-app-evidence"
FIGURE_WIDTH = Inches(6.1)
MAX_VISIBLE_SCREENSHOT_HEIGHT = 1000
SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.com")


def insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def append_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def normalized_text(paragraph: Paragraph) -> str:
    return " ".join(paragraph.text.split())


def is_page_break_only(paragraph: Paragraph) -> bool:
    return (
        not normalized_text(paragraph)
        and bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))
        and not bool(paragraph._p.xpath(".//w:drawing"))
    )


def is_empty_spacer(paragraph: Paragraph) -> bool:
    return (
        not normalized_text(paragraph)
        and not bool(paragraph._p.xpath(".//w:drawing"))
        and not bool(paragraph._p.xpath(".//w:sectPr"))
    )


def remove_existing_web_sections(doc: Document) -> None:
    paragraphs = doc.paragraphs
    remove_ranges: list[tuple[int, int]] = []
    for index, paragraph in enumerate(paragraphs):
        if normalized_text(paragraph).startswith("4.6 Web Application Evidence for Current APIs"):
            start = index
            while start > 0 and (
                is_page_break_only(paragraphs[start - 1])
                or is_empty_spacer(paragraphs[start - 1])
            ):
                start -= 1
            end = index + 1
            while end < len(paragraphs) and not normalized_text(paragraphs[end]).startswith("Chapter 5: Conclusion"):
                end += 1
            remove_ranges.append((start, end))

    for start, end in reversed(remove_ranges):
        for paragraph in paragraphs[start:end]:
            parent = paragraph._element.getparent()
            if parent is not None:
                parent.remove(paragraph._element)


def find_chapter_5(doc: Document) -> Paragraph:
    matches = [
        paragraph
        for paragraph in doc.paragraphs
        if normalized_text(paragraph).startswith("Chapter 5: Conclusion")
    ]
    if not matches:
        raise RuntimeError("Could not find Chapter 5 insertion point.")
    return matches[-1]


def add_picture_after(paragraph: Paragraph, image_path: Path, caption: str, page_break_before: bool = False) -> Paragraph:
    if page_break_before:
        break_para = append_after(paragraph)
        break_para.add_run().add_break(WD_BREAK.PAGE)
        paragraph = break_para

    picture_para = append_after(paragraph)
    run = picture_para.add_run()
    with Image.open(image_path) as image:
        visible_height = min(image.height, MAX_VISIBLE_SCREENSHOT_HEIGHT)
        displayed_height = int(FIGURE_WIDTH * visible_height / image.width)
        crop_bottom = int((image.height - visible_height) * 100000 / image.height)

    picture = run.add_picture(str(image_path), width=FIGURE_WIDTH, height=displayed_height)
    if crop_bottom:
        source_rect = OxmlElement("a:srcRect")
        source_rect.set("b", str(crop_bottom))
        picture._inline.graphic.graphicData.pic.blipFill.insert(1, source_rect)

    caption_para = append_after(picture_para, caption)
    for run in caption_para.runs:
        run.bold = True
    return caption_para


def export_pdf() -> None:
    subprocess.run(
        [
            str(SOFFICE),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUTPUTS),
            str(FINAL_REPORT_DOCX),
        ],
        check=True,
        cwd=ROOT,
    )


def repair_docx_for_word() -> None:
    """Round-trip the report through LibreOffice so Microsoft Word accepts the package."""
    with tempfile.TemporaryDirectory(prefix="final-report-word-", dir=OUTPUTS) as workspace:
        workspace_path = Path(workspace)
        input_dir = workspace_path / "input"
        output_dir = workspace_path / "output"
        input_dir.mkdir()
        output_dir.mkdir()
        input_path = input_dir / FINAL_REPORT_DOCX.name
        shutil.copy2(FINAL_REPORT_DOCX, input_path)

        subprocess.run(
            [
                str(SOFFICE),
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(output_dir),
                str(input_path),
            ],
            check=True,
            cwd=ROOT,
        )

        repaired_path = output_dir / FINAL_REPORT_DOCX.name
        if not repaired_path.exists():
            raise RuntimeError("LibreOffice did not produce a Word-compatible final report.")
        shutil.copy2(repaired_path, FINAL_REPORT_DOCX)


def main() -> None:
    doc = Document(FINAL_REPORT_DOCX)
    remove_existing_web_sections(doc)
    chapter_5 = find_chapter_5(doc)
    chapter_5.paragraph_format.page_break_before = True

    anchor = insert_paragraph_before(chapter_5, "4.6 Web Application Evidence for Current APIs")
    anchor.runs[0].bold = True
    anchor = append_after(
        anchor,
        "A React and Vite web application was added as a live validation artifact for the current ASP.NET Core financial accounting APIs. The application runs at http://localhost:5173 and connects to the API at http://localhost:5296, matching the CORS origins already configured in the backend. It demonstrates JWT login, role-aware navigation, account balances, expandable journal-entry debit and credit lines, accounting periods, reports, audit logs, user administration, and a research-evidence dashboard.",
    )
    anchor = append_after(
        anchor,
        "The local seed process creates four demo accounts for reproducible role testing: admin, finance-manager, normal-user, and auditor-user, all using the local research password Admin@123. The screenshots below were captured after seeding the local SQL Server database and logging in through the web application, so they validate the user-visible path rather than only the API response layer.",
    )

    screenshots = [
        ("01-login.png", "Figure 4.6.1 Web application login screen with demo role credentials."),
        ("02-dashboard-admin.png", "Figure 4.6.2 Admin dashboard showing live API counts and report summary."),
        ("03-reports-trial-balance.png", "Figure 4.6.3 Trial balance report generated from the live reports API."),
        ("04-audit-logs.png", "Figure 4.6.4 Audit-log screen available to Admin and Auditor roles."),
        ("05-auditor-role-view.png", "Figure 4.6.5 Auditor role view with restricted navigation."),
        ("06-research-evidence.png", "Figure 4.6.6 Research-evidence page mapping screens to API coverage."),
        ("07-journal-entry-details.png", "Figure 4.6.7 Expanded journal-entry lines showing account-level debit and credit evidence."),
    ]

    for index, (filename, caption) in enumerate(screenshots):
        image_path = EVIDENCE_DIR / filename
        if not image_path.exists():
            raise FileNotFoundError(image_path)
        anchor = add_picture_after(anchor, image_path, caption, page_break_before=index > 0)

    doc.save(FINAL_REPORT_DOCX)
    repair_docx_for_word()
    export_pdf()
    print(FINAL_REPORT_DOCX)
    print(FINAL_REPORT_PDF)


if __name__ == "__main__":
    main()
